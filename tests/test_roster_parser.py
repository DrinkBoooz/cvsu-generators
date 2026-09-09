import os
import csv
import pytest
import roster_parser
import process_schedule
from executable.main import ScriptAPI

def test_ceit_prefix_directory_completeness():
    """Verify all 15 official CEIT prefixes and department mappings are present."""
    expected_prefixes = [
        "AGEN", "ABEN", "ARCH", "CENG", "CIVL", "COSC", "CPEN",
        "DCEE", "DCIT", "ECEN", "EENG", "IENG", "INDT", "SMT", "ITEC"
    ]
    for prefix in expected_prefixes:
        assert prefix in roster_parser.CEIT_PREFIX_MAP, f"Missing CEIT prefix: {prefix}"
        meta = roster_parser.CEIT_PREFIX_MAP[prefix]
        assert "dept" in meta
        assert "dept_code" in meta
        assert "badge" in meta
        assert "name" in meta

def test_get_prefix_metadata_extraction():
    """Verify prefix extraction from subjects, filenames, and courses."""
    # Direct prefix match
    meta = roster_parser.get_prefix_metadata("COSC 101 - ADVANCED DATABASE")
    assert meta is not None
    assert meta["prefix"] == "COSC"
    assert meta["dept_code"] == "DIT"

    # From filename
    meta_arch = roster_parser.get_prefix_metadata("BSARCH1-1 List of Students for 202612099-ARCH 11 - THEORY.xlsx")
    assert meta_arch is not None
    assert meta_arch["prefix"] == "ARCH"
    assert meta_arch["dept_code"] == "DCE"

    # From custom name
    meta_cpen = roster_parser.get_prefix_metadata("students_CPEN_50_final.csv")
    assert meta_cpen is not None
    assert meta_cpen["prefix"] == "CPEN"
    assert meta_cpen["dept_code"] == "DCEE"

    # Unknown
    assert roster_parser.get_prefix_metadata("RANDOM_FILE_NO_PREFIX.csv") is None

def test_inspect_roster_csv_auto_and_overrides(tmp_path):
    """Test inspecting raw CSV rows, auto-detecting headers, and applying manual overrides."""
    csv_file = tmp_path / "custom_students.csv"
    csv_file.write_text(
        "COL_EXTRA,STUDENT_IDENTIFIER,STUDENT_FULL_NAME,REMARKS\n"
        "X1,20261001,\"DE LA CRUZ, JUAN\",Regular\n"
        "X2,20261002,\"DELOS REYES, MARIA\",Regular\n",
        encoding="utf-8"
    )

    # Auto inspection
    res = roster_parser.inspect_roster(str(csv_file))
    assert res["status"] == "success"
    assert len(res["columns"]) == 4
    assert res["student_count"] == 2
    assert res["active_name_col"] == 2  # STUDENT_FULL_NAME
    assert res["active_id_col"] == 1    # STUDENT_IDENTIFIER
    assert res["active_header_row"] == 0

    # Inspection with manual overrides (e.g. user overrides columns)
    overrides = {
        "name_col": 2,
        "id_col": 0,  # force extra col as ID
        "header_row": 0
    }
    res_ovr = roster_parser.inspect_roster(str(csv_file), overrides=overrides)
    assert res_ovr["status"] == "success"
    assert res_ovr["active_id_col"] == 0
    assert res_ovr["parsed_preview"][0][1] == "X1"

def test_load_students_with_overrides(tmp_path):
    """Test loading students with custom column/row overrides."""
    csv_file = tmp_path / "custom_layout.csv"
    csv_file.write_text(
        "HEADER ROW 0 - TITLE\n"
        "HEADER ROW 1 - METADATA\n"
        "CustomID,CustomName,Note\n"
        "20269999,\"DELA CRUZ, JUAN P.\",Active\n"
        "20268888,\"SANTOS, MARIA B.\",Active\n",
        encoding="utf-8"
    )

    students = roster_parser.load_students(
        str(csv_file),
        name_col=1,
        id_col=0,
        header_row=2
    )
    assert len(students) == 2
    assert students[0] == ("DELA CRUZ, JUAN P.", "20269999")
    assert students[1] == ("SANTOS, MARIA B.", "20268888")

def test_validate_rosters_with_manual_linking(tmp_path):
    """Test validate_rosters with non-standard roster filenames linked via roster_configs."""
    schedule_path = os.path.join(os.path.dirname(__file__), "..", "ORTEGA_SCHEDULE.xls")
    
    # Non-standard filename that fails default regex
    odd_file = tmp_path / "CS1-4_Official_Roster.csv"
    odd_file.write_text("Student ID,Full Name\n20261001,\"CRUZ, JUAN\"\n", encoding="utf-8")

    # Without configs: incomplete_filename, but has can_link: True and recommended_filename
    res_unlinked = process_schedule.validate_rosters(schedule_path, [str(odd_file)])
    assert len(res_unlinked) == 1
    assert res_unlinked[0]["status"] == "error"
    assert res_unlinked[0]["issue"] == "incomplete_filename"
    assert res_unlinked[0]["can_link"] is True
    assert "recommended_filename" in res_unlinked[0]

    # With configs: user manually links it to schedule code 202612040
    roster_configs = {
        str(odd_file): {
            "linked_schedule_code": "202612040",
            "course_sec": "CS1-4",
            "subject_name": "DCIT 21A",
            "name_col": 1,
            "id_col": 0,
            "header_row": 0
        }
    }
    res_linked = process_schedule.validate_rosters(schedule_path, [str(odd_file)], roster_configs=roster_configs)
    assert len(res_linked) == 1
    assert res_linked[0]["status"] == "valid"
    assert "custom_linked" in res_linked[0]["issue"]
    assert res_linked[0]["student_count"] == 1
    assert res_linked[0]["course_sec"] == "CS1-4"
    assert res_linked[0]["schedule_code"] == "202612040"
    assert res_linked[0]["ceit_metadata"]["prefix"] == "DCIT"

def test_script_api_inspect_and_ceit_directory(tmp_path):
    """Test ScriptAPI inspect_roster and get_ceit_prefix_directory bridge methods."""
    api = ScriptAPI()
    ceit_dir = api.get_ceit_prefix_directory()
    assert "COSC" in ceit_dir
    assert "CPEN" in ceit_dir
    assert "AGEN" in ceit_dir
    assert "IENG" in ceit_dir

    sample_csv = tmp_path / "BSCS1-1 List of Students for 202612040-DCIT 21A.csv"
    sample_csv.write_text("Student Number,Student Name\n20261001,JUAN DELA CRUZ\n", encoding="utf-8")

    api.rosters = [str(sample_csv)]
    inspect_res = api.inspect_roster(os.path.basename(str(sample_csv)))
    assert inspect_res["status"] == "success"
    assert inspect_res["student_count"] == 1
    assert inspect_res["active_name_col"] == 1
    assert inspect_res["active_id_col"] == 0

def test_roster_xlsx_with_multiple_columns_and_multi_row_headers(tmp_path):
    """Verify Excel roster with 10 columns and 3 metadata header rows before table headers."""
    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Class Roster"

    # 3 metadata / banner rows before the column headers
    ws.append(["CAVITE STATE UNIVERSITY", "", "", "", "", "", "", "", "", ""])
    ws.append(["College of Engineering and Information Technology", "", "", "", "", "", "", "", "", ""])
    ws.append(["First Semester, Academic Year 2026-2027", "", "", "", "", "", "", "", "", ""])

    # 10 column headers at row index 3
    headers = ["#", "Student ID", "Full Name", "Sex", "Course", "Year", "Section", "Email Address", "Contact Number", "Remarks"]
    ws.append(headers)

    # 3 student rows
    ws.append(["1", "202610001", "DELA CRUZ, JUAN A.", "M", "BSCS", "1", "4", "juan@cvsu.edu.ph", "09123456789", "Regular"])
    ws.append(["2", "202610002", "SANTOS, MARIA B.", "F", "BSCS", "1", "4", "maria@cvsu.edu.ph", "09987654321", "Regular"])
    ws.append(["3", "202610003", "REYES, CARLOS C.", "M", "BSCS", "1", "4", "carlos@cvsu.edu.ph", "09112233445", "Irregular"])

    xlsx_path = tmp_path / "multi_header_10_cols.xlsx"
    wb.save(str(xlsx_path))

    # 1. Automatic header and column detection
    res = roster_parser.inspect_roster(str(xlsx_path))
    assert res["status"] == "success"
    assert res["detected_header_row"] == 3
    assert res["active_header_row"] == 3
    assert res["active_id_col"] == "B"
    assert res["active_name_col"] == "C"
    assert len(res["columns"]) == 10
    assert res["columns"] == ["A", "B", "C", "D", "E", "F", "G", "H", "I", "J"]
    assert len(res["available_columns"]) == 10
    assert res["available_columns"][0]["name"] == "Col A: #"
    assert res["available_columns"][1]["name"] == "Col B: Student ID"
    assert res["available_columns"][2]["name"] == "Col C: Full Name"
    assert res["student_count"] == 3
    assert res["parsed_students"][0]["name"] == "DELA CRUZ, JUAN A."
    assert res["parsed_students"][0]["student_number"] == "202610001"

    # 2. Unified load_students directly
    students = roster_parser.load_students(str(xlsx_path))
    assert len(students) == 3
    assert students[0] == ("DELA CRUZ, JUAN A.", "202610001")
    assert students[1] == ("SANTOS, MARIA B.", "202610002")
    assert students[2] == ("REYES, CARLOS C.", "202610003")

    # 3. Manual override on 10-column layout
    overrides = {
        "header_row": 3,
        "name_col": "C",
        "id_col": "A"  # user maps column A (#) as ID
    }
    res_ovr = roster_parser.inspect_roster(str(xlsx_path), overrides=overrides)
    assert res_ovr["status"] == "success"
    assert res_ovr["active_header_row"] == 3
    assert res_ovr["active_id_col"] == "A"
    assert res_ovr["parsed_students"][0]["student_number"] == "1"
    assert res_ovr["parsed_students"][1]["student_number"] == "2"

def test_roster_xlsx_with_stacked_two_tier_headers(tmp_path):
    """Verify Excel roster with 2-tier stacked headers (category headers on row 0, column headers on row 1)."""
    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active

    # Row 0: Group / Category headers
    ws.append(["STUDENT INFORMATION", "", "", "ACADEMIC ENROLLMENT", "", "", "COMMUNICATION", "", ""])
    # Row 1: Sub / Column headers
    ws.append(["#", "Student Number", "Full Name", "Degree / Program", "Year", "Section", "Email Address", "Mobile No.", "Remarks"])
    # Student rows
    ws.append(["1", "202610001", "DELA CRUZ, JUAN A.", "BSCS", "1", "4", "juan@cvsu.edu.ph", "09123456789", "Regular"])
    ws.append(["2", "202610002", "SANTOS, MARIA B.", "BSCS", "1", "4", "maria@cvsu.edu.ph", "09987654321", "Regular"])

    path = tmp_path / "stacked_headers.xlsx"
    wb.save(str(path))

    res = roster_parser.inspect_roster(str(path))
    assert res["status"] == "success"
    assert res["detected_header_row"] == 1
    assert res["active_id_col"] == "B"
    assert res["active_name_col"] == "C"
    assert len(res["columns"]) == 9
    assert res["student_count"] == 2
    assert res["parsed_students"][0]["name"] == "DELA CRUZ, JUAN A."
    assert res["parsed_students"][0]["student_number"] == "202610001"

def test_roster_csv_with_multiple_columns_and_multi_row_headers(tmp_path):
    """Verify CSV roster with 9 columns and 2 banner metadata rows before headers."""
    csv_file = tmp_path / "multi_header_9_cols.csv"
    csv_file.write_text(
        "CAVITE STATE UNIVERSITY - MAIN CAMPUS\n"
        "OFFICE OF THE UNIVERSITY REGISTRAR - OFFICIAL ROSTER\n"
        "Index,Department,Student ID,Degree,Student Full Name,Gender,YearLevel,Section,Status\n"
        "1,DIT,20261001,BSCS,\"DELA CRUZ, JUAN A.\",M,1,4,Enrolled\n"
        "2,DIT,20261002,BSCS,\"SANTOS, MARIA B.\",F,1,4,Enrolled\n",
        encoding="utf-8"
    )

    res = roster_parser.inspect_roster(str(csv_file))
    assert res["status"] == "success"
    assert res["detected_header_row"] == 2
    assert res["active_header_row"] == 2
    assert res["active_id_col"] == 2  # Student ID is Col 2
    assert res["active_name_col"] == 4  # Student Full Name is Col 4
    assert len(res["columns"]) == 9
    assert res["student_count"] == 2
    assert res["parsed_students"][0]["name"] == "DELA CRUZ, JUAN A."
    assert res["parsed_students"][0]["student_number"] == "20261001"

def test_process_schedule_with_multi_column_multi_header_roster(tmp_path):
    """Verify end-to-end processing with a 10-column multi-header XLSX roster."""
    import openpyxl

    schedule_path = os.path.join(os.path.dirname(__file__), "..", "ORTEGA_SCHEDULE.xls")

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["CAVITE STATE UNIVERSITY", "", "", "", "", "", "", "", "", ""])
    ws.append(["Department of Information Technology", "", "", "", "", "", "", "", "", ""])
    ws.append(["#", "Student ID", "Full Name", "Sex", "Course", "Year", "Section", "Email Address", "Contact", "Status"])
    ws.append(["1", "202110001", "DELA CRUZ, JUAN A.", "M", "BSCS", "4", "1", "juan@cvsu.edu.ph", "09123456789", "Regular"])
    ws.append(["2", "202110002", "SANTOS, MARIA B.", "F", "BSCS", "4", "1", "maria@cvsu.edu.ph", "09987654321", "Regular"])

    roster_path = tmp_path / "BSCS4-1 List of Students for 202612731-COSC 111A - C S ELECTIVE 3 (INTERNET OF THINGS).xlsx"
    wb.save(str(roster_path))

    # Pre-flight validation
    val = process_schedule.validate_rosters(schedule_path, [str(roster_path)])
    assert len(val) == 1
    assert val[0]["status"] == "warning"
    assert val[0]["issue"] == "extra_columns"
    assert val[0]["column_count"] == 10
    assert val[0]["student_count"] == 2

    # End-to-end generation
    out_dir = tmp_path / "output"
    res = process_schedule.process_all(
        schedule_path=schedule_path,
        xlsx_files=[str(roster_path)],
        output_dir_base=str(out_dir)
    )
    assert len(res["generated"]["ceit"]) > 0
    assert len(res["generated"]["attendance"]) > 0
    assert len(res["generated"]["grades"]) > 0
    assert len(res["errors"]["ceit"]) == 0
    assert len(res["errors"]["attendance"]) == 0
    assert len(res["errors"]["grades"]) == 0

