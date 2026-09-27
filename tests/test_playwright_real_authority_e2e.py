"""
Browser-Hosted ScriptAPI Bridge E2E Test Suite (Playwright Chromium).

IMPORTANT ARCHITECTURAL DISTINCTION:
  This suite tests ui.html hosted inside Playwright Chromium with a JavaScript
  `window.pywebview.api` Proxy bridged directly to the real Python `ScriptAPI`.
  It provides comprehensive end-to-end integration of the REAL Python backend
  (parsers, inspectors, role detectors, recipe validators, generators)
  with the REAL frontend UI.

  This is DISTINCT from the Native PyWebView Host Smoke/Integration suite
  (`tests/test_native_pywebview_host.py`), which launches the real Windows
  WebView2 desktop container and native DnD event loop.

Verifies:
  1. Application startup, release badge synchronization, dark/light theme toggle.
  2. Template set creation, physical template inspection, and candidate cards discovery.
  3. XLSX ambiguous template manual sheet confirmation dropdowns (discovered candidates only).
  4. Blank formula-free template manual sheet confirmation and validation.
  5. Invalid/corrupt template rejection and fail-closed UI behavior.

Failure Artifacts:
  If any test fails, complete failure artifacts are deterministically written to:
    test_reports/playwright/<test-name>/
      - screenshot.png
      - trace.zip
      - video.webm
      - console.log
      - page_errors.log
"""

import os
import sys
import json
import shutil
import pytest
import openpyxl
import docx
from pathlib import Path
from contextlib import contextmanager
from playwright.sync_api import sync_playwright, expect

from executable_test.main import ScriptAPI
from modules.services.template_set_manager import TemplateSetManager
from modules.models.template_set import (
    ROLE_GRADE_SHEET_LECTURE,
    ROLE_ATTENDANCE_LECTURE,
    ROLE_SYLLABUS,
    BUILTIN_SET_ID,
)

WORKSPACE_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
UI_HTML_PATH = os.path.join(WORKSPACE_DIR, "executable_test", "ui.html")
FILE_URL = f"file:///{UI_HTML_PATH.replace(os.sep, '/')}"


class ControlledWindowMock:
    """
    Controlled test seam providing simulated OS file selection to ScriptAPI._window
    for browser-hosted Playwright tests.

    Architectural boundary documentation:
      - Browser-hosted D2: Real ui.html DOM + real JavaScript event handlers +
        real window.pywebview.api bridge + real Python ScriptAPI + real template inspection/role detection.
      - Native OS file picker: Not exercisable in headless Chromium without a native desktop window.
        Simulated via ControlledWindowMock.create_file_dialog.
      - Native desktop picker: Covered in Layer D3 / native host tests.
    """

    def __init__(self, selected_files: list[str]):
        self.selected_files = selected_files

    def create_file_dialog(self, dialog_type=None, allow_multiple=False, file_types=()):
        return tuple(self.selected_files)


def inject_test_file_selection(api_instance: ScriptAPI, file_paths: list[str]) -> None:
    """
    Named, controlled test seam: attaches ControlledWindowMock to api_instance._window.
    Crucially, does NOT replace or monkeypatch ScriptAPI.browse_template_set_files,
    allowing the real production browse implementation to execute its validation,
    dialog invocation, and delegation to inspect_template_set_files().
    """
    api_instance._window = ControlledWindowMock(file_paths)


@contextmanager
def playwright_browser_session(playwright_ctx, api_instance, test_name: str = "playwright_test"):
    """
    Launches headless Chromium, bridges window.pywebview.api to the real Python ScriptAPI,
    and captures complete failure artifacts (screenshot, trace.zip, video.webm, console.log,
    page_errors.log) into test_reports/playwright/<test-name>/ strictly upon test failure.
    """
    artifact_dir = Path(WORKSPACE_DIR) / "test_reports" / "playwright" / test_name
    temp_video_dir = Path(WORKSPACE_DIR) / "test_reports" / "playwright" / ".temp_videos"
    temp_video_dir.mkdir(parents=True, exist_ok=True)

    browser = playwright_ctx.chromium.launch(headless=True)
    context = browser.new_context(
        record_video_dir=str(temp_video_dir),
        record_video_size={"width": 1280, "height": 720},
    )
    context.tracing.start(screenshots=True, snapshots=True, sources=True)
    page = context.new_page()

    console_logs = []
    page_errors = []
    page.on("console", lambda msg: console_logs.append(f"[{msg.type}] {msg.text}"))
    page.on("pageerror", lambda err: page_errors.append(str(err)))

    # Route all pywebview.api calls to the real Python ScriptAPI instance
    def handle_api_call(method_name, args):
        if not hasattr(api_instance, method_name):
            raise AttributeError(f"ScriptAPI has no method '{method_name}'")
        method = getattr(api_instance, method_name)
        return method(*args)

    page.expose_function("pywebview_call", handle_api_call)
    page.add_init_script("""
        window.pywebview = {
            api: new Proxy({}, {
                get(target, prop) {
                    return (...args) => window.pywebview_call(prop, args);
                }
            })
        };
    """)
    page.goto(FILE_URL)
    page.wait_for_load_state("domcontentloaded")

    artifact_errors = []
    test_failed = False
    try:
        yield page
    except Exception:
        test_failed = True
        artifact_dir.mkdir(parents=True, exist_ok=True)
        try:
            page.screenshot(path=str(artifact_dir / "screenshot.png"), full_page=True)
        except Exception as e:
            artifact_errors.append(f"screenshot: {e}")
        try:
            context.tracing.stop(path=str(artifact_dir / "trace.zip"))
        except Exception as e:
            artifact_errors.append(f"trace: {e}")
        try:
            with open(artifact_dir / "console.log", "w", encoding="utf-8") as f:
                f.write("\n".join(console_logs))
        except Exception as e:
            artifact_errors.append(f"console: {e}")
        try:
            with open(artifact_dir / "page_errors.log", "w", encoding="utf-8") as f:
                f.write("\n".join(page_errors))
        except Exception as e:
            artifact_errors.append(f"page_errors: {e}")
        raise
    finally:
        video = page.video
        video_path = None
        try:
            if video:
                video_path = video.path()
        except Exception:
            pass

        try:
            context.close()
        except Exception as e:
            if test_failed:
                artifact_errors.append(f"context_close: {e}")
        try:
            browser.close()
        except Exception as e:
            if test_failed:
                artifact_errors.append(f"browser_close: {e}")

        if test_failed:
            if video_path and os.path.exists(video_path):
                try:
                    shutil.move(video_path, str(artifact_dir / "video.webm"))
                except Exception as e:
                    artifact_errors.append(f"video_move: {e}")
            else:
                artifact_errors.append("video: path not created or missing on failure")

            if artifact_errors:
                try:
                    with open(artifact_dir / "artifact_errors.log", "w", encoding="utf-8") as f:
                        f.write("\n".join(artifact_errors))
                except Exception:
                    pass
        elif video_path and os.path.exists(video_path):
            try:
                os.remove(video_path)
            except Exception:
                pass


def _setup_browser_page(playwright_ctx, api_instance):
    """Backward-compatible helper for existing callers."""
    browser = playwright_ctx.chromium.launch(headless=True)
    page = browser.new_page()

    def handle_api_call(method_name, args):
        if not hasattr(api_instance, method_name):
            raise AttributeError(f"ScriptAPI has no method '{method_name}'")
        method = getattr(api_instance, method_name)
        return method(*args)

    page.expose_function("pywebview_call", handle_api_call)
    page.add_init_script("""
        window.pywebview = {
            api: new Proxy({}, {
                get(target, prop) {
                    return (...args) => window.pywebview_call(prop, args);
                }
            })
        };
    """)
    page.goto(FILE_URL)
    page.wait_for_load_state("domcontentloaded")
    return browser, page


# ── Test Cases ────────────────────────────────────────────────────────────────

def test_playwright_app_startup_and_ui_responsiveness():
    """Verify application startup, release badge, theme toggle, and navigation."""
    api = ScriptAPI()
    with sync_playwright() as p:
        with playwright_browser_session(p, api, "startup_and_ui_responsiveness") as page:
            # 1. Release badge visible
            badge = page.locator("header .badge-version")
            expect(badge).to_be_visible()
            assert "Release v" in badge.inner_text()

            # 2. Main workflow steps visible
            step1 = page.locator("#cardStep1")
            expect(step1).to_be_visible()

            # 3. Theme toggling
            initial_theme = page.evaluate("() => document.documentElement.getAttribute('data-theme') || 'dark'")
            assert initial_theme in ("dark", "light")
            page.click("#btnToggleTheme")
            page.wait_for_timeout(400)
            new_theme = page.evaluate("() => document.documentElement.getAttribute('data-theme')")
            assert new_theme != initial_theme
            page.click("#btnToggleTheme")
            page.wait_for_timeout(400)

            # 4. Settings modal opens and closes
            page.click("#btnOpenSettings")
            modal = page.locator("#modalParserSettingsBackdrop")
            expect(modal).to_be_visible()
            page.click("#btnCloseSettingsModal")
            expect(modal).to_be_hidden()


def test_playwright_template_set_lifecycle_and_discovery(tmp_path):
    """
    Verify complete template-set lifecycle in UI:
    Open settings -> Template Sets -> Create -> Inspect real files -> Save -> Verify in list.
    """
    mgr = TemplateSetManager.get_instance()
    original_sets_dir = mgr.template_sets_dir
    temp_sets_dir = str(tmp_path / "custom_sets")
    os.makedirs(temp_sets_dir, exist_ok=True)
    mgr.template_sets_dir = temp_sets_dir

    # Create synthetic test templates
    docx_path = str(tmp_path / "syllabus_template.docx")
    doc = docx.Document()
    doc.add_heading("COURSE SYLLABUS", level=1)
    t = doc.add_table(rows=4, cols=2)
    t.rows[0].cells[0].text = "Instructor:"
    t.rows[0].cells[1].text = "Prof. Dan Joseph Ortega"
    t.rows[1].cells[0].text = "Course & Section:"
    t.rows[1].cells[1].text = "BSCS 4-1"
    t.rows[2].cells[0].text = "Schedule Code:"
    t.rows[2].cells[1].text = "99112"
    t.rows[3].cells[0].text = "Subject:"
    t.rows[3].cells[1].text = "COSC 111"
    rt = doc.add_table(rows=2, cols=3)
    rt.rows[0].cells[0].text = "#"
    rt.rows[0].cells[1].text = "Student Name"
    rt.rows[0].cells[2].text = "Student Number"
    rt.rows[1].cells[0].text = "1"
    rt.rows[1].cells[1].text = "Sample Student"
    rt.rows[1].cells[2].text = "2026-0001"
    doc.save(docx_path)

    api = ScriptAPI()
    # Attach named controlled test seam: executes real production browse_template_set_files()
    inject_test_file_selection(api, [docx_path])

    with sync_playwright() as p:
        with playwright_browser_session(p, api, "template_set_lifecycle_and_discovery") as page:
            try:
                # 1. Open settings modal and navigate to Template Sets tab
                page.click("#btnOpenSettings")
                page.click("#cfgTabTemplateSets")
                sets_pane = page.locator("#cfgPaneTemplateSets")
                expect(sets_pane).to_be_visible()

                # 2. Click Create Template Set
                page.click("text=➕ Create Template Set")
                editor_modal = page.locator("#templateSetEditorModal")
                expect(editor_modal).to_be_visible()

                # 3. Fill in name and description
                page.fill("#templateSetNameInput", "Engineering Custom Set")
                page.fill("#templateSetDescInput", "Department test template set")

                # 4. Trigger inspection via REAL user UI click on #templateSetDropzone
                page.click("#templateSetDropzone")
                page.wait_for_timeout(400)

                # Verify inspection card rendered
                card_list = page.locator("#templateSetInspectedFilesList")
                expect(card_list).to_contain_text("syllabus_template.docx")
                expect(card_list).to_contain_text("Syllabus Acceptance")

                # 5. Save the template set
                page.click("text=✓ Save Template Set")
                page.wait_for_timeout(400)
                expect(editor_modal).to_be_hidden()

                # 6. Verify newly saved set appears in Available Template Sets list
                sets_list = page.locator("#templateSetsListContainer")
                expect(sets_list).to_contain_text("Engineering Custom Set")

                # 7. Verify persisted in backend TemplateSetManager
                user_sets = [s.display_name for s in mgr.list_template_sets()]
                assert "Engineering Custom Set" in user_sets

            finally:
                mgr.template_sets_dir = original_sets_dir


def test_playwright_ambiguous_xlsx_sheet_confirmation_flow(tmp_path):
    """
    Verify ambiguous XLSX template manual confirmation flow:
    - Template with 2 candidate rosters is discovered as ambiguous.
    - Discovered sheet confirmation dropdowns appear.
    - User selects discovered candidate sheet.
    - Template is validated and saved.
    """
    mgr = TemplateSetManager.get_instance()
    original_sets_dir = mgr.template_sets_dir
    temp_sets_dir = str(tmp_path / "ambig_sets")
    os.makedirs(temp_sets_dir, exist_ok=True)
    mgr.template_sets_dir = temp_sets_dir

    # Create an ambiguous XLSX template (2 identical candidate rosters without formula lineage)
    xlsx_path = str(tmp_path / "ambiguous_grade_sheet.xlsx")
    wb = openpyxl.Workbook()
    ws1 = wb.active
    ws1.title = "Lecture_Section_A"
    ws1.cell(1, 1, "Instructor: Prof. Dan Joseph Ortega")
    ws1.cell(2, 1, "Course & Section: BSCS-4A")
    ws1.cell(3, 1, "Schedule Code: 99112")
    ws1.cell(6, 1, "#")
    ws1.cell(6, 2, "Student Name")
    ws1.cell(6, 3, "Student Number")
    ws1.cell(6, 4, "Quiz 1")
    for r in range(7, 12):
        ws1.cell(r, 1, r - 6)
        ws1.cell(r, 2, f"Student {r - 6}")
        ws1.cell(r, 3, f"2026-CS-{r - 6:03d}")
        ws1.cell(r, 4, 85)

    ws2 = wb.create_sheet("Lecture_Section_B")
    ws2.cell(1, 1, "Instructor: Prof. Dan Joseph Ortega")
    ws2.cell(2, 1, "Course & Section: BSCS-4B")
    ws2.cell(3, 1, "Schedule Code: 99113")
    ws2.cell(6, 1, "#")
    ws2.cell(6, 2, "Student Name")
    ws2.cell(6, 3, "Student Number")
    ws2.cell(6, 4, "Quiz 1")
    for r in range(7, 12):
        ws2.cell(r, 1, r - 6)
        ws2.cell(r, 2, f"Student {r - 6}")
        ws2.cell(r, 3, f"2026-CS-{r - 6:03d}")
        ws2.cell(r, 4, 85)

    ws_sum = wb.create_sheet("Grading Sheet")
    ws_sum.cell(1, 1, "Cavite State University - Summary")
    ws_sum.cell(6, 1, "Student Number")
    ws_sum.cell(6, 2, "Final Rating")
    for r in range(7, 12):
        ws_sum.cell(r, 1, f"2026-CS-{r - 6:03d}")
        ws_sum.cell(r, 2, 85.0)

    wb.save(xlsx_path)

    api = ScriptAPI()
    inject_test_file_selection(api, [xlsx_path])
    with sync_playwright() as p:
        with playwright_browser_session(p, api, "ambiguous_xlsx_sheet_confirmation") as page:
            try:
                page.click("#btnOpenSettings")
                page.click("#cfgTabTemplateSets")
                page.click("text=➕ Create Template Set")

                page.fill("#templateSetNameInput", "Ambiguous Set Test")
                page.check("#templateSetFallbackCheck")

                # Trigger real UI click on dropzone with injected test file selection
                page.click("#templateSetDropzone")
                page.wait_for_timeout(400)

                # Verify card rendered with Discovered Sheet Confirmation
                card_list = page.locator("#templateSetInspectedFilesList")
                expect(card_list).to_contain_text("ambiguous_grade_sheet.xlsx")
                expect(card_list).to_contain_text("Discovered Sheet Confirmation")

                # Verify dropdown options contain discovered roster candidates
                roster_select = card_list.locator("select").filter(has_text="Lecture_Section_A")
                expect(roster_select).to_be_visible()

                # Select Lecture_Section_A explicitly
                roster_select.select_option("Lecture_Section_A")

                # Assign role to Grade Sheet (Lecture)
                role_select = card_list.locator("select").filter(has_text="-- Choose Canonical Role --")
                role_select.select_option(ROLE_GRADE_SHEET_LECTURE)

                # Save template set
                page.click("text=✓ Save Template Set")
                page.wait_for_timeout(500)

                # Backend verification: recipe was validated and assigned
                user_set = next((s for s in mgr.list_template_sets() if s.display_name == "Ambiguous Set Test"), None)
                assert user_set is not None
                assert ROLE_GRADE_SHEET_LECTURE in user_set.templates

            finally:
                mgr.template_sets_dir = original_sets_dir


def test_playwright_blank_formula_free_template_flow(tmp_path):
    """
    Verify blank formula-free template workflow through UI:
    - Zero formulas in template.
    - Discovered in UI, manually confirmed via discovered sheet dropdown.
    - Validates successfully without coordinate input.
    """
    mgr = TemplateSetManager.get_instance()
    original_sets_dir = mgr.template_sets_dir
    temp_sets_dir = str(tmp_path / "blank_sets")
    os.makedirs(temp_sets_dir, exist_ok=True)
    mgr.template_sets_dir = temp_sets_dir

    xlsx_path = str(tmp_path / "blank_custom_grading.xlsx")
    wb = openpyxl.Workbook()
    ws1 = wb.active
    ws1.title = "Lecture_Sheet"
    ws1.cell(1, 1, "Instructor: Prof. Dan Joseph Ortega")
    ws1.cell(2, 1, "Course & Section: BSCS-4A")
    ws1.cell(3, 1, "Schedule Code: 99112")
    ws1.cell(4, 1, "Subject: Software Architecture")
    ws1.cell(6, 1, "#")
    ws1.cell(6, 2, "Student Name")
    ws1.cell(6, 3, "Student Number")
    ws1.cell(6, 4, "Midterm Grade")
    for r in range(7, 12):
        ws1.cell(r, 1, r - 6)
        ws1.cell(r, 2, f"Student {r - 6}")
        ws1.cell(r, 3, f"2026-CS-{r - 6:03d}")
        ws1.cell(r, 4, 88.5)

    ws_sum = wb.create_sheet("Official_Summary")
    ws_sum.cell(1, 1, "Cavite State University - Summary")
    ws_sum.cell(6, 1, "Student Number")
    ws_sum.cell(6, 2, "Final Rating")
    for r in range(7, 12):
        ws_sum.cell(r, 1, f"2026-CS-{r - 6:03d}")
        ws_sum.cell(r, 2, 88.5)

    wb.save(xlsx_path)

    api = ScriptAPI()
    inject_test_file_selection(api, [xlsx_path])
    with sync_playwright() as p:
        with playwright_browser_session(p, api, "blank_formula_free_template") as page:
            try:
                page.click("#btnOpenSettings")
                page.click("#cfgTabTemplateSets")
                page.click("text=➕ Create Template Set")
                page.fill("#templateSetNameInput", "Blank Formula-Free Set")
                page.check("#templateSetFallbackCheck")

                # Trigger real UI click on dropzone with injected test file selection
                page.click("#templateSetDropzone")
                page.wait_for_timeout(400)

                card_list = page.locator("#templateSetInspectedFilesList")
                expect(card_list).to_contain_text("blank_custom_grading.xlsx")

                # Set role to Grade Sheet Lecture
                role_select = card_list.locator("select").filter(has_text="-- Choose Canonical Role --")
                role_select.select_option(ROLE_GRADE_SHEET_LECTURE)

                # Confirm discovered sheets
                roster_sel = card_list.locator("select").filter(has_text="Lecture_Sheet")
                roster_sel.select_option("Lecture_Sheet")

                # Save
                page.click("text=✓ Save Template Set")
                page.wait_for_timeout(500)

                # Backend verification
                user_set = next((s for s in mgr.list_template_sets() if s.display_name == "Blank Formula-Free Set"), None)
                assert user_set is not None
                assert ROLE_GRADE_SHEET_LECTURE in user_set.templates

            finally:
                mgr.template_sets_dir = original_sets_dir


def test_playwright_invalid_template_rejection(tmp_path):
    """Verify that corrupt or invalid files fail closed and are rejected in the UI."""
    corrupt_file = str(tmp_path / "corrupt_template.docx")
    with open(corrupt_file, "w", encoding="utf-8") as f:
        f.write("This is not a valid zip/docx archive.")

    api = ScriptAPI()
    with sync_playwright() as p:
        with playwright_browser_session(p, api, "invalid_template_rejection") as page:
            page.click("#btnOpenSettings")
            page.click("#cfgTabTemplateSets")
            page.click("text=➕ Create Template Set")

            page.evaluate(f"""async () => {{
                const res = await window.pywebview.api.inspect_template_set_files(['{corrupt_file.replace(os.sep, "/")}']);
                if (res && res.status === 'success') {{
                    addInspectedFiles(res.inspections || []);
                }}
            }}""")
            page.wait_for_timeout(300)

            card_list = page.locator("#templateSetInspectedFilesList")
            expect(card_list).to_contain_text("corrupt_template.docx")
            expect(card_list).to_contain_text("❌")


def test_playwright_failure_artifact_infrastructure_verification():
    """
    Infrastructure self-test: Intentionally triggers a failure inside a controlled
    playwright_browser_session and proves the explicit failure-artifact contract:
      1. screenshot.png: REQUIRED + NON-EMPTY (valid PNG signature \x89PNG).
      2. trace.zip: REQUIRED + NON-EMPTY (valid ZIP archive).
      3. video.webm: REQUIRED + NON-EMPTY (valid EBML WebM header \x1a\x45\xdf\xa3).
      4. console.log: REQUIRED + NON-EMPTY (captures browser console messages).
      5. page_errors.log: REQUIRED + NON-EMPTY (captures unhandled browser page exceptions).
    Cleans up the test artifact directory upon completion so test suite remains clean.
    """
    import zipfile

    test_session_name = "infra_failure_verification_test"
    artifact_dir = Path(WORKSPACE_DIR) / "test_reports" / "playwright" / test_session_name
    if artifact_dir.exists():
        shutil.rmtree(artifact_dir, ignore_errors=True)

    api = ScriptAPI()
    caught_failure = False

    with sync_playwright() as p:
        try:
            with playwright_browser_session(p, api, test_session_name) as page:
                page.evaluate("console.log('Artifact verification probe message')")
                page.evaluate("setTimeout(() => { window.__artifact_probe_function_that_does_not_exist__(); }, 0)")
                page.wait_for_timeout(150)
                assert False, "Controlled assertion failure for artifact pipeline verification"
        except AssertionError as e:
            caught_failure = True
            assert "Controlled assertion failure" in str(e)

    assert caught_failure is True, "Controlled failure must be raised and caught"
    assert artifact_dir.is_dir(), "Artifact directory must be created on test failure"

    screenshot_path = artifact_dir / "screenshot.png"
    trace_path = artifact_dir / "trace.zip"
    video_path = artifact_dir / "video.webm"
    console_path = artifact_dir / "console.log"
    page_errors_path = artifact_dir / "page_errors.log"

    # 1. screenshot.png: required + non-empty + valid PNG signature
    assert screenshot_path.is_file(), "screenshot.png must exist on test failure"
    assert screenshot_path.stat().st_size > 0, "screenshot.png must not be empty"
    with open(screenshot_path, "rb") as f:
        assert f.read(8).startswith(b"\x89PNG"), "screenshot.png must have valid PNG magic bytes"

    # 2. trace.zip: required + non-empty + valid ZIP archive
    assert trace_path.is_file(), "trace.zip must exist on test failure"
    assert trace_path.stat().st_size > 0, "trace.zip must not be empty"
    assert zipfile.is_zipfile(str(trace_path)), "trace.zip must be a valid zip archive"

    # 3. video.webm: required + non-empty + valid EBML header
    assert video_path.is_file(), "video.webm must exist and be finalized on test failure"
    assert video_path.stat().st_size > 0, "video.webm must not be empty"
    with open(video_path, "rb") as f:
        assert f.read(4) == b"\x1a\x45\xdf\xa3", "video.webm must have valid EBML header"

    # 4. console.log: required + non-empty + contains probe message
    assert console_path.is_file(), "console.log must exist on test failure"
    assert console_path.stat().st_size > 0, "console.log must not be empty"
    with open(console_path, "r", encoding="utf-8") as f:
        console_content = f.read()
    assert "Artifact verification probe message" in console_content

    # 5. page_errors.log: required + non-empty + contains probe exception
    assert page_errors_path.is_file(), "page_errors.log must exist on test failure"
    assert page_errors_path.stat().st_size > 0, "page_errors.log must not be empty"
    with open(page_errors_path, "r", encoding="utf-8") as f:
        page_errors_content = f.read()
    assert "__artifact_probe_function_that_does_not_exist__" in page_errors_content

    # Clean up test artifact directory after successful verification
    shutil.rmtree(artifact_dir, ignore_errors=True)


def test_playwright_browse_function_not_monkeypatched():
    """
    Regression assertion:
    Verifies that ScriptAPI.browse_template_set_files is the genuine production method
    from TemplateMixin and has NOT been replaced or monkeypatched by an anonymous lambda.
    """
    from executable_test.api.templates import TemplateMixin

    api = ScriptAPI()
    assert api.browse_template_set_files.__func__ is TemplateMixin.browse_template_set_files, (
        "ScriptAPI.browse_template_set_files must remain the authoritative production method "
        "and not be monkeypatched with a test lambda."
    )
