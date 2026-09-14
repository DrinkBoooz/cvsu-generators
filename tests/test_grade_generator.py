import os
import hashlib
import tempfile
import unittest
import openpyxl

from grade_generator import GradeGenerator

class GradeGeneratorTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.repo_dir = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
        cls.templates_dir = os.path.join(cls.repo_dir, "templates")
        cls.generator = GradeGenerator(cls.templates_dir)
        
        # Original hashes of root reference files
        cls.ll_root = os.path.join(cls.repo_dir, "Lecture and Lab.xlsx")
        cls.l_root = os.path.join(cls.repo_dir, "Lecture only.xlsx")
        with open(cls.ll_root, "rb") as f:
            cls.ll_root_hash = hashlib.sha256(f.read()).hexdigest()
        with open(cls.l_root, "rb") as f:
            cls.l_root_hash = hashlib.sha256(f.read()).hexdigest()

    def test_parse_subject(self):
        code, title = self.generator._parse_subject("DCIT 21A - INTRODUCTION TO COMPUTING")
        self.assertEqual(code, "DCIT 21A")
        self.assertEqual(title, "INTRODUCTION TO COMPUTING")

        # Unicode en-dash
        code_en, title_en = self.generator._parse_subject("ITEC 50 – WEB SYSTEMS AND TECHNOLOGY")
        self.assertEqual(code_en, "ITEC 50")
        self.assertEqual(title_en, "WEB SYSTEMS AND TECHNOLOGY")

        # Unicode em-dash
        code_em, title_em = self.generator._parse_subject("COSC 101 — ADVANCED DATABASE")
        self.assertEqual(code_em, "COSC 101")
        self.assertEqual(title_em, "ADVANCED DATABASE")

        code2, title2 = self.generator._parse_subject("COSC 111A-C S ELECTIVE 3 (IOT)")
        self.assertEqual(code2, "COSC 111A")
        self.assertEqual(title2, "C S ELECTIVE 3 (IOT)")

        code3, title3 = self.generator._parse_subject("CVSU 101")
        self.assertEqual(code3, "CVSU 101")
        self.assertEqual(title3, "")

    def test_parse_semester_and_year(self):
        sem, yr = self.generator._parse_semester_and_year("FIRST SEMESTER, AY 2026 - 2027")
        self.assertEqual(sem, "1st Semester")
        self.assertEqual(yr, "2026-2027")

        sem2, yr2 = self.generator._parse_semester_and_year("SECOND SEMESTER, SY 2025 - 2026")
        self.assertEqual(sem2, "2nd Semester")
        self.assertEqual(yr2, "2025-2026")

        sem3, yr3 = self.generator._parse_semester_and_year("1st Semester / 2026-2027")
        self.assertEqual(sem3, "1st Semester")
        self.assertEqual(yr3, "2026-2027")

    def test_normalize_course_section(self):
        self.assertEqual(self.generator._normalize_course_section("BSCS1-4"), "BSCS 1-4")
        self.assertEqual(self.generator._normalize_course_section("BSCS 1-4"), "BSCS 1-4")
        self.assertEqual(self.generator._normalize_course_section("BSIT1-1"), "BSIT 1-1")

    def test_generate_lecture_only(self):
        with tempfile.TemporaryDirectory() as tmpdir:
            out_path = os.path.join(tmpdir, "test_lecture.xlsx")
            info = {
                "instructor": "DAN JOSEPH A. ORTEGA",
                "course": "BSCS 1-4",
                "sched": "202612040",
                "subject": "DCIT 21A - INTRODUCTION TO COMPUTING",
                "semester": "FIRST SEMESTER, AY 2026 - 2027",
                "time": "Mon: 07:00AM-09:00AM / LEC: CL2",
                "has_lab": False
            }
            # Test with list of tuples
            students = [
                ("ARCA, BRENCH LORENZ L.", "261014253"),
                ("BERNAL, RUTHERFORD Q.", "261013992"),
            ]
            success = self.generator.generate(info, students, out_path)
            self.assertTrue(success)
            self.assertTrue(os.path.exists(out_path))

            wb = openpyxl.load_workbook(out_path, data_only=False)
            ws_lec = wb["Lecture"]
            self.assertEqual(ws_lec["C1"].value, 202612040)
            self.assertEqual(ws_lec["M1"].value, "BSCS 1-4")
            self.assertEqual(ws_lec["C2"].value, "DCIT 21A")
            self.assertEqual(ws_lec["M2"].value, "1st Semester")
            self.assertEqual(ws_lec["C3"].value, "INTRODUCTION TO COMPUTING")
            self.assertEqual(ws_lec["M3"].value, "2026-2027")
            self.assertEqual(ws_lec["M4"].value, "DAN JOSEPH A. ORTEGA")

            # Student 1 starts at row 11 in Lecture Only
            self.assertEqual(ws_lec.cell(11, 1).value, 1)
            self.assertEqual(ws_lec.cell(11, 2).value, "ARCA, BRENCH LORENZ L.")
            self.assertEqual(ws_lec.cell(11, 3).value, 261014253)

            # Student 2 at row 12
            self.assertEqual(ws_lec.cell(12, 1).value, 2)
            self.assertEqual(ws_lec.cell(12, 2).value, "BERNAL, RUTHERFORD Q.")

            # Row 13 should be blank for name and number
            self.assertIsNone(ws_lec.cell(13, 2).value)
            self.assertIsNone(ws_lec.cell(13, 3).value)

            # Grading Sheet checks
            ws_grd = wb["Grading Sheet"]
            self.assertEqual(ws_grd["A18"].value, "=Lecture!B11")
            self.assertEqual(ws_grd["B18"].value, "=Lecture!C11")
            self.assertEqual(ws_grd["C18"].value, "=Lecture!CA11")
            self.assertEqual(ws_grd["D60"].value, "=UPPER(Lecture!M4)")
            self.assertEqual(ws_grd["A77"].value, "PROF. CHARLOTTE B. CARANDANG")

    def test_generate_lecture_and_lab(self):
        with tempfile.TemporaryDirectory() as tmpdir:
            out_path = os.path.join(tmpdir, "test_lecture_lab.xlsx")
            info = {
                "instructor": "DAN JOSEPH A. ORTEGA",
                "course": "BSCS 4-2",
                "sched": "202612736",
                "subject": "COSC 111A - C S ELECTIVE 3 (INTERNET OF THINGS)",
                "semester": "FIRST SEMESTER, AY 2026 - 2027",
                "time": "Tue: 09:00AM-11:00AM / LAB: CCL 303; Wed: 07:00AM-10:30AM / LEC: HOURS",
                "has_lab": True
            }
            # Test with list of dicts
            students = [
                {"student_name": "ACASIO, ARRON M.", "student_number": "202300022"},
                {"student_name": "AMBROCIO, EMMAN S.", "student_number": "202300348"},
            ]
            success = self.generator.generate(info, students, out_path)
            self.assertTrue(success)
            self.assertTrue(os.path.exists(out_path))

            wb = openpyxl.load_workbook(out_path, data_only=False)
            ws_lec = wb["Lecture"]
            self.assertEqual(ws_lec["C1"].value, 202612736)
            self.assertEqual(ws_lec["M1"].value, "BSCS 4-2")
            self.assertEqual(ws_lec["C2"].value, "COSC 111A")
            self.assertEqual(ws_lec["M2"].value, "1st Semester")
            self.assertEqual(ws_lec["C3"].value, "C S ELECTIVE 3 (INTERNET OF THINGS)")
            self.assertEqual(ws_lec["M3"].value, "2026-2027")
            self.assertEqual(ws_lec["M4"].value, "DAN JOSEPH A. ORTEGA")
            self.assertEqual(ws_lec["BI57"].value, "DAN JOSEPH A. ORTEGA")

            # Student 1 starts at row 12 in Lecture and Lab
            self.assertEqual(ws_lec.cell(12, 1).value, 1)
            self.assertEqual(ws_lec.cell(12, 2).value, "ACASIO, ARRON M.")
            self.assertEqual(ws_lec.cell(12, 3).value, 202300022)

            # Laboratory Sheet
            ws_lab = wb["Laboratory"]
            self.assertEqual(ws_lab.cell(12, 1).value, 1)
            self.assertEqual(ws_lab.cell(12, 2).value, "=Lecture!B12")
            self.assertEqual(ws_lab.cell(12, 3).value, "=Lecture!C12")
            self.assertEqual(ws_lab["AO59"].value, "DAN JOSEPH A. ORTEGA")

            # Consolidated Sheet
            ws_con = wb["Consolidated"]
            self.assertEqual(ws_con.cell(10, 1).value, 1)
            self.assertEqual(ws_con.cell(10, 2).value, "=Lecture!B12")
            self.assertEqual(ws_con.cell(10, 3).value, "=Lecture!C12")
            self.assertEqual(ws_con["J56"].value, "DAN JOSEPH A. ORTEGA")

            # Grading Sheet
            ws_grd = wb["Grading Sheet"]
            self.assertEqual(ws_grd["A18"].value, "=Lecture!B12")
            self.assertEqual(ws_grd["B18"].value, "=Lecture!C12")
            self.assertEqual(ws_grd["C18"].value, "=Consolidated!V10")
            self.assertEqual(ws_grd["D60"].value, "=UPPER(Lecture!M4)")
            self.assertEqual(ws_grd["A77"].value, "PROF. CHARLOTTE B. CARANDANG")
            # Row 44 formula should be formula, not hardcoded 'DRP'
            self.assertEqual(ws_grd["C44"].value, "=Consolidated!V36")

    def test_detect_classes(self):
        import glob
        import process_schedule
        sched_path = os.path.join(self.repo_dir, "dev", "1st sem 25-26.xls")
        rosters = glob.glob(os.path.join(self.repo_dir, "schedules", "*.xlsx"))
        detected = process_schedule.detect_classes(sched_path, rosters)
        self.assertEqual(len(detected), len(rosters))
        
        # Check specific classes
        lookup = {d["course_sec"]: d for d in detected}
        
        # BSCS4-2 has both LAB and LEC in ORTEGA_SCHEDULE
        if "BSCS4-2" in lookup:
            self.assertEqual(lookup["BSCS4-2"]["detected_type"], "lecture_lab")
            self.assertTrue(lookup["BSCS4-2"]["has_lab"])
            
        # BSCS1-4 has only LEC
        if "BSCS1-4" in lookup:
            self.assertEqual(lookup["BSCS1-4"]["detected_type"], "lecture_only")
            self.assertFalse(lookup["BSCS1-4"]["has_lab"])

    def test_subject_type_override(self):
        with tempfile.TemporaryDirectory() as tmpdir:
            out_path = os.path.join(tmpdir, "test_override.xlsx")
            info = {
                "instructor": "DAN JOSEPH A. ORTEGA",
                "course": "BSCS 1-4",
                "sched": "202612040",
                "subject": "DCIT 21A",
                "semester": "1st Semester / 2026-2027",
                "time": "SEE SCHEDULE",
                "subject_type": "lecture_lab"  # User override from lecture_only to lecture_lab
            }
            students = [("TEST STUDENT", "123456")]
            success = self.generator.generate(info, students, out_path)
            self.assertTrue(success)
            wb = openpyxl.load_workbook(out_path)
            # When forced to lecture_lab, it must contain Laboratory sheet
            self.assertIn("Laboratory", wb.sheetnames)
            self.assertIn("Consolidated", wb.sheetnames)

    def test_clamp_max_students_overflow(self):
        with tempfile.TemporaryDirectory() as tmpdir:
            out_path = os.path.join(tmpdir, "test_overflow.xlsx")
            info = {
                "instructor": "DAN JOSEPH A. ORTEGA",
                "course": "BSCS 4-2",
                "sched": "202612736",
                "subject": "COSC 111A - C S ELECTIVE 3 (IOT)",
                "semester": "FIRST SEMESTER, AY 2026 - 2027",
                "time": "Tue: 09:00AM-11:00AM / LAB: CCL 303; Wed: 07:00AM-10:30AM / LEC: HOURS",
                "has_lab": True
            }
            # 45 students (exceeds max capacity of 40)
            students = [
                {"student_name": f"STUDENT {i:02d}", "student_number": f"2026000{i:02d}"}
                for i in range(1, 46)
            ]
            success = self.generator.generate(info, students, out_path)
            self.assertTrue(success)

            wb = openpyxl.load_workbook(out_path, data_only=False)
            ws_lec = wb["Lecture"]
            # Student 1 is at row 12
            self.assertEqual(ws_lec.cell(12, 2).value, "STUDENT 01")
            # Student 40 is at row 51
            self.assertEqual(ws_lec.cell(51, 2).value, "STUDENT 40")
            # Student 41 (row 52) must NOT be written - row 52 remains None
            self.assertIsNone(ws_lec.cell(52, 2).value)
            self.assertIsNone(ws_lec.cell(52, 3).value)

    def test_clamp_max_students_overflow_lecture_only(self):
        with tempfile.TemporaryDirectory() as tmpdir:
            out_path = os.path.join(tmpdir, "test_overflow_lec.xlsx")
            info = {
                "instructor": "DAN JOSEPH A. ORTEGA",
                "course": "BSCS 1-4",
                "sched": "202612040",
                "subject": "DCIT 21A - INTRODUCTION TO COMPUTING",
                "semester": "FIRST SEMESTER, AY 2026 - 2027",
                "time": "Mon: 07:00AM-09:00AM / LEC: CL2",
                "has_lab": False
            }
            # 65 students (exceeds lecture only max capacity of 60)
            students = [
                {"student_name": f"STUDENT {i:02d}", "student_number": f"2026000{i:02d}"}
                for i in range(1, 66)
            ]
            success = self.generator.generate(info, students, out_path)
            self.assertTrue(success)

            wb = openpyxl.load_workbook(out_path, data_only=False)
            ws_lec = wb["Lecture"]
            # Student 1 is at row 11 in Lecture Only
            self.assertEqual(ws_lec.cell(11, 2).value, "STUDENT 01")
            # Student 60 is at row 70
            self.assertEqual(ws_lec.cell(70, 2).value, "STUDENT 60")
            # Student 61 (row 71) must NOT be written - row 71 remains None
            self.assertIsNone(ws_lec.cell(71, 2).value)
            self.assertIsNone(ws_lec.cell(71, 3).value)

    def test_metadata_formula_injection_sanitization(self):
        with tempfile.TemporaryDirectory() as tmpdir:
            out_path = os.path.join(tmpdir, "test_injection.xlsx")
            info = {
                "instructor": "=CMD|' /C calc'!A0",
                "course": "+BSCS 1-4",
                "sched": "-202612040",
                "subject": "@DCIT 21A - INTRODUCTION TO COMPUTING",
                "semester": "=FIRST SEMESTER, AY 2026 - 2027",
                "time": "Mon: 07:00AM-09:00AM / LEC: CL2",
                "has_lab": False
            }
            students = [("=DDE('cmd';'calc')", "+261014253")]
            success = self.generator.generate(info, students, out_path)
            self.assertTrue(success)

            wb = openpyxl.load_workbook(out_path, data_only=False)
            ws_lec = wb["Lecture"]
            # Metadata fields must be prefixed with single quote
            self.assertTrue(str(ws_lec["M4"].value).startswith("'="))
            self.assertTrue(str(ws_lec["M1"].value).startswith("'+"))
            self.assertTrue(str(ws_lec["C1"].value).startswith("'-"))
            self.assertTrue(str(ws_lec["C2"].value).startswith("'@"))
            # Student fields must be prefixed with single quote
            self.assertTrue(str(ws_lec.cell(11, 2).value).startswith("'="))

    def test_root_files_remain_untouched(self):
        with open(self.ll_root, "rb") as f:
            current_ll_hash = hashlib.sha256(f.read()).hexdigest()
        with open(self.l_root, "rb") as f:
            current_l_hash = hashlib.sha256(f.read()).hexdigest()

        self.assertEqual(current_ll_hash, self.ll_root_hash, "Lecture and Lab.xlsx was modified!")
        self.assertEqual(current_l_hash, self.l_root_hash, "Lecture only.xlsx was modified!")


if __name__ == "__main__":
    unittest.main()
