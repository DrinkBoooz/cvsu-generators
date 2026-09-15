#!/usr/bin/env python3
"""
tests/test_synthetic_non_ceit_header.py

Tests generic DOCX generation on a synthetic non-CEIT template with a nonstandard metadata arrangement:
- Department:
- Program:
- Academic Year:
- Course Code:
- Course Title:
- Instructor:

Proves:
1. DocxTemplateInspector detects these non-CEIT fields.
2. RecipeValidator produces an authoritative recipe without CEIT requirements.
3. FieldResolver dynamically resolves authoritative and derived values.
4. ConfigurableDocumentGenerator populates all target cells driven purely by the recipe,
   with zero CEIT-specific or form-specific branching in the generator.
"""

import os
import docx
import pytest

from modules.models.schedule import ClassInfo
from modules.parsers.template_inspector import DocxTemplateInspector
from modules.parsers.recipe_validator import RecipeValidator
from modules.generators.document_generator import ConfigurableDocumentGenerator, DocumentGenerator


def build_synthetic_non_ceit_docx(file_path: str) -> None:
    """Creates a synthetic Word document with a 4-column non-CEIT header table and roster."""
    doc = docx.Document()
    doc.add_heading("DEPARTMENT ACADEMIC RECORD", level=1)

    # 4-column header table: [Label1, Value1, Label2, Value2]
    tbl = doc.add_table(rows=3, cols=4)
    # Row 0
    tbl.rows[0].cells[0].text = "Department:"
    tbl.rows[0].cells[1].text = ""
    tbl.rows[0].cells[2].text = "Program:"
    tbl.rows[0].cells[3].text = ""

    # Row 1
    tbl.rows[1].cells[0].text = "Academic Year:"
    tbl.rows[1].cells[1].text = ""
    tbl.rows[1].cells[2].text = "Course Code:"
    tbl.rows[1].cells[3].text = ""

    # Row 2
    tbl.rows[2].cells[0].text = "Course Title:"
    tbl.rows[2].cells[1].text = ""
    tbl.rows[2].cells[2].text = "Instructor:"
    tbl.rows[2].cells[3].text = ""

    doc.add_paragraph("")

    # Roster table
    r_tbl = doc.add_table(rows=2, cols=3)
    r_tbl.rows[0].cells[0].text = "No."
    r_tbl.rows[0].cells[1].text = "Student Number"
    r_tbl.rows[0].cells[2].text = "Student Name"

    r_tbl.rows[1].cells[0].text = "1"
    r_tbl.rows[1].cells[1].text = "2026-0001"
    r_tbl.rows[1].cells[2].text = "Sample Student"

    doc.save(file_path)


def test_synthetic_non_ceit_template_generation(tmp_path):
    tmpl_path = str(tmp_path / "synthetic_non_ceit_template.docx")
    out_path = str(tmp_path / "out_synthetic_non_ceit.docx")

    # 1. Build synthetic template
    build_synthetic_non_ceit_docx(tmpl_path)
    assert os.path.exists(tmpl_path)

    # 2. Inspect candidate observations
    inspector = DocxTemplateInspector()
    candidate = inspector.inspect(tmpl_path, profile_id="custom_docx")
    assert candidate is not None

    detected_fields = {c["field"] for c in candidate.header_candidates}
    assert "department" in detected_fields, f"department not detected in {detected_fields}"
    assert "program" in detected_fields, f"program not detected in {detected_fields}"
    assert "school_year" in detected_fields, f"school_year not detected in {detected_fields}"
    assert "subject_code" in detected_fields, f"subject_code not detected in {detected_fields}"
    assert "subject_title" in detected_fields, f"subject_title not detected in {detected_fields}"
    assert "instructor" in detected_fields, f"instructor not detected in {detected_fields}"

    # 3. Supply metadata for context fields (department, program)
    candidate.metadata["department"] = "Department of Biological Sciences"
    candidate.metadata["program"] = "BS Biology"

    # 4. Validate recipe
    recipe = RecipeValidator.validate(candidate, profile="custom_docx")
    assert recipe.verified_safe is True
    assert "department" in recipe.header_bindings
    assert "program" in recipe.header_bindings
    assert "school_year" in recipe.header_bindings
    assert "subject_code" in recipe.header_bindings
    assert "subject_title" in recipe.header_bindings
    assert "instructor" in recipe.header_bindings

    # 5. Runtime ClassInfo with realistic data
    info = ClassInfo(
        instructor="Dr. Jane Goodall",
        course_section="BSBIO 2-1",
        subject="BIO 101 - General Biology",
        semester_ay="1st Semester / 2026-2027",
        students=[
            ("Dela Cruz, Juan", "2026-1001"),
            ("Santos, Maria", "2026-1002"),
        ],
    )

    # 6. Instantiate extracted neutral engine
    gen = ConfigurableDocumentGenerator(tmpl_path, recipe)
    assert isinstance(gen, DocumentGenerator)

    # 7. Generate output
    gen.generate(info, out_path)
    assert os.path.exists(out_path)

    # 8. Readback and verify populated fields
    out_doc = docx.Document(out_path)
    h_tbl = out_doc.tables[0]

    # Row 0: Department & Program
    assert h_tbl.rows[0].cells[1].text.strip() == "Department of Biological Sciences"
    assert h_tbl.rows[0].cells[3].text.strip() == "BS Biology"

    # Row 1: Academic Year & Course Code
    assert h_tbl.rows[1].cells[1].text.strip() == "2026-2027"
    assert h_tbl.rows[1].cells[3].text.strip() == "BIO 101"

    # Row 2: Course Title & Instructor
    assert h_tbl.rows[2].cells[1].text.strip() == "General Biology"
    assert h_tbl.rows[2].cells[3].text.strip() == "Dr. Jane Goodall"

    # Roster table verification
    r_tbl = out_doc.tables[1]
    assert len(r_tbl.rows) == 1 + len(info.students)
    assert "Dela Cruz, Juan" in r_tbl.rows[1].cells[2].text
    assert "2026-1001" in r_tbl.rows[1].cells[1].text
    assert "Santos, Maria" in r_tbl.rows[2].cells[2].text
    assert "2026-1002" in r_tbl.rows[2].cells[1].text
