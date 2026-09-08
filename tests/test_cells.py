import openpyxl

if __name__ == "__main__":
    wb = openpyxl.load_workbook("c:/Users/danjo/OneDrive/CVSU GENERATORS/Lecture and Lab.xlsx", data_only=True)
    for sheet_name in wb.sheetnames:
        print(f"--- {sheet_name} ---")
        ws = wb[sheet_name]
        for r in range(40, 70):
            for c in range(1, 60):
                val = ws.cell(row=r, column=c).value
                if isinstance(val, str) and ("instructor" in val.lower() or "prepared by" in val.lower() or "faculty" in val.lower()):
                    print(f"Found {val} at {ws.cell(row=r, column=c).coordinate}")

