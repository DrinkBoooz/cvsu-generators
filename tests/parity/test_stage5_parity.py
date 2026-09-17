import os
import tempfile
import pytest
import docx
from modules.generators.ceit_gen import ClassInfo, GeneratorFactory
from tests.parity.parity_harness import ParityHarness

TEMPLATES_DIR = os.path.join(
    os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__)))),
    "templates",
)

SAMPLE_INFO = ClassInfo(
    instructor="PROF. JUAN DELA CRUZ",
    course_section="BSCS 3-1",
    schedule_code="202612345",
    subject="DCIT 50 - OBJECT-ORIENTED PROGRAMMING",
    time_days_room="Mon/Wed 07:00AM-09:00AM / CL2",
    semester_ay="1st Semester 2026-2027",
    students=[
        ("SANTOS, MARIA CLARA", "202610001"),
        ("REYES, PEDRO PENDUKO", "202610002"),
    ],
)


def test_all_seven_native_generators_instantiate_and_generate(tmp_path):
    """Verify all 7 native generators produce valid, populated DOCX files via GeneratorFactory."""
    factory = GeneratorFactory(TEMPLATES_DIR)
    gens = factory.get_all(include_custom=False)

    assert len(gens) == 7, f"Expected exactly 7 native generators, got {len(gens)}"

    expected_suffixes = [
        "SYLLABUS_ACCEPTANCE",
        "EXAM_RETURNS_MIDTERM",
        "EXAM_RETURNS_FINALS",
        "TOS_MIDTERM",
        "TOS_FINALS",
        "GRADE_DISCUSSION_MIDTERM",
        "GRADE_DISCUSSION_FINALS",
    ]

    actual_suffixes = [suffix for _, suffix in gens]
    assert actual_suffixes == expected_suffixes

    for gen_factory_fn, suffix in gens:
        gen = gen_factory_fn()
        out_file1 = str(tmp_path / f"test_{suffix}_run1.docx")
        out_file2 = str(tmp_path / f"test_{suffix}_run2.docx")

        # Generate run 1
        gen.generate(SAMPLE_INFO, out_file1)
        assert os.path.exists(out_file1), f"Failed to generate {out_file1}"

        # Generate run 2 (idempotency check)
        gen.generate(SAMPLE_INFO, out_file2)
        assert os.path.exists(out_file2), f"Failed to generate {out_file2}"

        # 1. Structural check
        doc1 = docx.Document(out_file1)
        assert len(doc1.tables) >= 1, f"{suffix} output has no tables"

        # 2. Check student roster population
        full_text = " ".join(p.text for p in doc1.paragraphs)
        for tbl in doc1.tables:
            for row in tbl.rows:
                full_text += " " + " ".join(c.text for c in row.cells)

        assert "SANTOS, MARIA CLARA" in full_text, f"{suffix} missing student 1"
        assert "202610001" in full_text, f"{suffix} missing student number 1"
        assert "REYES, PEDRO PENDUKO" in full_text, f"{suffix} missing student 2"
        assert "202610002" in full_text, f"{suffix} missing student number 2"

        # 3. Check metadata population
        assert "PROF. JUAN DELA CRUZ" in full_text, f"{suffix} missing instructor"
        assert "BSCS 3-1" in full_text, f"{suffix} missing course_section"
        assert "202612345" in full_text, f"{suffix} missing schedule_code"

        # 4. ParityHarness comparison between two runs of the same generator
        diff = ParityHarness.compare_docx(out_file1, out_file2)
        assert not diff.has_errors, f"Self-parity mismatch for {suffix}: {diff.summary()}"
