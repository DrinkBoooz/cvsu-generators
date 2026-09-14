import os
import pytest
import docx
import openpyxl
from tests.parity.parity_harness import ParityHarness

TEMPLATES_DIR = os.path.join(os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__)))), "templates")
LEC_XLSX = os.path.join(TEMPLATES_DIR, "GRADING_LECTURE_TEMPLATE.xlsx")
SYLLABUS_DOCX = os.path.join(TEMPLATES_DIR, "template_syllabus.docx")


def test_parity_harness_identical_docx():
    """Verify ParityHarness reports zero diffs for identical DOCX."""
    diff = ParityHarness.compare_docx(SYLLABUS_DOCX, SYLLABUS_DOCX)
    assert not diff.has_errors
    assert diff.summary() == ""


def test_parity_harness_identical_xlsx():
    """Verify ParityHarness reports zero diffs for identical XLSX."""
    diff = ParityHarness.compare_xlsx(LEC_XLSX, LEC_XLSX)
    assert not diff.has_errors
    assert diff.summary() == ""


def test_parity_harness_detects_docx_difference(tmp_path):
    """Verify ParityHarness detects differences in DOCX cell text."""
    doc1 = docx.Document(SYLLABUS_DOCX)
    # Modify a cell in copy
    doc2 = docx.Document(SYLLABUS_DOCX)
    doc2.tables[0].rows[0].cells[0].text = "MODIFIED TEXT FOR TEST"

    p1 = tmp_path / "doc1.docx"
    p2 = tmp_path / "doc2.docx"
    doc1.save(str(p1))
    doc2.save(str(p2))

    diff = ParityHarness.compare_docx(str(p1), str(p2))
    assert diff.has_errors
    assert "Cell text mismatch" in diff.summary()


def test_parity_harness_detects_xlsx_difference(tmp_path):
    """Verify ParityHarness detects differences in XLSX cell values."""
    wb1 = openpyxl.load_workbook(LEC_XLSX)
    wb2 = openpyxl.load_workbook(LEC_XLSX)
    wb2["Lecture"]["A1"] = "MODIFIED SCHEDULE CODE"

    p1 = tmp_path / "wb1.xlsx"
    p2 = tmp_path / "wb2.xlsx"
    wb1.save(str(p1))
    wb2.save(str(p2))

    diff = ParityHarness.compare_xlsx(str(p1), str(p2))
    assert diff.has_errors
    assert "Cell value mismatch" in diff.summary()
