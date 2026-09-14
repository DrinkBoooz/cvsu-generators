import os
import openpyxl
import pytest

TEMPLATES_DIR = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "templates")
LEC_TEMPLATE_PATH = os.path.join(TEMPLATES_DIR, "GRADING_LECTURE_TEMPLATE.xlsx")
LL_TEMPLATE_PATH = os.path.join(TEMPLATES_DIR, "GRADING_LECTURE_LAB_TEMPLATE.xlsx")


def test_xlsx_canonical_templates_exist():
    """Verify both physical canonical grading sheet templates exist."""
    assert os.path.exists(LEC_TEMPLATE_PATH), f"Missing {LEC_TEMPLATE_PATH}"
    assert os.path.exists(LL_TEMPLATE_PATH), f"Missing {LL_TEMPLATE_PATH}"


def test_lecture_only_xlsx_template_contract():
    """Verify structural contract of GRADING_LECTURE_TEMPLATE.xlsx."""
    wb = openpyxl.load_workbook(LEC_TEMPLATE_PATH, data_only=True)
    assert "Lecture" in wb.sheetnames
    ws = wb["Lecture"]

    # 1. Header Labels and Coordinates
    assert ws["A1"].value.strip() == "Schedule Code"
    assert ws["I1"].value.strip() == "Course & Section"
    assert ws["A2"].value.strip() == "Subject Code"
    assert ws["I2"].value.strip() == "Semester"
    assert ws["A3"].value.strip() == "Subject Title"
    assert ws["I3"].value.strip() == "School Year"
    assert ws["A4"].value.strip() == "Units"
    assert "Instructor" in ws["I4"].value

    # Destination target cells for data
    # C1 (Schedule Code), M1 (Course & Section), C2 (Subject Code), M2 (Semester),
    # C3 (Subject Title), M3 (School Year), C4 (Units), M4 (Instructor)
    expected_targets = ["C1", "M1", "C2", "M2", "C3", "M3", "C4", "M4"]
    for tgt in expected_targets:
        cell = ws[tgt]
        assert cell is not None

    # 2. Roster Columns and First Data Row
    # Row 7 is column headers
    assert ws.cell(7, 1).value.strip() == "#"
    assert "Student Name" in ws.cell(7, 2).value
    assert "Student Number" in ws.cell(7, 3).value

    # First data row is row 11 (0-indexed 10)
    assert ws.cell(11, 1).value == 1
    assert ws.cell(12, 1).value == 2

    # Capacity is at least 40
    numbered_rows = [r for r in range(11, 70) if isinstance(ws.cell(r, 1).value, int)]
    assert len(numbered_rows) >= 40
    assert ws.cell(50, 1).value == 40


def test_lecture_lab_xlsx_template_contract():
    """Verify structural contract of GRADING_LECTURE_LAB_TEMPLATE.xlsx."""
    wb = openpyxl.load_workbook(LL_TEMPLATE_PATH, data_only=True)
    assert "Lecture" in wb.sheetnames
    assert "Laboratory" in wb.sheetnames
    assert "Consolidated" in wb.sheetnames

    ws_lec = wb["Lecture"]

    # 1. Header Labels and Coordinates
    assert ws_lec["A1"].value.strip() == "Schedule Code"
    assert ws_lec["I1"].value.strip() == "Course & Section"
    assert ws_lec["A2"].value.strip() == "Subject Code"
    assert ws_lec["I2"].value.strip() == "Semester"
    assert ws_lec["A3"].value.strip() == "Subject Title"
    assert ws_lec["I3"].value.strip() == "School Year"
    assert ws_lec["A4"].value.strip() == "Units"
    assert "Instructor" in ws_lec["I4"].value

    # 2. Roster Columns and First Data Row
    # Row 8 is column headers
    assert ws_lec.cell(8, 1).value.strip() == "#"
    assert "Student Name" in ws_lec.cell(8, 2).value
    assert "Student Number" in ws_lec.cell(8, 3).value

    # First data row is row 12 (0-indexed 11)
    assert ws_lec.cell(12, 1).value == 1
    assert ws_lec.cell(13, 1).value == 2

    # Capacity is at least 40
    numbered_rows = [r for r in range(12, 70) if isinstance(ws_lec.cell(r, 1).value, int)]
    assert len(numbered_rows) >= 40
    assert ws_lec.cell(51, 1).value == 40

    # 3. Signature Blocks - Structural Merged Cell Stack
    # Lecture sheet signature:
    merged_lec = {str(m) for m in ws_lec.merged_cells.ranges}
    assert "BI60:BR62" in merged_lec, "Expected BI60:BR62 signature label merged range in Lecture"
    assert ws_lec["BI60"].value.strip() == "INSTRUCTOR"
    assert "BI57:BR59" in merged_lec, "Expected BI57:BR59 signature input merged range directly above label"

    # Laboratory sheet signature:
    ws_lab = wb["Laboratory"]
    merged_lab = {str(m) for m in ws_lab.merged_cells.ranges}
    assert "AO62:AX64" in merged_lab, "Expected AO62:AX64 signature label merged range in Laboratory"
    assert ws_lab["AO62"].value.strip() == "INSTRUCTOR"
    assert "AO59:AX61" in merged_lab, "Expected AO59:AX61 signature input merged range directly above label"
