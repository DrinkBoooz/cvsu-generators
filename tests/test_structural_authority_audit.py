"""
tests/test_structural_authority_audit.py

Comprehensive test suite verifying the complete structural authority closure:
1. Generic DOCX physical geometry bounds checks (headers, roster, signatures, corrupt files)
2. Serialized generic DOCX geometry bounds checks
3. Physical XLSX geometry bounds checks (worksheets, headers, roster, signatures, corrupt files)
4. Serialized XLSX geometry bounds checks
5. Structural ambiguity rejection across headers, signatures, and roster tables
6. Generator-side fail-closed execution safeguards
"""

import os
import pytest
import docx
import openpyxl

from modules.models.recipe import (
    RECIPE_SCHEMA_VERSION,
    TemplateError,
    AmbiguousTemplateError,
    InvalidRecipeError,
    RosterBinding,
    HeaderCellBinding,
    SignatureBinding,
    RawTemplateRecipeCandidate,
    ValidatedTemplateRecipe,
    PROFILE_ACADEMIC_DOCX,
    PROFILE_GRADE_SHEET_XLSX,
    _PRIVATE_CONSTRUCTION_SENTINEL,
)
from modules.parsers.recipe_validator import RecipeValidator
from modules.generators.document_generator import DocumentGenerator
from modules.generators.grade_gen import GradeGenerator


# ═══════════════════════════════════════════════════════════════════════════════
# 1. Generic DOCX Physical Geometry Validation
# ═══════════════════════════════════════════════════════════════════════════════

def _create_minimal_docx(path: str, num_tables: int = 1, rows: int = 3, cols: int = 3, paras: int = 2):
    doc = docx.Document()
    for _ in range(paras):
        doc.add_paragraph("Sample paragraph text")
    for _ in range(num_tables):
        doc.add_table(rows=rows, cols=cols)
    doc.save(path)


def test_docx_physical_header_table_index_out_of_range(tmp_path):
    """Physical header table index >= number of tables must raise InvalidRecipeError."""
    doc_path = str(tmp_path / "test_doc.docx")
    _create_minimal_docx(doc_path, num_tables=1, rows=3, cols=3)

    cand = RawTemplateRecipeCandidate(
        template_path=doc_path,
        profile_id="academic_docx",
        fingerprint="fp1",
        roster_candidate={"table_index": 0, "first_data_row_index": 1, "name_col": 0, "id_col": 1},
        header_candidates=[
            {"cell_type": "docx_table", "target": (5, 0, 0), "field": "instructor", "confidence": 1.0},
            {"cell_type": "docx_table", "target": (0, 0, 1), "field": "course_section", "confidence": 1.0},
            {"cell_type": "docx_table", "target": (0, 0, 2), "field": "schedule_code", "confidence": 1.0},
            {"cell_type": "docx_table", "target": (0, 1, 0), "field": "subject", "confidence": 1.0},
        ],
    )
    with pytest.raises(InvalidRecipeError) as exc:
        RecipeValidator.validate(cand, PROFILE_ACADEMIC_DOCX)
    assert "table index 5 out of range" in str(exc.value)


def test_docx_physical_header_row_out_of_range(tmp_path):
    """Physical header row index >= row count must raise InvalidRecipeError."""
    doc_path = str(tmp_path / "test_doc.docx")
    _create_minimal_docx(doc_path, num_tables=1, rows=3, cols=3)

    cand = RawTemplateRecipeCandidate(
        template_path=doc_path,
        profile_id="academic_docx",
        fingerprint="fp1",
        roster_candidate={"table_index": 0, "first_data_row_index": 1, "name_col": 0, "id_col": 1},
        header_candidates=[
            {"cell_type": "docx_table", "target": (0, 10, 0), "field": "instructor", "confidence": 1.0},
            {"cell_type": "docx_table", "target": (0, 0, 1), "field": "course_section", "confidence": 1.0},
            {"cell_type": "docx_table", "target": (0, 0, 2), "field": "schedule_code", "confidence": 1.0},
            {"cell_type": "docx_table", "target": (0, 1, 0), "field": "subject", "confidence": 1.0},
        ],
    )
    with pytest.raises(InvalidRecipeError) as exc:
        RecipeValidator.validate(cand, PROFILE_ACADEMIC_DOCX)
    assert "row index 10 out of range" in str(exc.value)


def test_docx_physical_header_cell_out_of_range(tmp_path):
    """Physical header cell index >= cell count must raise InvalidRecipeError."""
    doc_path = str(tmp_path / "test_doc.docx")
    _create_minimal_docx(doc_path, num_tables=1, rows=3, cols=3)

    cand = RawTemplateRecipeCandidate(
        template_path=doc_path,
        profile_id="academic_docx",
        fingerprint="fp1",
        roster_candidate={"table_index": 0, "first_data_row_index": 1, "name_col": 0, "id_col": 1},
        header_candidates=[
            {"cell_type": "docx_table", "target": (0, 0, 10), "field": "instructor", "confidence": 1.0},
            {"cell_type": "docx_table", "target": (0, 0, 1), "field": "course_section", "confidence": 1.0},
            {"cell_type": "docx_table", "target": (0, 0, 2), "field": "schedule_code", "confidence": 1.0},
            {"cell_type": "docx_table", "target": (0, 1, 0), "field": "subject", "confidence": 1.0},
        ],
    )
    with pytest.raises(InvalidRecipeError) as exc:
        RecipeValidator.validate(cand, PROFILE_ACADEMIC_DOCX)
    assert "cell index 10 out of range" in str(exc.value)


def test_docx_physical_paragraph_index_out_of_range(tmp_path):
    """Physical paragraph index >= paragraph count must raise InvalidRecipeError."""
    doc_path = str(tmp_path / "test_doc.docx")
    _create_minimal_docx(doc_path, num_tables=1, rows=3, cols=3, paras=2)

    cand = RawTemplateRecipeCandidate(
        template_path=doc_path,
        profile_id="academic_docx",
        fingerprint="fp1",
        roster_candidate={"table_index": 0, "first_data_row_index": 1, "name_col": 0, "id_col": 1},
        header_candidates=[
            {"cell_type": "docx_paragraph", "target": 15, "field": "instructor", "confidence": 1.0},
            {"cell_type": "docx_table", "target": (0, 0, 1), "field": "course_section", "confidence": 1.0},
            {"cell_type": "docx_table", "target": (0, 0, 2), "field": "schedule_code", "confidence": 1.0},
            {"cell_type": "docx_table", "target": (0, 1, 0), "field": "subject", "confidence": 1.0},
        ],
    )
    with pytest.raises(InvalidRecipeError) as exc:
        RecipeValidator.validate(cand, PROFILE_ACADEMIC_DOCX)
    assert "paragraph index 15 out of range" in str(exc.value)


def test_docx_physical_roster_table_index_out_of_range(tmp_path):
    """Physical roster table index out of range must raise InvalidRecipeError."""
    doc_path = str(tmp_path / "test_doc.docx")
    _create_minimal_docx(doc_path, num_tables=1, rows=3, cols=3)

    cand = RawTemplateRecipeCandidate(
        template_path=doc_path,
        profile_id="academic_docx",
        fingerprint="fp1",
        roster_candidate={"table_index": 5, "first_data_row_index": 1, "name_col": 0, "id_col": 1},
        header_candidates=[
            {"cell_type": "docx_table", "target": (0, 0, 0), "field": "instructor", "confidence": 1.0},
            {"cell_type": "docx_table", "target": (0, 0, 1), "field": "course_section", "confidence": 1.0},
            {"cell_type": "docx_table", "target": (0, 0, 2), "field": "schedule_code", "confidence": 1.0},
            {"cell_type": "docx_table", "target": (0, 1, 0), "field": "subject", "confidence": 1.0},
        ],
    )
    with pytest.raises(InvalidRecipeError) as exc:
        RecipeValidator.validate(cand, PROFILE_ACADEMIC_DOCX)
    assert "Roster binding table index 5 out of range" in str(exc.value)


def test_docx_physical_roster_row_out_of_range(tmp_path):
    """Physical roster first_data_row_index out of range must raise InvalidRecipeError."""
    doc_path = str(tmp_path / "test_doc.docx")
    _create_minimal_docx(doc_path, num_tables=1, rows=3, cols=3)

    cand = RawTemplateRecipeCandidate(
        template_path=doc_path,
        profile_id="academic_docx",
        fingerprint="fp1",
        roster_candidate={"table_index": 0, "first_data_row_index": 10, "name_col": 0, "id_col": 1},
        header_candidates=[
            {"cell_type": "docx_table", "target": (0, 0, 0), "field": "instructor", "confidence": 1.0},
            {"cell_type": "docx_table", "target": (0, 0, 1), "field": "course_section", "confidence": 1.0},
            {"cell_type": "docx_table", "target": (0, 0, 2), "field": "schedule_code", "confidence": 1.0},
            {"cell_type": "docx_table", "target": (0, 1, 0), "field": "subject", "confidence": 1.0},
        ],
    )
    with pytest.raises(InvalidRecipeError) as exc:
        RecipeValidator.validate(cand, PROFILE_ACADEMIC_DOCX)
    assert "Roster first_data_row_index 10 out of range" in str(exc.value)


def test_docx_physical_roster_name_col_out_of_range(tmp_path):
    """Physical roster name_col out of range must raise InvalidRecipeError."""
    doc_path = str(tmp_path / "test_doc.docx")
    _create_minimal_docx(doc_path, num_tables=1, rows=3, cols=3)

    cand = RawTemplateRecipeCandidate(
        template_path=doc_path,
        profile_id="academic_docx",
        fingerprint="fp1",
        roster_candidate={"table_index": 0, "first_data_row_index": 1, "name_col": 10, "id_col": 1},
        header_candidates=[
            {"cell_type": "docx_table", "target": (0, 0, 0), "field": "instructor", "confidence": 1.0},
            {"cell_type": "docx_table", "target": (0, 0, 1), "field": "course_section", "confidence": 1.0},
            {"cell_type": "docx_table", "target": (0, 0, 2), "field": "schedule_code", "confidence": 1.0},
            {"cell_type": "docx_table", "target": (0, 1, 0), "field": "subject", "confidence": 1.0},
        ],
    )
    with pytest.raises(InvalidRecipeError) as exc:
        RecipeValidator.validate(cand, PROFILE_ACADEMIC_DOCX)
    assert "Roster name_col (10) out of range" in str(exc.value)


def test_docx_physical_roster_id_col_out_of_range(tmp_path):
    """Physical roster id_col out of range must raise InvalidRecipeError."""
    doc_path = str(tmp_path / "test_doc.docx")
    _create_minimal_docx(doc_path, num_tables=1, rows=3, cols=3)

    cand = RawTemplateRecipeCandidate(
        template_path=doc_path,
        profile_id="academic_docx",
        fingerprint="fp1",
        roster_candidate={"table_index": 0, "first_data_row_index": 1, "name_col": 0, "id_col": 10},
        header_candidates=[
            {"cell_type": "docx_table", "target": (0, 0, 0), "field": "instructor", "confidence": 1.0},
            {"cell_type": "docx_table", "target": (0, 0, 1), "field": "course_section", "confidence": 1.0},
            {"cell_type": "docx_table", "target": (0, 0, 2), "field": "schedule_code", "confidence": 1.0},
            {"cell_type": "docx_table", "target": (0, 1, 0), "field": "subject", "confidence": 1.0},
        ],
    )
    with pytest.raises(InvalidRecipeError) as exc:
        RecipeValidator.validate(cand, PROFILE_ACADEMIC_DOCX)
    assert "Roster id_col (10) out of range" in str(exc.value)


def test_docx_physical_signature_table_coordinate_out_of_range(tmp_path):
    """Physical signature table coordinate out of range must raise InvalidRecipeError."""
    doc_path = str(tmp_path / "test_doc.docx")
    _create_minimal_docx(doc_path, num_tables=1, rows=3, cols=3)

    cand = RawTemplateRecipeCandidate(
        template_path=doc_path,
        profile_id="academic_docx",
        fingerprint="fp1",
        roster_candidate={"table_index": 0, "first_data_row_index": 1, "name_col": 0, "id_col": 1},
        header_candidates=[
            {"cell_type": "docx_table", "target": (0, 0, 0), "field": "instructor", "confidence": 1.0},
            {"cell_type": "docx_table", "target": (0, 0, 1), "field": "course_section", "confidence": 1.0},
            {"cell_type": "docx_table", "target": (0, 0, 2), "field": "schedule_code", "confidence": 1.0},
            {"cell_type": "docx_table", "target": (0, 1, 0), "field": "subject", "confidence": 1.0},
        ],
        signature_candidates=[
            {"role": "instructor_signature", "target": (5, 0, 0), "confidence": 1.0},
        ],
    )
    with pytest.raises(InvalidRecipeError) as exc:
        RecipeValidator.validate(cand, PROFILE_ACADEMIC_DOCX)
    assert "Signature binding for 'instructor_signature' table index 5 out of range" in str(exc.value)


def test_docx_corrupt_file_fails_closed(tmp_path):
    """Corrupt physical DOCX template must fail closed with TemplateError."""
    doc_path = str(tmp_path / "corrupt.docx")
    with open(doc_path, "wb") as f:
        f.write(b"not a valid zip file or docx")

    cand = RawTemplateRecipeCandidate(
        template_path=doc_path,
        profile_id="academic_docx",
        fingerprint="fp_corrupt",
        roster_candidate={"table_index": 0, "first_data_row_index": 1, "name_col": 0, "id_col": 1},
        header_candidates=[
            {"cell_type": "docx_table", "target": (0, 0, 0), "field": "instructor", "confidence": 1.0},
            {"cell_type": "docx_table", "target": (0, 0, 1), "field": "course_section", "confidence": 1.0},
            {"cell_type": "docx_table", "target": (0, 0, 2), "field": "schedule_code", "confidence": 1.0},
            {"cell_type": "docx_table", "target": (0, 1, 0), "field": "subject", "confidence": 1.0},
        ],
    )
    with pytest.raises(TemplateError) as exc:
        RecipeValidator.validate(cand, PROFILE_ACADEMIC_DOCX)
    assert "Physical template inspection failed" in str(exc.value)


# ═══════════════════════════════════════════════════════════════════════════════
# 2. Serialized Generic DOCX Geometry Validation
# ═══════════════════════════════════════════════════════════════════════════════

def test_docx_serialized_missing_geometry_fails():
    """Serialized recipe with coordinates and missing physical file requires docx_geometry."""
    cand = RawTemplateRecipeCandidate(
        template_path="nonexistent.docx",
        profile_id="academic_docx",
        fingerprint="fp1",
        roster_candidate={"table_index": 0, "first_data_row_index": 1, "name_col": 0, "id_col": 1},
        header_candidates=[
            {"cell_type": "docx_table", "target": (0, 0, 0), "field": "instructor", "confidence": 1.0},
            {"cell_type": "docx_table", "target": (0, 0, 1), "field": "course_section", "confidence": 1.0},
            {"cell_type": "docx_table", "target": (0, 0, 2), "field": "schedule_code", "confidence": 1.0},
            {"cell_type": "docx_table", "target": (0, 1, 0), "field": "subject", "confidence": 1.0},
        ],
        metadata={},
    )
    with pytest.raises(InvalidRecipeError) as exc:
        RecipeValidator.validate(cand, PROFILE_ACADEMIC_DOCX)
    assert "missing required 'docx_geometry' metadata" in str(exc.value)


def test_docx_serialized_table_coordinate_out_of_bounds():
    """Serialized recipe with out-of-bounds table coordinate against docx_geometry fails."""
    cand = RawTemplateRecipeCandidate(
        template_path="nonexistent.docx",
        profile_id="academic_docx",
        fingerprint="fp1",
        roster_candidate={"table_index": 0, "first_data_row_index": 1, "name_col": 0, "id_col": 1},
        header_candidates=[
            {"cell_type": "docx_table", "target": (2, 0, 0), "field": "instructor", "confidence": 1.0},
            {"cell_type": "docx_table", "target": (0, 0, 1), "field": "course_section", "confidence": 1.0},
            {"cell_type": "docx_table", "target": (0, 0, 2), "field": "schedule_code", "confidence": 1.0},
            {"cell_type": "docx_table", "target": (0, 1, 0), "field": "subject", "confidence": 1.0},
        ],
        metadata={
            "docx_geometry": {
                "table_row_cell_counts": [[3, 3]],  # 1 table with 2 rows, 3 cells each
                "paragraph_count": 2,
            }
        },
    )
    with pytest.raises(InvalidRecipeError) as exc:
        RecipeValidator.validate(cand, PROFILE_ACADEMIC_DOCX)
    assert "table index 2 out of range" in str(exc.value)


def test_docx_serialized_row_coordinate_out_of_bounds():
    """Serialized recipe with out-of-bounds row coordinate against docx_geometry fails."""
    cand = RawTemplateRecipeCandidate(
        template_path="nonexistent.docx",
        profile_id="academic_docx",
        fingerprint="fp1",
        roster_candidate={"table_index": 0, "first_data_row_index": 1, "name_col": 0, "id_col": 1},
        header_candidates=[
            {"cell_type": "docx_table", "target": (0, 5, 0), "field": "instructor", "confidence": 1.0},
            {"cell_type": "docx_table", "target": (0, 0, 1), "field": "course_section", "confidence": 1.0},
            {"cell_type": "docx_table", "target": (0, 0, 2), "field": "schedule_code", "confidence": 1.0},
            {"cell_type": "docx_table", "target": (0, 1, 0), "field": "subject", "confidence": 1.0},
        ],
        metadata={
            "docx_geometry": {
                "table_row_cell_counts": [[3, 3]],
                "paragraph_count": 2,
            }
        },
    )
    with pytest.raises(InvalidRecipeError) as exc:
        RecipeValidator.validate(cand, PROFILE_ACADEMIC_DOCX)
    assert "row index 5 out of range" in str(exc.value)


def test_docx_serialized_cell_coordinate_out_of_bounds():
    """Serialized recipe with out-of-bounds cell coordinate against docx_geometry fails."""
    cand = RawTemplateRecipeCandidate(
        template_path="nonexistent.docx",
        profile_id="academic_docx",
        fingerprint="fp1",
        roster_candidate={"table_index": 0, "first_data_row_index": 1, "name_col": 0, "id_col": 1},
        header_candidates=[
            {"cell_type": "docx_table", "target": (0, 0, 8), "field": "instructor", "confidence": 1.0},
            {"cell_type": "docx_table", "target": (0, 0, 1), "field": "course_section", "confidence": 1.0},
            {"cell_type": "docx_table", "target": (0, 0, 2), "field": "schedule_code", "confidence": 1.0},
            {"cell_type": "docx_table", "target": (0, 1, 0), "field": "subject", "confidence": 1.0},
        ],
        metadata={
            "docx_geometry": {
                "table_row_cell_counts": [[3, 3]],
                "paragraph_count": 2,
            }
        },
    )
    with pytest.raises(InvalidRecipeError) as exc:
        RecipeValidator.validate(cand, PROFILE_ACADEMIC_DOCX)
    assert "cell index 8 out of range" in str(exc.value)


# ═══════════════════════════════════════════════════════════════════════════════
# 3. XLSX Physical and Serialized Geometry Validation
# ═══════════════════════════════════════════════════════════════════════════════

def _create_minimal_xlsx(path: str, sheets=("Lecture", "Grading Sheet"), max_r=20, max_c=10):
    wb = openpyxl.Workbook()
    default_ws = wb.active
    default_ws.title = sheets[0]
    for r in range(1, max_r + 1):
        for c in range(1, max_c + 1):
            default_ws.cell(r, c).value = f"cell_{r}_{c}"

    for s in sheets[1:]:
        ws = wb.create_sheet(title=s)
        for r in range(1, max_r + 1):
            for c in range(1, max_c + 1):
                ws.cell(r, c).value = f"cell_{r}_{c}"
    wb.save(path)


def test_xlsx_physical_worksheet_target_invalid(tmp_path):
    """Physical XLSX with invalid worksheet target raises InvalidRecipeError."""
    xlsx_path = str(tmp_path / "test.xlsx")
    _create_minimal_xlsx(xlsx_path, sheets=("Lecture",))

    cand = RawTemplateRecipeCandidate(
        template_path=xlsx_path,
        profile_id="grade_sheet_xlsx",
        fingerprint="fp1",
        roster_candidate={
            "table_index": 0,
            "worksheet_name": "NonExistentSheet",
            "first_data_row_index": 7,
            "name_col": 3,
            "id_col": 1,
            "index_col": 1,
            "capacity_limit": 10,
        },
        header_candidates=[
            {"cell_type": "xlsx_cell", "target": "Lecture!B3", "field": "instructor", "confidence": 1.0},
            {"cell_type": "xlsx_cell", "target": "Lecture!B1", "field": "course_section", "confidence": 1.0},
            {"cell_type": "xlsx_cell", "target": "Lecture!C1", "field": "schedule_code", "confidence": 1.0},
        ],
    )
    with pytest.raises(InvalidRecipeError) as exc:
        RecipeValidator.validate(cand, PROFILE_GRADE_SHEET_XLSX)
    assert "Worksheet 'NonExistentSheet' not found" in str(exc.value)


def test_xlsx_physical_header_cell_out_of_bounds(tmp_path):
    """Physical XLSX with header cell coordinate out of worksheet bounds raises InvalidRecipeError."""
    xlsx_path = str(tmp_path / "test.xlsx")
    _create_minimal_xlsx(xlsx_path, sheets=("Lecture",), max_r=20, max_c=10)

    cand = RawTemplateRecipeCandidate(
        template_path=xlsx_path,
        profile_id="grade_sheet_xlsx",
        fingerprint="fp1",
        roster_candidate={
            "table_index": 0,
            "worksheet_name": "Lecture",
            "first_data_row_index": 7,
            "name_col": 3,
            "id_col": 1,
            "index_col": 1,
            "capacity_limit": 10,
        },
        header_candidates=[
            {"cell_type": "xlsx_cell", "target": "Lecture!Z100", "field": "instructor", "confidence": 1.0},
            {"cell_type": "xlsx_cell", "target": "Lecture!B1", "field": "course_section", "confidence": 1.0},
            {"cell_type": "xlsx_cell", "target": "Lecture!C1", "field": "schedule_code", "confidence": 1.0},
        ],
    )
    with pytest.raises(InvalidRecipeError) as exc:
        RecipeValidator.validate(cand, PROFILE_GRADE_SHEET_XLSX)
    assert "out of bounds for sheet 'Lecture'" in str(exc.value)


def test_xlsx_physical_roster_row_out_of_bounds(tmp_path):
    """Physical XLSX with roster row out of worksheet bounds raises InvalidRecipeError."""
    xlsx_path = str(tmp_path / "test.xlsx")
    _create_minimal_xlsx(xlsx_path, sheets=("Lecture",), max_r=20, max_c=10)

    cand = RawTemplateRecipeCandidate(
        template_path=xlsx_path,
        profile_id="grade_sheet_xlsx",
        fingerprint="fp1",
        roster_candidate={
            "table_index": 0,
            "worksheet_name": "Lecture",
            "first_data_row_index": 50,  # max_r is 20
            "name_col": 3,
            "id_col": 1,
            "index_col": 1,
            "capacity_limit": 10,
        },
        header_candidates=[
            {"cell_type": "xlsx_cell", "target": "Lecture!B3", "field": "instructor", "confidence": 1.0},
            {"cell_type": "xlsx_cell", "target": "Lecture!B1", "field": "course_section", "confidence": 1.0},
            {"cell_type": "xlsx_cell", "target": "Lecture!C1", "field": "schedule_code", "confidence": 1.0},
        ],
    )
    with pytest.raises(InvalidRecipeError) as exc:
        RecipeValidator.validate(cand, PROFILE_GRADE_SHEET_XLSX)
    assert "Roster first_data_row_index 50 out of bounds" in str(exc.value)


def test_xlsx_physical_roster_col_out_of_bounds(tmp_path):
    """Physical XLSX with roster column out of worksheet bounds raises InvalidRecipeError."""
    xlsx_path = str(tmp_path / "test.xlsx")
    _create_minimal_xlsx(xlsx_path, sheets=("Lecture",), max_r=20, max_c=10)

    cand = RawTemplateRecipeCandidate(
        template_path=xlsx_path,
        profile_id="grade_sheet_xlsx",
        fingerprint="fp1",
        roster_candidate={
            "table_index": 0,
            "worksheet_name": "Lecture",
            "first_data_row_index": 7,
            "name_col": 50,  # max_c is 10
            "id_col": 1,
            "index_col": 1,
            "capacity_limit": 10,
        },
        header_candidates=[
            {"cell_type": "xlsx_cell", "target": "Lecture!B3", "field": "instructor", "confidence": 1.0},
            {"cell_type": "xlsx_cell", "target": "Lecture!B1", "field": "course_section", "confidence": 1.0},
            {"cell_type": "xlsx_cell", "target": "Lecture!C1", "field": "schedule_code", "confidence": 1.0},
        ],
    )
    with pytest.raises(InvalidRecipeError) as exc:
        RecipeValidator.validate(cand, PROFILE_GRADE_SHEET_XLSX)
    assert "Roster name_col 50 out of bounds" in str(exc.value)


def test_xlsx_physical_signature_target_out_of_bounds(tmp_path):
    """Physical XLSX with signature coordinate out of bounds raises InvalidRecipeError."""
    xlsx_path = str(tmp_path / "test.xlsx")
    _create_minimal_xlsx(xlsx_path, sheets=("Lecture",), max_r=20, max_c=10)

    cand = RawTemplateRecipeCandidate(
        template_path=xlsx_path,
        profile_id="grade_sheet_xlsx",
        fingerprint="fp1",
        roster_candidate={
            "table_index": 0,
            "worksheet_name": "Lecture",
            "first_data_row_index": 7,
            "name_col": 3,
            "id_col": 1,
            "index_col": 1,
            "capacity_limit": 10,
        },
        header_candidates=[
            {"cell_type": "xlsx_cell", "target": "Lecture!B3", "field": "instructor", "confidence": 1.0},
            {"cell_type": "xlsx_cell", "target": "Lecture!B1", "field": "course_section", "confidence": 1.0},
            {"cell_type": "xlsx_cell", "target": "Lecture!C1", "field": "schedule_code", "confidence": 1.0},
        ],
        signature_candidates=[
            {"role": "instructor", "target": "Lecture!Z100", "confidence": 1.0},
        ],
    )
    with pytest.raises(InvalidRecipeError) as exc:
        RecipeValidator.validate(cand, PROFILE_GRADE_SHEET_XLSX)
    assert "Signature binding for 'instructor' coordinate 'Lecture!Z100' out of bounds" in str(exc.value)


def test_xlsx_corrupt_file_fails_closed(tmp_path):
    """Corrupt physical XLSX template must fail closed with TemplateError."""
    xlsx_path = str(tmp_path / "corrupt.xlsx")
    with open(xlsx_path, "wb") as f:
        f.write(b"corrupt xlsx bytes")

    cand = RawTemplateRecipeCandidate(
        template_path=xlsx_path,
        profile_id="grade_sheet_xlsx",
        fingerprint="fp_corrupt",
        roster_candidate={
            "table_index": 0,
            "worksheet_name": "Lecture",
            "first_data_row_index": 7,
            "name_col": 3,
            "id_col": 1,
            "index_col": 1,
            "capacity_limit": 10,
        },
        header_candidates=[
            {"cell_type": "xlsx_cell", "target": "Lecture!B3", "field": "instructor", "confidence": 1.0},
            {"cell_type": "xlsx_cell", "target": "Lecture!B1", "field": "course_section", "confidence": 1.0},
            {"cell_type": "xlsx_cell", "target": "Lecture!C1", "field": "schedule_code", "confidence": 1.0},
        ],
    )
    with pytest.raises(TemplateError) as exc:
        RecipeValidator.validate(cand, PROFILE_GRADE_SHEET_XLSX)
    assert "Physical template inspection failed" in str(exc.value)


def test_xlsx_serialized_missing_geometry_fails():
    """Serialized XLSX recipe without physical file and missing xlsx_geometry raises InvalidRecipeError."""
    cand = RawTemplateRecipeCandidate(
        template_path="nonexistent.xlsx",
        profile_id="grade_sheet_xlsx",
        fingerprint="fp1",
        roster_candidate={
            "table_index": 0,
            "worksheet_name": "Lecture",
            "first_data_row_index": 7,
            "name_col": 3,
            "id_col": 1,
            "index_col": 1,
            "capacity_limit": 10,
        },
        header_candidates=[
            {"cell_type": "xlsx_cell", "target": "Lecture!B3", "field": "instructor", "confidence": 1.0},
            {"cell_type": "xlsx_cell", "target": "Lecture!B1", "field": "course_section", "confidence": 1.0},
            {"cell_type": "xlsx_cell", "target": "Lecture!C1", "field": "schedule_code", "confidence": 1.0},
        ],
        metadata={},
    )
    with pytest.raises(InvalidRecipeError) as exc:
        RecipeValidator.validate(cand, PROFILE_GRADE_SHEET_XLSX)
    assert "missing required 'xlsx_geometry' metadata" in str(exc.value)


def test_xlsx_serialized_coordinate_out_of_bounds():
    """Serialized XLSX recipe with out-of-bounds coordinate against xlsx_geometry raises InvalidRecipeError."""
    cand = RawTemplateRecipeCandidate(
        template_path="nonexistent.xlsx",
        profile_id="grade_sheet_xlsx",
        fingerprint="fp1",
        roster_candidate={
            "table_index": 0,
            "worksheet_name": "Lecture",
            "first_data_row_index": 7,
            "name_col": 3,
            "id_col": 1,
            "index_col": 1,
            "capacity_limit": 10,
        },
        header_candidates=[
            {"cell_type": "xlsx_cell", "target": "Lecture!M50", "field": "instructor", "confidence": 1.0},
            {"cell_type": "xlsx_cell", "target": "Lecture!B1", "field": "course_section", "confidence": 1.0},
            {"cell_type": "xlsx_cell", "target": "Lecture!C1", "field": "schedule_code", "confidence": 1.0},
        ],
        metadata={
            "xlsx_geometry": {
                "worksheets": {
                    "Lecture": {"max_row": 20, "max_column": 10},
                }
            }
        },
    )
    with pytest.raises(InvalidRecipeError) as exc:
        RecipeValidator.validate(cand, PROFILE_GRADE_SHEET_XLSX)
    assert "out of bounds for sheet 'Lecture'" in str(exc.value)


# ═══════════════════════════════════════════════════════════════════════════════
# 4. Structural Ambiguity Rejection (Requirement 6)
# ═══════════════════════════════════════════════════════════════════════════════

def test_ambiguity_two_plausible_instructor_cells_rejected(tmp_path):
    """Two plausible instructor cells without a unique structural discriminator must raise AmbiguousTemplateError."""
    doc_path = str(tmp_path / "ambig_inst.docx")
    _create_minimal_docx(doc_path, num_tables=1, rows=4, cols=3)

    cand = RawTemplateRecipeCandidate(
        template_path=doc_path,
        profile_id="academic_docx",
        fingerprint="fp_ambig",
        roster_candidate={"table_index": 0, "first_data_row_index": 2, "name_col": 0, "id_col": 1},
        header_candidates=[
            {"cell_type": "docx_table", "target": (0, 0, 1), "field": "instructor", "confidence": 0.91},
            {"cell_type": "docx_table", "target": (0, 1, 1), "field": "instructor", "confidence": 0.87},
            {"cell_type": "docx_table", "target": (0, 0, 0), "field": "course_section", "confidence": 0.95},
            {"cell_type": "docx_table", "target": (0, 0, 2), "field": "schedule_code", "confidence": 0.95},
            {"cell_type": "docx_table", "target": (0, 1, 0), "field": "subject", "confidence": 0.95},
        ],
    )
    with pytest.raises(AmbiguousTemplateError) as exc:
        RecipeValidator.validate(cand, PROFILE_ACADEMIC_DOCX)
    assert "Ambiguous template binding for field 'instructor'" in str(exc.value)


def test_ambiguity_two_plausible_course_cells_rejected(tmp_path):
    """Two plausible course cells without structural discriminator must raise AmbiguousTemplateError."""
    doc_path = str(tmp_path / "ambig_course.docx")
    _create_minimal_docx(doc_path, num_tables=1, rows=4, cols=3)

    cand = RawTemplateRecipeCandidate(
        template_path=doc_path,
        profile_id="academic_docx",
        fingerprint="fp_ambig",
        roster_candidate={"table_index": 0, "first_data_row_index": 2, "name_col": 0, "id_col": 1},
        header_candidates=[
            {"cell_type": "docx_table", "target": (0, 0, 0), "field": "course_section", "confidence": 0.93},
            {"cell_type": "docx_table", "target": (0, 1, 0), "field": "course_section", "confidence": 0.88},
            {"cell_type": "docx_table", "target": (0, 0, 1), "field": "instructor", "confidence": 0.95},
            {"cell_type": "docx_table", "target": (0, 0, 2), "field": "schedule_code", "confidence": 0.95},
            {"cell_type": "docx_table", "target": (0, 1, 2), "field": "subject", "confidence": 0.95},
        ],
    )
    with pytest.raises(AmbiguousTemplateError) as exc:
        RecipeValidator.validate(cand, PROFILE_ACADEMIC_DOCX)
    assert "Ambiguous template binding for field 'course_section'" in str(exc.value)


def test_ambiguity_two_plausible_signature_targets_rejected(tmp_path):
    """Two plausible signature targets for the same role without structural dominance must raise AmbiguousTemplateError."""
    doc_path = str(tmp_path / "ambig_sig.docx")
    _create_minimal_docx(doc_path, num_tables=1, rows=4, cols=3)

    cand = RawTemplateRecipeCandidate(
        template_path=doc_path,
        profile_id="academic_docx",
        fingerprint="fp_ambig",
        roster_candidate={"table_index": 0, "first_data_row_index": 2, "name_col": 0, "id_col": 1},
        header_candidates=[
            {"cell_type": "docx_table", "target": (0, 0, 0), "field": "course_section", "confidence": 0.95},
            {"cell_type": "docx_table", "target": (0, 0, 1), "field": "instructor", "confidence": 0.95},
            {"cell_type": "docx_table", "target": (0, 0, 2), "field": "schedule_code", "confidence": 0.95},
            {"cell_type": "docx_table", "target": (0, 1, 2), "field": "subject", "confidence": 0.95},
        ],
        signature_candidates=[
            {"role": "instructor_signature", "target": (0, 2, 1), "confidence": 0.90},
            {"role": "instructor_signature", "target": (0, 3, 1), "confidence": 0.85},
        ],
    )
    with pytest.raises(AmbiguousTemplateError) as exc:
        RecipeValidator.validate(cand, PROFILE_ACADEMIC_DOCX)
    assert "Ambiguous signature binding for role 'instructor_signature'" in str(exc.value)


# ═══════════════════════════════════════════════════════════════════════════════
# 5. Generator Fail-Closed Safeguards
# ═══════════════════════════════════════════════════════════════════════════════

class DummyDocGenerator(DocumentGenerator):
    """Concrete DocumentGenerator for testing fail-closed execution."""
    def _fill_student_row(self, cells, idx, name, stnum):
        pass


def test_document_generator_fails_closed_on_impossible_header_coordinate(tmp_path):
    """DocumentGenerator must raise TemplateError if an impossible coordinate is consumed."""
    doc_path = str(tmp_path / "doc.docx")
    _create_minimal_docx(doc_path, num_tables=1, rows=2, cols=2)

    # Construct recipe with out-of-bounds target via internal token
    recipe = ValidatedTemplateRecipe(
        schema_version=RECIPE_SCHEMA_VERSION,
        profile_id="academic_docx",
        fingerprint="fp1",
        template_path=doc_path,
        roster_binding=RosterBinding(table_index=0, first_data_row_index=1, name_col=0, id_col=1),
        header_bindings={
            "instructor": HeaderCellBinding(cell_type="docx_table", target=(0, 99, 0), field="instructor"),
        },
        signature_bindings={},
        metadata={},
        verified_safe=True,
        _construction_token=_PRIVATE_CONSTRUCTION_SENTINEL,
    )

    from modules.models.schedule import ClassInfo

    info = ClassInfo(
        instructor="Dr. Smith",
        course_section="BSCS 3-1",
        schedule_code="12345",
        subject="CS 101",
        students=[("Student One", "123")],
    )
    gen = DummyDocGenerator(doc_path, recipe)
    with pytest.raises(TemplateError) as exc:
        gen.generate(info, str(tmp_path / "out.docx"))
    assert "out of range" in str(exc.value)


def test_grade_generator_rejects_missing_index_col(tmp_path):
    """GradeGenerator fails closed with TemplateError if recipe.roster_binding.index_col is missing."""
    template_path = str(tmp_path / "grading.xlsx")
    _create_minimal_xlsx(template_path, sheets=("Lecture", "Grading Sheet"))
    recipe = ValidatedTemplateRecipe(
        schema_version=RECIPE_SCHEMA_VERSION,
        profile_id="grade_sheet_xlsx",
        fingerprint="fp1",
        template_path=template_path,
        roster_binding=RosterBinding(
            table_index=0,
            worksheet_name="Lecture",
            first_data_row_index=7,
            name_col=3,
            id_col=1,
            index_col=None,  # Missing!
            capacity_limit=50,
        ),
        header_bindings={},
        signature_bindings={},
        metadata={},
        verified_safe=True,
        _construction_token=_PRIVATE_CONSTRUCTION_SENTINEL,
    )

    gen = GradeGenerator(template_path, recipe)
    info = {"semester": "1st Semester 2025-2026", "subject": "CS 101 - Intro", "course": "BSCS 3-1", "sched": "12345"}
    with pytest.raises(TemplateError) as exc:
        gen.generate(info, [("Student", "123")], str(tmp_path / "out.xlsx"))
    assert "missing mandatory index_col" in str(exc.value)


def test_grade_generator_rejects_missing_worksheet_name(tmp_path):
    """GradeGenerator fails closed with TemplateError if recipe.roster_binding.worksheet_name is missing."""
    template_path = str(tmp_path / "grading.xlsx")
    _create_minimal_xlsx(template_path, sheets=("Lecture", "Grading Sheet"))
    recipe = ValidatedTemplateRecipe(
        schema_version=RECIPE_SCHEMA_VERSION,
        profile_id="grade_sheet_xlsx",
        fingerprint="fp1",
        template_path=template_path,
        roster_binding=RosterBinding(
            table_index=0,
            worksheet_name=None,  # Missing!
            first_data_row_index=7,
            name_col=3,
            id_col=1,
            index_col=1,
            capacity_limit=50,
        ),
        header_bindings={},
        signature_bindings={},
        metadata={},
        verified_safe=True,
        _construction_token=_PRIVATE_CONSTRUCTION_SENTINEL,
    )

    gen = GradeGenerator(template_path, recipe)
    info = {"semester": "1st Semester 2025-2026", "subject": "CS 101 - Intro", "course": "BSCS 3-1", "sched": "12345"}
    with pytest.raises(TemplateError) as exc:
        gen.generate(info, [("Student", "123")], str(tmp_path / "out.xlsx"))
    assert "missing mandatory worksheet_name" in str(exc.value)
