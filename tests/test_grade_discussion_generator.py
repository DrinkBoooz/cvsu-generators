import os
import tempfile
import unittest
import docx

import ceit_generator


class GradeDiscussionGeneratorTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.repo_dir = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
        cls.templates_dir = os.path.join(cls.repo_dir, "templates")
        cls.factory = ceit_generator.GeneratorFactory(cls.templates_dir)

    def setUp(self):
        self.temp_dir = tempfile.TemporaryDirectory()
        self.info = ceit_generator.ClassInfo(
            instructor="PROF. DAN MICHAEL A. JOCSON",
            course_section="BSCS 4-2",
            schedule_code="202612736",
            subject="COSC 111 - SOFTWARE ENGINEERING II",
            time_days_room="7:00AM-12:00PM / F / CL4",
            semester_ay="1st Semester / 2026-2027",
            students=[
                ("AGUSTIN, JUAN P.", "202310001"),
                ("BONDOC, MARIA S.", "202310002"),
                ("DELA CRUZ, PEDRO T.", "202310003"),
            ],
        )

    def tearDown(self):
        self.temp_dir.cleanup()

    def test_factory_includes_grade_discussion(self):
        all_gens = self.factory.get_all()
        suffixes = [s for _, s in all_gens]
        self.assertIn("GRADE_DISCUSSION_MIDTERM", suffixes)
        self.assertIn("GRADE_DISCUSSION_FINALS", suffixes)

    def test_generate_midterm_grade_discussion(self):
        out_path = os.path.join(self.temp_dir.name, "midterm_gd.docx")
        recipe = self.factory._get_validated_recipe("grade_midterm", self.factory._path("grade_midterm"))
        gen = ceit_generator.GradeDiscussionGenerator(
            self.factory._path("grade_midterm"), recipe, "Midterm"
        )
        gen.generate(self.info, out_path)
        self.assertTrue(os.path.exists(out_path))

        doc = docx.Document(out_path)
        t0 = doc.tables[0]
        # Check Table 0 header mapping
        row_dict = {}
        for r in t0.rows:
            label = r.cells[0].text.strip()
            val = r.cells[2].text.strip() if len(r.cells) > 2 else ""
            row_dict[label] = val

        self.assertEqual(row_dict.get("INSTRUCTOR’S NAME/SIGNATURE") or row_dict.get("INSTRUCTOR\u2019S NAME/SIGNATURE") or row_dict.get("INSTRUCTORS NAME/SIGNATURE"), self.info.instructor)
        self.assertEqual(row_dict.get("COURSE / YEAR / SECTION"), self.info.course_section)
        self.assertEqual(row_dict.get("SCHEDULE CODE"), self.info.schedule_code)
        self.assertEqual(row_dict.get("SUBJECT CODE / TITLE"), self.info.subject)
        self.assertEqual(row_dict.get("SEMESTER / ACADEMIC YEAR"), self.info.semester_ay)
        self.assertEqual(row_dict.get("DATE"), "")

        # Check Table 1 student roster
        t1 = doc.tables[1]
        self.assertEqual(t1.rows[1].cells[0].text.strip(), "AGUSTIN, JUAN P.")
        self.assertEqual(t1.rows[1].cells[1].text.strip(), "202310001")
        self.assertEqual(t1.rows[2].cells[0].text.strip(), "BONDOC, MARIA S.")
        self.assertEqual(t1.rows[3].cells[0].text.strip(), "DELA CRUZ, PEDRO T.")

    def test_generate_finals_grade_discussion(self):
        out_path = os.path.join(self.temp_dir.name, "finals_gd.docx")
        recipe = self.factory._get_validated_recipe("grade_finals", self.factory._path("grade_finals"))
        gen = ceit_generator.GradeDiscussionGenerator(
            self.factory._path("grade_finals"), recipe, "Finals"
        )
        gen.generate(self.info, out_path)
        self.assertTrue(os.path.exists(out_path))

        doc = docx.Document(out_path)
        t0 = doc.tables[0]
        row_dict = {}
        for r in t0.rows:
            label = r.cells[0].text.strip()
            val = r.cells[2].text.strip() if len(r.cells) > 2 else ""
            row_dict[label] = val

        self.assertEqual(row_dict.get("INSTRUCTOR’S NAME/SIGNATURE") or row_dict.get("INSTRUCTOR\u2019S NAME/SIGNATURE") or row_dict.get("INSTRUCTORS NAME/SIGNATURE"), self.info.instructor)
        self.assertEqual(row_dict.get("COURSE / YEAR / SECTION"), self.info.course_section)
        self.assertEqual(row_dict.get("SCHEDULE CODE"), self.info.schedule_code)
        self.assertEqual(row_dict.get("SUBJECT CODE / TITLE"), self.info.subject)
        self.assertEqual(row_dict.get("SEMESTER / ACADEMIC YEAR"), self.info.semester_ay)
        self.assertEqual(row_dict.get("DATE"), "")

        t1 = doc.tables[1]
        self.assertEqual(t1.rows[1].cells[0].text.strip(), "AGUSTIN, JUAN P.")
        self.assertEqual(t1.rows[1].cells[1].text.strip(), "202310001")


if __name__ == "__main__":
    unittest.main()
