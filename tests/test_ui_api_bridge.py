import os
import sys
import pytest

# Ensure workspace root is on sys.path
WORKSPACE_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
if WORKSPACE_DIR not in sys.path:
    sys.path.insert(0, WORKSPACE_DIR)

import process_schedule
from executable.main import ScriptAPI

def test_inspect_schedule_file_valid():
    schedule_path = os.path.join(WORKSPACE_DIR, "ORTEGA_SCHEDULE.xls")
    assert os.path.exists(schedule_path), "Test schedule ORTEGA_SCHEDULE.xls must exist"

    info = process_schedule.inspect_schedule_file(schedule_path)
    assert info is not None
    assert "DAN JOSEPH A. ORTEGA" in info["instructor"]
    assert "COLLEGE OF ENGINEERING" in info["college"].upper()
    assert info["total_slots"] > 0
    assert len(info["detected_sections"]) > 0

def test_inspect_schedule_file_invalid():
    res = process_schedule.inspect_schedule_file("non_existent_file.xls")
    assert res is None

def test_validate_rosters_valid_and_invalid(tmp_path):
    schedule_path = os.path.join(WORKSPACE_DIR, "ORTEGA_SCHEDULE.xls")
    
    # Valid filename
    valid_roster = tmp_path / "BSCS1-4 List of Students for 202612040-DCIT 21A - INTRODUCTION TO COMPUTING.csv"
    valid_roster.write_text("Name,Student number\nDoe, John,20261001\nSmith, Jane,20261002\n", encoding="utf-8")

    # Invalid filename
    invalid_roster = tmp_path / "Random_Student_List.csv"
    invalid_roster.write_text("Name,Student number\nDoe, John,20261001\n", encoding="utf-8")

    # Extra columns roster
    extra_col_roster = tmp_path / "BSCS3-1 List of Students for 202612041-ITEC 50 - WEB SYSTEMS.csv"
    extra_col_roster.write_text("Name,Student number,Email,Remarks\nDoe, John,20261001,john@cvsu.edu.ph,Regular\n", encoding="utf-8")

    reports = process_schedule.validate_rosters(schedule_path, [
        str(valid_roster),
        str(invalid_roster),
        str(extra_col_roster)
    ])

    assert len(reports) == 3

    # Report 0: Valid
    assert reports[0]["status"] == "valid"
    assert reports[0]["student_count"] == 2
    assert reports[0]["course_sec"] == "CS1-4"
    assert reports[0]["schedule_code"] == "202612040"

    # Report 1: Invalid filename
    assert reports[1]["status"] == "error"
    assert reports[1]["issue"] == "invalid_filename"

    # Report 2: Extra columns warning
    assert reports[2]["status"] == "warning"
    assert reports[2]["issue"] == "extra_columns"
    assert reports[2]["column_count"] == 4
    assert reports[2]["student_count"] == 1

def test_validate_and_process_xlsx_extra_columns(tmp_path):
    import openpyxl
    import ceit_generator
    import attendancegen

    schedule_path = os.path.join(WORKSPACE_DIR, "ORTEGA_SCHEDULE.xls")
    assert os.path.exists(schedule_path)

    # Create an actual .xlsx file with 8 columns (Name, Student number + 6 extra columns)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    headers = ["Name", "Student number", "Email", "Course", "Year", "Section", "Status", "Remarks"]
    ws.append(headers)
    ws.append(["DELA CRUZ, JUAN A.", "202110001", "juan.delacruz@cvsu.edu.ph", "BSCS", "4", "1", "Enrolled", "Regular"])
    ws.append(["SANTOS, MARIA B.", "202110002", "maria.santos@cvsu.edu.ph", "BSCS", "4", "1", "Enrolled", "Irregular"])

    xlsx_path = os.path.join(tmp_path, "BSCS4-1 List of Students for 202612731-COSC 111A - C S ELECTIVE 3 (INTERNET OF THINGS).xlsx")
    wb.save(xlsx_path)

    # 1. Test load_students isolation
    students_ceit = ceit_generator.load_students(xlsx_path)
    assert len(students_ceit) == 2
    assert students_ceit[0] == ("DELA CRUZ, JUAN A.", "202110001")
    assert students_ceit[1] == ("SANTOS, MARIA B.", "202110002")

    students_att = attendancegen.load_students(xlsx_path)
    assert len(students_att) == 2
    assert students_att[0] == ("DELA CRUZ, JUAN A.", "202110001")

    # 2. Test pre-flight roster validation
    val = process_schedule.validate_rosters(schedule_path, [xlsx_path])
    assert len(val) == 1
    assert val[0]["status"] == "warning"
    assert val[0]["issue"] == "extra_columns"
    assert val[0]["column_count"] == 8
    assert val[0]["student_count"] == 2
    assert "auto-cleaned" in val[0]["message"]

    # 3. Test end-to-end execution with extra columns
    out_dir = os.path.join(tmp_path, "output")
    res = process_schedule.process_all(
        schedule_path=schedule_path,
        xlsx_files=[xlsx_path],
        output_dir_base=out_dir
    )
    assert len(res["generated"]["ceit"]) > 0
    assert len(res["generated"]["attendance"]) > 0
    assert len(res["generated"]["grades"]) > 0
    assert len(res["errors"]["ceit"]) == 0
    assert len(res["errors"]["attendance"]) == 0
    assert len(res["errors"]["grades"]) == 0

def test_intelligent_header_detection_arbitrary_columns(tmp_path):
    import openpyxl
    import ceit_generator
    import attendancegen

    schedule_path = os.path.join(WORKSPACE_DIR, "ORTEGA_SCHEDULE.xls")

    # Case 1: Swapped columns in XLSX (Col A: Student Number, Col B: Student Name)
    wb1 = openpyxl.Workbook()
    ws1 = wb1.active
    ws1.append(["Student Number", "Student Name"])
    ws1.append(["202110001", "DELA CRUZ, JUAN A."])
    p1 = os.path.join(tmp_path, "BSCS4-1 List of Students for 202612731-COSC 111A - C S ELECTIVE 3 (INTERNET OF THINGS).xlsx")
    wb1.save(p1)

    s1 = ceit_generator.load_students(p1)
    assert len(s1) == 1
    assert s1[0] == ("DELA CRUZ, JUAN A.", "202110001")
    s1_att = attendancegen.load_students(p1)
    assert s1_att[0] == ("DELA CRUZ, JUAN A.", "202110001")

    # Case 2: Leading '#' / 'No.' column in XLSX (Col A: '#', Col B: Student ID, Col C: Full Name, Col D: Email)
    wb2 = openpyxl.Workbook()
    ws2 = wb2.active
    ws2.append(["#", "Student ID", "Full Name", "Email"])
    ws2.append(["1", "202110002", "SANTOS, MARIA B.", "maria@cvsu.edu.ph"])
    p2 = os.path.join(tmp_path, "case2_leading_no.xlsx")
    wb2.save(p2)

    s2 = ceit_generator.load_students(p2)
    assert len(s2) == 1
    assert s2[0] == ("SANTOS, MARIA B.", "202110002")

    # Case 3: Top title metadata in CSV before actual headers
    p3 = os.path.join(tmp_path, "case3_top_metadata.csv")
    with open(p3, "w", encoding="utf-8") as f:
        f.write("CAVITE STATE UNIVERSITY\n")
        f.write("OFFICIAL CLASS LIST 2026\n")
        f.write("Department,Student No.,Student's Name,Status\n")
        f.write('DIT,202110003,"REYES, CARLOS C.",Regular\n')

    s3 = ceit_generator.load_students(p3)
    assert len(s3) == 1
    assert s3[0] == ("REYES, CARLOS C.", "202110003")

    # Case 4: Headerless swapped CSV (Col 0: ID, Col 1: Name)
    p4 = os.path.join(tmp_path, "case4_headerless_swapped.csv")
    with open(p4, "w", encoding="utf-8") as f:
        f.write('202110004,"GARCIA, ANA D."\n')
        f.write('202110005,"LOPEZ, MARK E."\n')


    s4 = ceit_generator.load_students(p4)
    assert len(s4) == 2
    assert s4[0] == ("GARCIA, ANA D.", "202110004")
    assert s4[1] == ("LOPEZ, MARK E.", "202110005")

    # Case 5: Empty roster validation check (headers only, no students)
    p5 = os.path.join(tmp_path, "BSCS1-4 List of Students for 202612040-DCIT 21A - INTRODUCTION TO COMPUTING.csv")
    with open(p5, "w", encoding="utf-8") as f:
        f.write("Name,Student number\n")

    val = process_schedule.validate_rosters(schedule_path, [p5])
    assert len(val) == 1
    assert val[0]["status"] == "warning"
    assert val[0]["issue"] == "no_students"
    assert "No students detected" in val[0]["message"]

def test_script_api_methods(tmp_path):
    api = ScriptAPI()
    assert api.schedule_path == ""
    assert api.rosters == []
    assert api.output_dir == ""

    # Test clear and remove roster
    api.rosters = ["path1.xlsx", "path2.xlsx", "path3.xlsx"]
    res = api.remove_roster(1)
    assert res["count"] == 2
    assert api.rosters == ["path1.xlsx", "path3.xlsx"]

    res = api.clear_rosters()
    assert res["count"] == 0
    assert api.rosters == []

def test_process_all_engine_and_class_filter(tmp_path):
    schedule_path = os.path.join(WORKSPACE_DIR, "ORTEGA_SCHEDULE.xls")
    
    roster_file = tmp_path / "BSCS1-4 List of Students for 202612040-DCIT 21A - INTRODUCTION TO COMPUTING.csv"
    roster_file.write_text("Name,Student number\nOrtega, Dan,20261001\n", encoding="utf-8")

    out_dir = tmp_path / "output"
    out_dir.mkdir()

    progress_events = []
    def on_progress(event):
        progress_events.append(event)

    # Test running ONLY grade sheet generator
    results = process_schedule.process_all(
        schedule_path=schedule_path,
        xlsx_files=[str(roster_file)],
        output_dir_base=str(out_dir),
        class_filter=["202612040_CS1-4"],
        engine_filter=["grades"],
        progress_callback=on_progress
    )

    assert len(results["generated"]["grades"]) == 1
    assert len(results["generated"]["ceit"]) == 0
    assert len(results["generated"]["attendance"]) == 0
    assert len(progress_events) > 0
    assert progress_events[-1]["percent"] > 0

def test_handle_dropped_schedule_and_rosters(tmp_path):
    import base64
    schedule_path = os.path.join(WORKSPACE_DIR, "ORTEGA_SCHEDULE.xls")
    
    api = ScriptAPI()
    
    # 1. Test dropped schedule with direct path
    res = api.handle_dropped_schedule(
        filename="ORTEGA_SCHEDULE.xls",
        original_path=schedule_path
    )
    assert res["path"] == schedule_path
    assert res["metadata"] is not None
    assert "ORTEGA" in res["metadata"]["instructor"]

    # 2. Test dropped schedule with base64 payload
    with open(schedule_path, "rb") as f:
        b64_sched = base64.b64encode(f.read()).decode("utf-8")
    
    res2 = api.handle_dropped_schedule(
        filename="dropped_sched.xls",
        base64_data=b64_sched
    )
    assert res2["path"] != ""
    assert os.path.exists(res2["path"])
    assert res2["metadata"] is not None

    # 3. Test dropped rosters with base64 payload
    roster_content = b"Name,Student number\nTest Student,20261234\n"
    b64_roster = base64.b64encode(roster_content).decode("utf-8")

    roster_filename = "BSCS1-4 List of Students for 202612040-DCIT 21A - INTRODUCTION TO COMPUTING.csv"
    res3 = api.handle_dropped_rosters([
        {
            "filename": roster_filename,
            "data": b64_roster,
            "path": None
        }
    ])
    assert res3["count"] == 1
    assert len(res3["rosters"]) == 1
    assert os.path.exists(res3["rosters"][0])
    assert len(res3["validation"]) == 1
    assert res3["validation"][0]["status"] == "valid"

def test_dynamic_total_steps_telemetry(tmp_path):
    schedule_path = os.path.join(WORKSPACE_DIR, "ORTEGA_SCHEDULE.xls")
    roster_file = tmp_path / "BSCS1-4 List of Students for 202612040-DCIT 21A - INTRODUCTION TO COMPUTING.csv"
    roster_file.write_text("Name,Student number\nOrtega, Dan,20261001\n", encoding="utf-8")

    out_dir = tmp_path / "output"
    out_dir.mkdir()

    progress_events = []
    def on_progress(event):
        progress_events.append(event)

    # 6-month date overrides (Aug to Jan): 7 CEIT + 6 Attendance + 1 Grades = 14 steps
    date_overrides = {
        "startYear": 2026, "startMonth": 8, "startDay": 1,
        "endYear": 2027, "endMonth": 1, "endDay": 31
    }

    results = process_schedule.process_all(
        schedule_path=schedule_path,
        xlsx_files=[str(roster_file)],
        output_dir_base=str(out_dir),
        class_filter=["202612040_CS1-4"],
        engine_filter=["ceit", "attendance", "grades"],
        date_overrides=date_overrides,
        progress_callback=on_progress
    )

    assert len(results["generated"]["ceit"]) == 7
    assert len(results["generated"]["attendance"]) == 6
    assert len(results["generated"]["grades"]) == 1
    assert len(progress_events) == 14

    for ev in progress_events:
        assert ev["step"] <= ev["total_steps"]
        assert ev["total_steps"] == 14

    assert progress_events[-1]["step"] == 14
    assert progress_events[-1]["total_steps"] == 14

def test_cancel_generation_and_by_class_artifacts(tmp_path):
    import threading
    schedule_path = os.path.join(WORKSPACE_DIR, "ORTEGA_SCHEDULE.xls")
    roster_file = tmp_path / "BSCS1-4 List of Students for 202612040-DCIT 21A - INTRODUCTION TO COMPUTING.csv"
    roster_file.write_text("Name,Student number\nOrtega, Dan,20261001\n", encoding="utf-8")

    out_dir = tmp_path / "output"
    out_dir.mkdir()

    cancel_evt = threading.Event()
    cancel_evt.set()  # Pre-cancel before execution

    results = process_schedule.process_all(
        schedule_path=schedule_path,
        xlsx_files=[str(roster_file)],
        output_dir_base=str(out_dir),
        class_filter=["202612040_CS1-4"],
        engine_filter=["ceit", "attendance", "grades"],
        cancel_event=cancel_evt
    )

    assert results["cancelled"] is True
    # Verify by_class exists in results
    assert "by_class" in results

