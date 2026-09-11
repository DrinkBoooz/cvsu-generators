"""
Script to create multiple distinct synthetic .docx test templates for verifying
the Deterministic Heuristic Template Analyzer and ConfigurableDocumentGenerator.
"""

import os
import docx
from docx.shared import Inches, Pt, RGBColor


def build_template_consultation_log(dest_path: str):
    """
    Format 1: Classic Academic Consultation Log
    - 2-column header table with colon metadata labels.
    - 4-column roster table: [No., Student Number, Name of Student, Signature].
    """
    doc = docx.Document()
    doc.add_heading("STUDENT CONSULTATION LOG SHEET", level=1)

    # Header metadata table (4 rows, 2 columns)
    h_tbl = doc.add_table(rows=4, cols=2)
    h_tbl.autofit = False

    metadata_rows = [
        ("Instructor Name:", "Dr. Jane Doe"),
        ("Course & Section:", "BSCS 3-1"),
        ("Subject Descriptive Title:", "COSC 101 - Advanced Software Engineering"),
        ("Schedule Code:", "12345"),
    ]
    for r_idx, (label, val) in enumerate(metadata_rows):
        row = h_tbl.rows[r_idx]
        row.cells[0].text = label
        row.cells[1].text = val

    doc.add_paragraph("")  # Spacer

    # Student roster table (header row + 1 template row)
    r_tbl = doc.add_table(rows=2, cols=4)
    r_tbl.autofit = False

    headers = ["No.", "Student Number", "Name of Student", "Signature"]
    for c_idx, h in enumerate(headers):
        r_tbl.rows[0].cells[c_idx].text = h

    # Template row
    template_cells = ["1", "202310001", "Dela Cruz, Juan M.", ""]
    for c_idx, val in enumerate(template_cells):
        r_tbl.rows[1].cells[c_idx].text = val

    doc.save(dest_path)


def build_template_guidance_advising(dest_path: str):
    """
    Format 2: Academic Advising & Guidance Form
    - 3-column colon header table (Label | : | Value).
    - 5-column roster table: [Item, Student ID, Student Name, Concern, Action Taken].
    """
    doc = docx.Document()
    doc.add_heading("FACULTY ADVISING RECORD", level=1)

    # Header metadata table (4 rows, 3 columns: Label, :, Value)
    h_tbl = doc.add_table(rows=4, cols=3)
    metadata_rows = [
        ("Faculty Member", ":", "Prof. John Smith"),
        ("Degree Program & Section", ":", "BSINFOTECH 2-2"),
        ("Class Schedule Code", ":", "67890"),
        ("Subject Code & Title", ":", "ITEC 50 - Web Systems"),
    ]
    for r_idx, (lbl, sep, val) in enumerate(metadata_rows):
        row = h_tbl.rows[r_idx]
        row.cells[0].text = lbl
        row.cells[1].text = sep
        row.cells[2].text = val

    doc.add_paragraph("")

    # Student roster table
    r_tbl = doc.add_table(rows=2, cols=5)
    headers = ["Item", "Student ID", "Student Name", "Concern", "Action Taken"]
    for c_idx, h in enumerate(headers):
        r_tbl.rows[0].cells[c_idx].text = h

    template_cells = ["1", "202420002", "Santos, Maria L.", "Prerequisite inquiry", "Advised"]
    for c_idx, val in enumerate(template_cells):
        r_tbl.rows[1].cells[c_idx].text = val

    doc.save(dest_path)


def build_template_tag_placeholders(dest_path: str):
    """
    Format 3: Tag-Based Document
    - Metadata embedded via explicit tags {{INSTRUCTOR}}, {{COURSE_SECTION}}, etc. in paragraphs.
    - 3-column roster table: [LRN, Name, Remarks].
    """
    doc = docx.Document()
    doc.add_heading("LABORATORY CLEARANCE REPORT", level=1)

    doc.add_paragraph("Instructor: {{INSTRUCTOR}}")
    doc.add_paragraph("Course and Section: {{COURSE_SECTION}} | Sched Code: {{SCHEDULE_CODE}}")
    doc.add_paragraph("Subject: {{SUBJECT}}")
    doc.add_paragraph("Class Schedule: {{TIME_DAYS_ROOM}} | Semester/AY: {{SEMESTER_AY}}")

    doc.add_paragraph("")

    # Student roster table
    r_tbl = doc.add_table(rows=2, cols=3)
    headers = ["LRN", "Name", "Remarks"]
    for c_idx, h in enumerate(headers):
        r_tbl.rows[0].cells[c_idx].text = h

    template_cells = ["100234567890", "Reyes, Jose A.", "Cleared"]
    for c_idx, val in enumerate(template_cells):
        r_tbl.rows[1].cells[c_idx].text = val

    doc.save(dest_path)


def build_template_laboratory_monitoring(dest_path: str):
    """
    Format 4: Computer Laboratory Utilization Log
    - 4-column header table with 2 pairs of (Label, Value) per row.
    - 6-column roster table: [#, Student's Name, Stud No, Workstation No., Time In, Time Out].
    """
    doc = docx.Document()
    doc.add_heading("COMPUTER LABORATORY LOG SHEET", level=1)

    h_tbl = doc.add_table(rows=3, cols=4)
    # Row 0
    h_tbl.rows[0].cells[0].text = "Teacher:"
    h_tbl.rows[0].cells[1].text = "Engr. Alan Turing"
    h_tbl.rows[0].cells[2].text = "Course & Section:"
    h_tbl.rows[0].cells[3].text = "BSCpE 4-1"

    # Row 1
    h_tbl.rows[1].cells[0].text = "Subject:"
    h_tbl.rows[1].cells[1].text = "CPEG 300 - Embedded Systems"
    h_tbl.rows[1].cells[2].text = "Schedule Code:"
    h_tbl.rows[1].cells[3].text = "99887"

    # Row 2
    h_tbl.rows[2].cells[0].text = "Time & Days & Room:"
    h_tbl.rows[2].cells[1].text = "TF 7:00-10:00 AM CEIT LAB 1"
    h_tbl.rows[2].cells[2].text = "Semester & Academic Year:"
    h_tbl.rows[2].cells[3].text = "First Semester, AY 2026-2027"

    doc.add_paragraph("")

    r_tbl = doc.add_table(rows=2, cols=6)
    headers = ["#", "Student's Name", "Stud No", "Workstation No.", "Time In", "Time Out"]
    for c_idx, h in enumerate(headers):
        r_tbl.rows[0].cells[c_idx].text = h

    template_cells = ["1", "Aquino, Benigno C.", "202110443", "PC-01", "", ""]
    for c_idx, val in enumerate(template_cells):
        r_tbl.rows[1].cells[c_idx].text = val

    doc.save(dest_path)


def build_template_faculty_eval(dest_path: str):
    """
    Format 5: Faculty Evaluation / Survey Form
    - Paragraph headers with colons.
    - Filipino / Tagalog roster tokens: [Index, Pangalan, Numero, Lagda].
    """
    doc = docx.Document()
    doc.add_heading("FACULTY TEACHING EVALUATION FORM", level=1)

    doc.add_paragraph("Professor: Prof. Ada Lovelace")
    doc.add_paragraph("Section: BSCS 1-1")
    doc.add_paragraph("Subject: COSC 50 - Discrete Mathematics")
    doc.add_paragraph("Schedule Code: 54321")

    doc.add_paragraph("")

    r_tbl = doc.add_table(rows=2, cols=4)
    headers = ["Index", "Pangalan", "Numero", "Lagda"]
    for c_idx, h in enumerate(headers):
        r_tbl.rows[0].cells[c_idx].text = h

    template_cells = ["1", "Luna, Antonio N.", "202510111", ""]
    for c_idx, val in enumerate(template_cells):
        r_tbl.rows[1].cells[c_idx].text = val

    doc.save(dest_path)


def generate_all_templates(base_dir: str):
    os.makedirs(base_dir, exist_ok=True)
    templates = {
        "template_consultation_log.docx": build_template_consultation_log,
        "template_guidance_advising.docx": build_template_guidance_advising,
        "template_tag_placeholders.docx": build_template_tag_placeholders,
        "template_laboratory_monitoring.docx": build_template_laboratory_monitoring,
        "template_faculty_eval.docx": build_template_faculty_eval,
    }

    generated_paths = {}
    for filename, builder in templates.items():
        filepath = os.path.join(base_dir, filename)
        builder(filepath)
        generated_paths[filename] = filepath
        print(f"Generated test template: {filepath}")

    return generated_paths


if __name__ == "__main__":
    current_dir = os.path.dirname(os.path.abspath(__file__))
    generate_all_templates(current_dir)
