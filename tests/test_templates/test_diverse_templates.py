"""
tests/test_templates/test_diverse_templates.py

Automated test suite verifying the Deterministic Heuristic Template Analyzer
and ConfigurableDocumentGenerator across diverse template layouts and formats
located in tests/test_templates/.
"""

import os
import pytest
import docx

from modules.models.schedule import ClassInfo
from modules.parsers.template_inspector import TemplateInspector
from modules.generators.generic_doc_gen import ConfigurableDocumentGenerator
from modules.common.config_manager import ConfigManager
from modules.generators.ceit_gen import GeneratorFactory


TEMPLATES_DIR = os.path.dirname(os.path.abspath(__file__))

TEMPLATES_TO_TEST = [
    "template_consultation_log.docx",
    "template_guidance_advising.docx",
    "template_tag_placeholders.docx",
    "template_laboratory_monitoring.docx",
    "template_faculty_eval.docx",
]


@pytest.fixture
def sample_class_info():
    return ClassInfo(
        instructor="Dr. Ada Lovelace",
        course_section="BSCS 3-1",
        schedule_code="77123",
        subject="COSC 110 - Compiler Construction",
        time_days_room="TF 1:00-3:00 PM CL3",
        semester_ay="Second Semester, AY 2026-2027",
        students=[
            ("Dela Cruz, Juan M.", "202310001"),
            ("Santos, Maria Clara L.", "202310002"),
            ("Rizal, Jose Protacio Mercado Y Alonso", "202310003"),  # Long name for font scaling
            ("Bonifacio, Andres C.", "202310004"),
        ],
    )


class TestDiverseTemplatesInspection:
    """Verifies that the TemplateInspector correctly parses diverse template structures."""

    @pytest.mark.parametrize("template_name", TEMPLATES_TO_TEST)
    def test_template_file_exists(self, template_name):
        path = os.path.join(TEMPLATES_DIR, template_name)
        assert os.path.isfile(path), f"Test template file missing: {path}"

    @pytest.mark.parametrize("template_name", TEMPLATES_TO_TEST)
    def test_inspection_confidence_and_validity(self, template_name):
        path = os.path.join(TEMPLATES_DIR, template_name)
        inspector = TemplateInspector()
        recipe = inspector.inspect_docx(path)

        assert recipe is not None
        assert recipe["confidence"] >= 90, f"Confidence too low ({recipe['confidence']}%) for {template_name}"
        assert recipe["title"], "Document title should not be empty"
        assert recipe["suffix"], "Document suffix should not be empty"

    @pytest.mark.parametrize("template_name", TEMPLATES_TO_TEST)
    def test_roster_table_detected(self, template_name):
        path = os.path.join(TEMPLATES_DIR, template_name)
        inspector = TemplateInspector()
        recipe = inspector.inspect_docx(path)

        roster = recipe.get("roster_table")
        assert roster is not None, f"Roster table not detected in {template_name}"
        assert roster["name_col"] is not None, f"name_col not detected in {template_name}"
        assert roster["name_col"] != roster["id_col"], f"name_col and id_col overlap in {template_name}"

    def test_consultation_log_specific_mappings(self):
        path = os.path.join(TEMPLATES_DIR, "template_consultation_log.docx")
        recipe = TemplateInspector().inspect_docx(path)
        roster = recipe["roster_table"]

        assert roster["index_col"] == 0  # No.
        assert roster["id_col"] == 1     # Student Number
        assert roster["name_col"] == 2   # Name of Student
        assert roster["signature_col"] == 3

    def test_guidance_advising_specific_mappings(self):
        path = os.path.join(TEMPLATES_DIR, "template_guidance_advising.docx")
        recipe = TemplateInspector().inspect_docx(path)
        roster = recipe["roster_table"]

        assert roster["index_col"] == 0  # Item
        assert roster["id_col"] == 1     # Student ID
        assert roster["name_col"] == 2   # Student Name

    def test_tag_placeholders_specific_mappings(self):
        path = os.path.join(TEMPLATES_DIR, "template_tag_placeholders.docx")
        recipe = TemplateInspector().inspect_docx(path)

        placeholders = {p["field"] for p in recipe.get("placeholders", [])}
        assert "instructor" in placeholders
        assert "course_section" in placeholders
        assert "schedule_code" in placeholders
        assert "subject" in placeholders
        assert "time_days_room" in placeholders
        assert "semester_ay" in placeholders

    def test_faculty_eval_specific_mappings(self):
        path = os.path.join(TEMPLATES_DIR, "template_faculty_eval.docx")
        recipe = TemplateInspector().inspect_docx(path)
        roster = recipe["roster_table"]

        assert roster["index_col"] == 0  # Index
        assert roster["name_col"] == 1   # Pangalan
        assert roster["id_col"] == 2     # Numero
        assert roster["signature_col"] == 3  # Lagda


class TestDiverseTemplatesGeneration:
    """Verifies document generation across all template variants."""

    @pytest.mark.parametrize("template_name", TEMPLATES_TO_TEST)
    def test_document_generation_and_readback(self, template_name, sample_class_info, tmp_path):
        tmpl_path = os.path.join(TEMPLATES_DIR, template_name)
        recipe = TemplateInspector().inspect_docx(tmpl_path)
        gen = ConfigurableDocumentGenerator(tmpl_path, recipe)

        out_path = str(tmp_path / f"out_{template_name}")
        gen.generate(sample_class_info, out_path)

        assert os.path.isfile(out_path), f"Generated file was not created: {out_path}"

        # Read back with docx and verify contents
        doc = docx.Document(out_path)

        # 1. Verify roster row count (1 header + 4 students)
        tbl_idx = recipe["roster_table"]["table_index"]
        roster_tbl = doc.tables[tbl_idx]
        assert len(roster_tbl.rows) == 1 + len(sample_class_info.students)

        # 2. Verify student names and IDs in rows
        name_col = recipe["roster_table"]["name_col"]
        id_col = recipe["roster_table"]["id_col"]

        for idx, (expected_name, expected_id) in enumerate(sample_class_info.students, 1):
            row_cells = roster_tbl.rows[idx].cells
            assert expected_name in row_cells[name_col].text
            if id_col is not None:
                assert expected_id in row_cells[id_col].text

        # 3. Verify metadata presence
        all_text = " ".join(p.text for p in doc.paragraphs)
        for tbl in doc.tables:
            for r in tbl.rows:
                all_text += " " + " ".join(c.text for c in r.cells)

        assert "Dr. Ada Lovelace" in all_text
        assert "BSCS 3-1" in all_text
        assert "77123" in all_text

    def test_generation_with_empty_student_roster(self, tmp_path):
        """Edge case: 0 students."""
        tmpl_path = os.path.join(TEMPLATES_DIR, "template_consultation_log.docx")
        recipe = TemplateInspector().inspect_docx(tmpl_path)
        gen = ConfigurableDocumentGenerator(tmpl_path, recipe)

        empty_info = ClassInfo(
            instructor="Prof. Test",
            course_section="BSIT 1-1",
            schedule_code="11111",
            subject="Test Subject",
            time_days_room="M 8-11 AM",
            semester_ay="1st Sem",
            students=[],
        )

        out_path = str(tmp_path / "out_empty.docx")
        gen.generate(empty_info, out_path)
        assert os.path.isfile(out_path)

        doc = docx.Document(out_path)
        roster_tbl = doc.tables[recipe["roster_table"]["table_index"]]
        assert len(roster_tbl.rows) == 1  # Only header row remains

    def test_generation_with_large_student_roster(self, tmp_path):
        """Edge case: 50 students."""
        tmpl_path = os.path.join(TEMPLATES_DIR, "template_laboratory_monitoring.docx")
        recipe = TemplateInspector().inspect_docx(tmpl_path)
        gen = ConfigurableDocumentGenerator(tmpl_path, recipe)

        large_students = [(f"Student {i}, Test M.", f"2023{i:05d}") for i in range(1, 51)]
        large_info = ClassInfo(
            instructor="Engr. Lab Admin",
            course_section="BSCpE 3-1",
            schedule_code="99001",
            subject="Microprocessors",
            time_days_room="W 1-4 PM",
            semester_ay="2nd Sem",
            students=large_students,
        )

        out_path = str(tmp_path / "out_large.docx")
        gen.generate(large_info, out_path)
        assert os.path.isfile(out_path)

        doc = docx.Document(out_path)
        roster_tbl = doc.tables[recipe["roster_table"]["table_index"]]
        assert len(roster_tbl.rows) == 51  # 1 header + 50 students


class TestCustomTemplateLifecycle:
    """Verifies end-to-end management of custom templates in config and factory."""

    def test_save_toggle_and_factory_instantiation(self, tmp_path, sample_class_info):
        from modules.common.config_manager import ParserConfigManager
        from modules.generators.ceit_gen import GeneratorFactory

        cfg_mgr = ParserConfigManager(config_dir=str(tmp_path / "config"))
        tmpl_path = os.path.join(TEMPLATES_DIR, "template_consultation_log.docx")
        recipe = TemplateInspector().inspect_docx(tmpl_path)

        # 1. Save custom template
        res = cfg_mgr.save_custom_template(
            source_path=tmpl_path,
            title="Student Consultation Sheet",
            suffix="CONSULTATION_LOG",
            recipe=recipe,
            enabled=True,
        )
        assert res["status"] == "success"
        template_id = res["template"]["id"]

        # 2. Verify registered
        templates = cfg_mgr.get_custom_templates()
        assert any(t["id"] == template_id for t in templates)

        # 3. Factory produces the custom generator
        factory = GeneratorFactory("templates", config_manager=cfg_mgr)
        all_gens = factory.get_all(include_custom=True)
        custom_suffixes = [suffix for _, suffix in all_gens]
        assert "CONSULTATION_LOG" in custom_suffixes

        # 4. Generate through factory instance
        gen_factory_fn = next(fn for fn, sfx in all_gens if sfx == "CONSULTATION_LOG")
        generator = gen_factory_fn()
        out_file = str(tmp_path / "factory_out.docx")
        generator.generate(sample_class_info, out_file)
        assert os.path.isfile(out_file)

        # 5. Toggle disable
        cfg_mgr.toggle_custom_template(template_id, False)
        all_gens_after_disable = factory.get_all(include_custom=True)
        assert "CONSULTATION_LOG" not in [sfx for _, sfx in all_gens_after_disable]

        # 6. Delete
        del_res = cfg_mgr.delete_custom_template(template_id)
        assert del_res["status"] == "success"
        assert len(cfg_mgr.get_custom_templates()) == 0

