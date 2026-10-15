#!/usr/bin/env python3
"""
tests/test_dynamic_org_and_signatories.py

Comprehensive TDD test suite for Phase 6:
100% Dynamic Organizational Topology & N-Tier Signatory Chain Engine.

Validates:
1. Deterministic extraction of OrganizationalContext (Institution, Campus, College, Department)
   across Main Campus colleges (CEIT, CAS, CED, etc.), Graduate School (GS-OLC), and Satellite Campuses (Tanza, Naic, Rosario, etc.)
   for both .docx and .xlsx templates with zero hardcoded campus/college lists.
2. Dynamic discovery of SignatorySlots across table grids and paragraph flows with role labels,
   default titles, and bound keys.
3. Style-preserving rendering of N-tier signatory chains (Main Campus with Dean, Satellite Campus with Campus Administrator,
   Graduate School with GS Dean) with 100% backward compatibility for existing department_chair and college_dean keys.
4. Shareable Department Presets (.cvsu-org.json) export and import with strict path validation and schema integrity.
5. Integration with TemplateMixin manifest_summary and ConfigMixin bridge APIs.
"""

import os
import json
import pytest
import docx
import openpyxl
from lxml import etree

from modules.models.organization import OrganizationalContext, SignatorySlot
from modules.models.schedule import ClassInfo
from modules.parsers.template_analyzer import DeterministicTemplateAnalyzer, TemplateManifest
from modules.generators.dynamic_renderer import DynamicDocumentRenderer
from executable_test.api.config import ConfigMixin
from executable_test.api.templates import TemplateMixin


# ─────────────────────────────────────────────────────────────────────────────
# Test Fixtures & Synthetic Template Builders
# ─────────────────────────────────────────────────────────────────────────────

def create_main_campus_ceit_docx(file_path: str) -> None:
    """Builds a DOCX with standard CvSU Main Campus CEIT header and 3-tier signatory footer."""
    doc = docx.Document()
    # Top banner paragraphs
    doc.add_paragraph("Republic of the Philippines")
    doc.add_paragraph("CAVITE STATE UNIVERSITY")
    doc.add_paragraph("Don Severino delas Alas Campus")
    doc.add_paragraph("Indang, Cavite")
    doc.add_paragraph("COLLEGE OF ENGINEERING AND INFORMATION TECHNOLOGY")
    doc.add_paragraph("Department of Information Technology")
    doc.add_paragraph("CLASS ACCEPTANCE FORM")

    # Header metadata table
    meta_tbl = doc.add_table(rows=3, cols=2)
    meta_tbl.rows[0].cells[0].text = "Instructor's Name and Signature:"
    meta_tbl.rows[0].cells[1].text = ""
    meta_tbl.rows[1].cells[0].text = "Course / Year / Section:"
    meta_tbl.rows[1].cells[1].text = ""
    meta_tbl.rows[2].cells[0].text = "Schedule Code:"
    meta_tbl.rows[2].cells[1].text = ""

    doc.add_paragraph("")

    # Roster table
    r_tbl = doc.add_table(rows=2, cols=3)
    r_tbl.rows[0].cells[0].text = "No."
    r_tbl.rows[0].cells[1].text = "Name of Student"
    r_tbl.rows[0].cells[2].text = "Student Number"
    r_tbl.rows[1].cells[0].text = "1"
    r_tbl.rows[1].cells[1].text = "Doe, Jane"
    r_tbl.rows[1].cells[2].text = "2026-0001"

    doc.add_paragraph("")

    # 3-tier Main Campus Signatory Footer Table
    sig_tbl = doc.add_table(rows=3, cols=3)
    # Row 0: Labels
    sig_tbl.rows[0].cells[0].text = "Prepared by:"
    sig_tbl.rows[0].cells[1].text = "Reviewed by:"
    sig_tbl.rows[0].cells[2].text = "Approved by:"
    # Row 1: Target cells for signatory names
    sig_tbl.rows[1].cells[0].text = ""
    sig_tbl.rows[1].cells[1].text = ""
    sig_tbl.rows[1].cells[2].text = ""
    # Row 2: Default titles
    sig_tbl.rows[2].cells[0].text = "Instructor"
    sig_tbl.rows[2].cells[1].text = "Department Chairperson"
    sig_tbl.rows[2].cells[2].text = "College Dean"

    doc.save(file_path)


def create_satellite_campus_tanza_docx(file_path: str) -> None:
    """Builds a DOCX with CvSU - Tanza Campus header and 2-tier satellite signatory chain (Chair -> Campus Admin)."""
    doc = docx.Document()
    # Top banner paragraphs
    doc.add_paragraph("Republic of the Philippines")
    doc.add_paragraph("CAVITE STATE UNIVERSITY")
    doc.add_paragraph("CvSU - Tanza Campus")
    doc.add_paragraph("Tanza, Cavite")
    doc.add_paragraph("Department of Information Technology")
    doc.add_paragraph("COURSE SYLLABUS ACCEPTANCE FORM")

    # Metadata table
    meta_tbl = doc.add_table(rows=2, cols=2)
    meta_tbl.rows[0].cells[0].text = "Instructor:"
    meta_tbl.rows[0].cells[1].text = ""
    meta_tbl.rows[1].cells[0].text = "Course & Section:"
    meta_tbl.rows[1].cells[1].text = ""

    doc.add_paragraph("")

    # Roster table
    r_tbl = doc.add_table(rows=2, cols=3)
    r_tbl.rows[0].cells[0].text = "Item"
    r_tbl.rows[0].cells[1].text = "Name of Student"
    r_tbl.rows[0].cells[2].text = "Student ID"
    r_tbl.rows[1].cells[0].text = "1"
    r_tbl.rows[1].cells[1].text = "Student A"
    r_tbl.rows[1].cells[2].text = "2026-9999"

    doc.add_paragraph("")

    # Satellite Campus 2-tier Signatory Table: Recommending Approval (Chair) -> Approved (Campus Administrator)
    sig_tbl = doc.add_table(rows=3, cols=2)
    sig_tbl.rows[0].cells[0].text = "Recommending Approval:"
    sig_tbl.rows[0].cells[1].text = "Approved by:"
    sig_tbl.rows[1].cells[0].text = ""
    sig_tbl.rows[1].cells[1].text = ""
    sig_tbl.rows[2].cells[0].text = "Department Chairperson"
    sig_tbl.rows[2].cells[1].text = "Campus Administrator"

    doc.save(file_path)


def create_graduate_school_docx(file_path: str) -> None:
    """Builds a DOCX with Graduate School (GS-OLC) header and Program Coordinator / GS Dean signatories."""
    doc = docx.Document()
    doc.add_paragraph("Republic of the Philippines")
    doc.add_paragraph("CAVITE STATE UNIVERSITY")
    doc.add_paragraph("Main Campus")
    doc.add_paragraph("GRADUATE SCHOOL AND OPEN LEARNING COLLEGE")
    doc.add_paragraph("Department of Open Learning")
    doc.add_paragraph("COMPREHENSIVE EXAMINATION CONSULTATION FORM")

    meta_tbl = doc.add_table(rows=2, cols=2)
    meta_tbl.rows[0].cells[0].text = "Faculty / Adviser:"
    meta_tbl.rows[0].cells[1].text = ""
    meta_tbl.rows[1].cells[0].text = "Program & Section:"
    meta_tbl.rows[1].cells[1].text = ""

    doc.add_paragraph("")

    # Signatory block
    sig_tbl = doc.add_table(rows=3, cols=2)
    sig_tbl.rows[0].cells[0].text = "Noted by:"
    sig_tbl.rows[0].cells[1].text = "Approved by:"
    sig_tbl.rows[1].cells[0].text = ""
    sig_tbl.rows[1].cells[1].text = ""
    sig_tbl.rows[2].cells[0].text = "Program Coordinator"
    sig_tbl.rows[2].cells[1].text = "Dean, GS-OLC"

    doc.save(file_path)


def create_satellite_xlsx_grading(file_path: str) -> None:
    """Builds an Excel grading sheet with Rosario Campus header and N-tier footer."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Grading Sheet"

    # Header rows
    ws["A1"] = "Republic of the Philippines"
    ws["A2"] = "CAVITE STATE UNIVERSITY"
    ws["A3"] = "CvSU - Rosario Campus"
    ws["A4"] = "Department of Engineering"

    # Metadata
    ws["A6"] = "Instructor:"
    ws["B6"] = ""
    ws["C6"] = "Course & Sec:"
    ws["D6"] = ""

    # Roster header
    ws["A8"] = "No."
    ws["B8"] = "Student Name"
    ws["C8"] = "Student Number"
    ws["A9"] = "1"
    ws["B9"] = "Cruz, Juan"
    ws["C9"] = "2026-1001"

    # Signatory block in rows 15-17
    ws["A15"] = "Prepared by:"
    ws["C15"] = "Recommending Approval:"
    ws["E15"] = "Approved by:"

    ws["A16"] = ""
    ws["C16"] = ""
    ws["E16"] = ""

    ws["A17"] = "Instructor"
    ws["C17"] = "Department Chairperson"
    ws["E17"] = "Campus Administrator"

    wb.save(file_path)


# ─────────────────────────────────────────────────────────────────────────────
# 1. Deterministic OrganizationalContext Extraction Tests
# ─────────────────────────────────────────────────────────────────────────────

def test_main_campus_ceit_header_discovery(tmp_path):
    """Verify deterministic discovery of Main Campus, CEIT, and DIT from banner paragraphs."""
    docx_path = str(tmp_path / "main_ceit.docx")
    create_main_campus_ceit_docx(docx_path)

    analyzer = DeterministicTemplateAnalyzer.get_instance()
    manifest = analyzer.analyze(docx_path, force_refresh=True)

    org = manifest.organizational_context
    assert org is not None, "OrganizationalContext must be detected"
    assert org.institution == "CAVITE STATE UNIVERSITY"
    assert "Don Severino" in org.campus or "Main" in org.campus
    assert org.college == "COLLEGE OF ENGINEERING AND INFORMATION TECHNOLOGY"
    assert org.department == "Department of Information Technology"


def test_satellite_campus_tanza_header_discovery(tmp_path):
    """Verify discovery of satellite campus (Tanza Campus) and department without hardcoded college."""
    docx_path = str(tmp_path / "tanza_campus.docx")
    create_satellite_campus_tanza_docx(docx_path)

    analyzer = DeterministicTemplateAnalyzer.get_instance()
    manifest = analyzer.analyze(docx_path, force_refresh=True)

    org = manifest.organizational_context
    assert org is not None
    assert org.institution == "CAVITE STATE UNIVERSITY"
    assert "Tanza Campus" in org.campus
    assert org.department == "Department of Information Technology"
    # College should be None or empty for satellite campuses without dedicated college divisions
    assert not org.college


def test_graduate_school_header_discovery(tmp_path):
    """Verify discovery of Graduate School and Open Learning College (GS-OLC)."""
    docx_path = str(tmp_path / "gs_olc.docx")
    create_graduate_school_docx(docx_path)

    analyzer = DeterministicTemplateAnalyzer.get_instance()
    manifest = analyzer.analyze(docx_path, force_refresh=True)

    org = manifest.organizational_context
    assert org is not None
    assert org.institution == "CAVITE STATE UNIVERSITY"
    assert "Main" in org.campus
    assert "GRADUATE SCHOOL" in org.college
    assert "Open Learning" in org.department


def test_excel_grading_sheet_header_discovery(tmp_path):
    """Verify header banner extraction on .xlsx spreadsheets (Rosario Campus)."""
    xlsx_path = str(tmp_path / "rosario_grading.xlsx")
    create_satellite_xlsx_grading(xlsx_path)

    analyzer = DeterministicTemplateAnalyzer.get_instance()
    manifest = analyzer.analyze(xlsx_path, force_refresh=True)

    org = manifest.organizational_context
    assert org is not None
    assert org.institution == "CAVITE STATE UNIVERSITY"
    assert "Rosario Campus" in org.campus
    assert "Department of Engineering" in org.department


# ─────────────────────────────────────────────────────────────────────────────
# 2. Dynamic SignatorySlot Discovery Tests
# ─────────────────────────────────────────────────────────────────────────────

def test_signatory_slot_discovery_main_campus_ceit(tmp_path):
    """Verify discovery of 3-tier signatory chain in Main Campus DOCX table."""
    docx_path = str(tmp_path / "main_ceit_sigs.docx")
    create_main_campus_ceit_docx(docx_path)

    analyzer = DeterministicTemplateAnalyzer.get_instance()
    manifest = analyzer.analyze(docx_path, force_refresh=True)

    slots = manifest.signatory_slots
    assert len(slots) == 3, f"Expected 3 signatory slots, found {len(slots)}"

    # Slot 0: Prepared by (Instructor)
    assert "Prepared by" in slots[0].role_label
    assert slots[0].default_title == "Instructor"
    assert slots[0].bound_key in ("prepared_by", "instructor")

    # Slot 1: Reviewed by (Department Chairperson)
    assert "Reviewed by" in slots[1].role_label
    assert slots[1].default_title == "Department Chairperson"
    assert slots[1].bound_key == "department_chair"

    # Slot 2: Approved by (College Dean)
    assert "Approved by" in slots[2].role_label
    assert slots[2].default_title == "College Dean"
    assert slots[2].bound_key == "college_dean"


def test_signatory_slot_discovery_satellite_campus(tmp_path):
    """Verify discovery of satellite campus signatory chain (Chairperson -> Campus Administrator)."""
    docx_path = str(tmp_path / "satellite_sigs.docx")
    create_satellite_campus_tanza_docx(docx_path)

    analyzer = DeterministicTemplateAnalyzer.get_instance()
    manifest = analyzer.analyze(docx_path, force_refresh=True)

    slots = manifest.signatory_slots
    assert len(slots) == 2, f"Expected 2 signatory slots, found {len(slots)}"

    # Slot 0: Recommending Approval (Chair)
    assert "Recommending Approval" in slots[0].role_label
    assert slots[0].bound_key == "department_chair"

    # Slot 1: Approved by (Campus Administrator)
    assert "Approved by" in slots[1].role_label
    assert slots[1].default_title == "Campus Administrator"
    assert slots[1].bound_key in ("campus_administrator", "approved_by")


def test_signatory_slot_discovery_excel(tmp_path):
    """Verify discovery of signatory slots in Excel (.xlsx) sheets."""
    xlsx_path = str(tmp_path / "excel_sigs.xlsx")
    create_satellite_xlsx_grading(xlsx_path)

    analyzer = DeterministicTemplateAnalyzer.get_instance()
    manifest = analyzer.analyze(xlsx_path, force_refresh=True)

    slots = manifest.signatory_slots
    assert len(slots) == 3, f"Expected 3 signatory slots in Excel, got {len(slots)}"
    labels = [s.role_label for s in slots]
    assert any("Prepared by" in lbl for lbl in labels)
    assert any("Recommending Approval" in lbl for lbl in labels)
    assert any("Approved by" in lbl for lbl in labels)


# ─────────────────────────────────────────────────────────────────────────────
# 3. Dynamic Signatory Rendering Tests
# ─────────────────────────────────────────────────────────────────────────────

def test_dynamic_signatory_rendering_main_campus(tmp_path):
    """Verify DynamicDocumentRenderer populates Main Campus signatories (Instructor, Chair, Dean)."""
    tmpl_path = str(tmp_path / "main_tmpl.docx")
    out_path = str(tmp_path / "main_out.docx")
    create_main_campus_ceit_docx(tmpl_path)

    info = ClassInfo(
        instructor="DAN JOSEPH A. ORTEGA",
        course_section="BSCS 3-1",
        schedule_code="20261001",
        subject="ITEC 50",
        students=[("Student A", "2026-0001")]
    )

    context = {
        "instructor": "DAN JOSEPH A. ORTEGA",
        "course_section": "BSCS 3-1",
        "schedule_code": "20261001",
        "subject": "ITEC 50",
        "students": [("Student A", "2026-0001")],
        "department_chair": "ENGR. EMIL C. ALVAREZ",
        "college_dean": "DR. MARLON R. PERA",
    }

    renderer = DynamicDocumentRenderer()
    renderer.render_docx(tmpl_path, context, output_path=out_path)

    # Inspect rendered docx
    res_doc = docx.Document(out_path)
    full_text = " ".join([p.text for p in res_doc.paragraphs] + [c.text for t in res_doc.tables for row in t.rows for c in row.cells])

    assert "DAN JOSEPH A. ORTEGA" in full_text
    assert "ENGR. EMIL C. ALVAREZ" in full_text
    assert "DR. MARLON R. PERA" in full_text


def test_dynamic_signatory_rendering_satellite_campus(tmp_path):
    """Verify DynamicDocumentRenderer populates Satellite Campus signatories (Chair, Campus Administrator)."""
    tmpl_path = str(tmp_path / "satellite_tmpl.docx")
    out_path = str(tmp_path / "satellite_out.docx")
    create_satellite_campus_tanza_docx(tmpl_path)

    context = {
        "instructor": "PROF. ALICE TAN",
        "course_section": "BSIT 2-1",
        "schedule_code": "20262002",
        "subject": "DCIT 21",
        "students": [("Student B", "2026-0002")],
        "department_chair": "PROF. BOBBY RAY",
        "campus_administrator": "DR. NOEL A. SEDIGO",
    }

    renderer = DynamicDocumentRenderer()
    renderer.render_docx(tmpl_path, context, output_path=out_path)

    res_doc = docx.Document(out_path)
    full_text = " ".join([p.text for p in res_doc.paragraphs] + [c.text for t in res_doc.tables for row in t.rows for c in row.cells])

    assert "PROF. ALICE TAN" in full_text
    assert "PROF. BOBBY RAY" in full_text
    assert "DR. NOEL A. SEDIGO" in full_text


# ─────────────────────────────────────────────────────────────────────────────
# 4. Shareable Department Presets (.cvsu-org.json) & Bridge API Tests
# ─────────────────────────────────────────────────────────────────────────────

class DummyWindow:
    def __init__(self, dialog_return=None):
        self.dialog_return = dialog_return

    def create_file_dialog(self, *args, **kwargs):
        return self.dialog_return


class DummyBridge(ConfigMixin, TemplateMixin):
    def __init__(self, window=None):
        self._window = window
        self.rosters = []
        self.schedule_path = None
        self.roster_configs = []


def test_cvsu_org_preset_export_and_import(tmp_path):
    """Verify export and import of .cvsu-org.json organizational presets."""
    preset_file = str(tmp_path / "tanza_it.cvsu-org.json")

    preset_data = {
        "version": "1.0",
        "campus": "CvSU - Tanza Campus",
        "college": "",
        "department": "Department of Information Technology",
        "signatories": [
            {
                "slot_id": "prepared_by",
                "label": "Prepared by:",
                "name": "DAN JOSEPH A. ORTEGA",
                "title": "Instructor"
            },
            {
                "slot_id": "recommending_approval",
                "label": "Recommending Approval:",
                "name": "PROF. BOBBY RAY",
                "title": "Department Chairperson"
            },
            {
                "slot_id": "approved_by",
                "label": "Approved by:",
                "name": "DR. NOEL A. SEDIGO",
                "title": "Campus Administrator"
            }
        ]
    }

    bridge = DummyBridge()

    # 1. Export preset
    export_res = bridge.export_org_preset(path=preset_file, preset_data=preset_data)
    assert export_res["status"] == "success"
    assert os.path.exists(preset_file)

    # 2. Import preset
    import_res = bridge.import_org_preset(path=preset_file)
    assert import_res["status"] == "success"
    loaded = import_res["preset"]

    assert loaded["campus"] == "CvSU - Tanza Campus"
    assert loaded["department"] == "Department of Information Technology"
    assert len(loaded["signatories"]) == 3
    assert loaded["signatories"][2]["title"] == "Campus Administrator"
    assert loaded["signatories"][2]["name"] == "DR. NOEL A. SEDIGO"


def test_preset_strict_path_validation(tmp_path):
    """Verify security controls and path validation on preset import/export."""
    bridge = DummyBridge()

    # Reject null byte
    res = bridge.export_org_preset(path="invalid\x00path.json", preset_data={})
    assert res["status"] == "error"

    # Reject non-json extension
    res = bridge.export_org_preset(path=str(tmp_path / "preset.exe"), preset_data={})
    assert res["status"] == "error"

    # Reject non-existent file on import
    res = bridge.import_org_preset(path=str(tmp_path / "non_existent.json"))
    assert res["status"] == "error"

    # Reject malformed json on import
    bad_json = str(tmp_path / "corrupt.json")
    with open(bad_json, "w") as f:
        f.write("{invalid_json:")
    res = bridge.import_org_preset(path=bad_json)
    assert res["status"] == "error"


def test_manifest_summary_includes_org_context_and_signatories(tmp_path):
    """Verify inspect_custom_template manifest_summary provides organizational_context and signatory_slots."""
    docx_path = str(tmp_path / "test_summary.docx")
    create_main_campus_ceit_docx(docx_path)

    bridge = DummyBridge()
    inspect_res = bridge.inspect_custom_template(docx_path)
    assert inspect_res["status"] == "success"

    summary = inspect_res["manifest_summary"]
    assert "organizational_context" in summary
    assert "signatory_slots" in summary

    org = summary["organizational_context"]
    assert org["institution"] == "CAVITE STATE UNIVERSITY"
    assert org["college"] == "COLLEGE OF ENGINEERING AND INFORMATION TECHNOLOGY"

    slots = summary["signatory_slots"]
    assert len(slots) == 3
    assert slots[0]["default_title"] == "Instructor"
    assert slots[1]["default_title"] == "Department Chairperson"
    assert slots[2]["default_title"] == "College Dean"
