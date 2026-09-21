#!/usr/bin/env python3
"""
modules/generators/grade_gen.py

Pure recipe-driven Excel Grading Sheet Generator.
Consumes ValidatedTemplateRecipe exclusively.
Zero hardcoded coordinates or positional fallback logic.
"""

import os
import re
import shutil
import tempfile
from typing import Optional, Union, Any, List, Tuple

import openpyxl

from modules.common.logger import logger
from modules.models.recipe import (
    TemplateError,
    ValidatedTemplateRecipe,
)


class GradeGenerator:
    """
    Pure recipe-driven Grade Sheet Generator for CvSU.
    Requires a ValidatedTemplateRecipe instance for construction.
    """

    def __init__(
        self,
        template_path: str,
        recipe: ValidatedTemplateRecipe,
    ):
        if not isinstance(recipe, ValidatedTemplateRecipe):
            raise TypeError("GradeGenerator requires a ValidatedTemplateRecipe instance")
        self.template_path = template_path
        self.recipe = recipe

    @staticmethod
    def eval_has_lab(info: dict) -> bool:
        """Business logic evaluation to determine whether class has laboratory."""
        subject_type = str(info.get("subject_type", "")).lower()
        template_type = str(info.get("template_type", "")).lower()
        time_str = str(info.get("time", "")).upper()
        subj_str = str(info.get("subject", "")).upper()
        has_lab_flag = info.get("has_lab")

        if subject_type:
            return "lab" in subject_type
        elif has_lab_flag is not None:
            return bool(has_lab_flag)
        elif "ll" in template_type or "lab" in template_type:
            return True
        elif "l" in template_type or "lec" in template_type:
            return False
        elif "LAB" in time_str or "LAB:" in time_str or "(LAB)" in subj_str or "LABORATORY" in subj_str:
            return True
        return False

    @classmethod
    def for_class(
        cls,
        templates_dir: str,
        info: dict,
        resolver: Optional[Any] = None,
    ) -> "GradeGenerator":
        """
        Factory method: selects template identity from business logic,
        resolves validated recipe via TemplateRecipeResolver,
        and constructs GradeGenerator with the validated recipe.
        """
        if resolver is None:
            from modules.services.template_recipe_service import TemplateRecipeResolver
            resolver = TemplateRecipeResolver.get_instance()
        has_lab = cls.eval_has_lab(info)
        template_filename = "GRADING_LECTURE_LAB_TEMPLATE.xlsx" if has_lab else "GRADING_LECTURE_TEMPLATE.xlsx"
        template_path = os.path.join(templates_dir, template_filename)
        recipe = resolver.resolve(template_path, "grade_sheet_xlsx")
        return cls(template_path, recipe)

    def _parse_semester_and_year(self, raw_sem: str) -> tuple:
        raw_sem = str(raw_sem).strip()
        sem_lower = raw_sem.lower()
        if "second" in sem_lower or "2nd" in sem_lower:
            sem_str = "2nd Semester"
        elif "first" in sem_lower or "1st" in sem_lower:
            sem_str = "1st Semester"
        elif "midyear" in sem_lower or "summer" in sem_lower:
            sem_str = "Midyear"
        else:
            sem_str = raw_sem

        year_match = re.search(r'(20\d{2})\s*[-–/]\s*(20\d{2})', raw_sem)
        if year_match:
            year_str = f"{year_match.group(1)}-{year_match.group(2)}"
        else:
            years = re.findall(r'20\d{2}', raw_sem)
            if len(years) >= 2:
                year_str = f"{years[0]}-{years[1]}"
            elif len(years) == 1:
                y = int(years[0])
                year_str = f"{y}-{y+1}"
            else:
                raise ValueError(f"Could not parse year from '{raw_sem}'")

        return sem_str, year_str

    def _parse_subject(self, raw_subject: str) -> tuple:
        raw_subject = str(raw_subject).strip()
        parts = re.split(r'\s+[-–—―−]\s+', raw_subject, maxsplit=1)
        if len(parts) == 2:
            return parts[0].strip(), parts[1].strip()
        parts = re.split(r'[-–—―−]', raw_subject, maxsplit=1)
        if len(parts) == 2:
            return parts[0].strip(), parts[1].strip()
        return raw_subject, ""

    def _normalize_course_section(self, raw_course: str) -> str:
        raw = str(raw_course).strip()
        norm = re.sub(r'([a-zA-Z]+)(\d)', r'\1 \2', raw)
        return re.sub(r'\s+', ' ', norm).strip()

    def generate(self, info: dict, students: list, output_path: str) -> bool:
        """
        Executes grade sheet generation purely driven by self.recipe.
        Coordinates, roster bounds, and signatures are resolved from recipe.
        """
        recipe = self.recipe
        rb = recipe.roster_binding
        if rb is None:
            raise TemplateError("GradeGenerator recipe is missing mandatory roster_binding")

        start_row = rb.first_data_row_index
        if rb.capacity_limit is None or rb.capacity_limit <= 0:
            raise TemplateError("GradeGenerator recipe is missing mandatory capacity_limit")
        max_capacity = rb.capacity_limit

        sem_str, year_str = self._parse_semester_and_year(info.get("semester", ""))
        subj_code, subj_title = self._parse_subject(info.get("subject", ""))
        course_str = self._normalize_course_section(info.get("course", ""))
        raw_sched = str(info.get("sched", "")).strip()
        sched_val = int(raw_sched) if raw_sched.isdigit() else raw_sched
        instructor = str(info.get("instructor", "")).strip().upper()

        cleaned_students = []
        for s in students:
            if isinstance(s, (list, tuple)):
                s_name = str(s[0]).strip() if len(s) > 0 and s[0] is not None else ""
                s_num = str(s[1]).strip() if len(s) > 1 and s[1] is not None else ""
            elif isinstance(s, dict):
                s_name = str(s.get("student_name") or s.get("name") or "").strip()
                s_num = str(s.get("student_number") or s.get("number") or "").strip()
            else:
                s_name = str(s).strip()
                s_num = ""
            if s_name:
                cleaned_students.append((s_name, s_num))

        dir_name = os.path.dirname(os.path.abspath(output_path))
        os.makedirs(dir_name, exist_ok=True)
        with tempfile.NamedTemporaryFile(dir=dir_name, delete=False, suffix=".tmp.xlsx") as fh:
            tmp_path = fh.name
        shutil.copy2(self.template_path, tmp_path)

        def sanitize_excel(val):
            if isinstance(val, str) and val.startswith(('=', '+', '-', '@')):
                return "'" + val
            return val

        try:
            wb = openpyxl.load_workbook(tmp_path, data_only=False)
            sheet_names = [s.lower() for s in wb.sheetnames]
            name_map = {s.lower(): s for s in wb.sheetnames}

            if not rb.worksheet_name:
                raise TemplateError("GradeGenerator recipe is missing mandatory worksheet_name")
            roster_sheet = rb.worksheet_name
            if roster_sheet.lower() not in name_map:
                raise TemplateError(
                    f"Grade template is missing roster worksheet '{roster_sheet}'."
                )
            ws = wb[name_map[roster_sheet.lower()]]

            # 1. Metadata Bindings
            field_values = {
                "schedule_code": sched_val if isinstance(sched_val, int) else sanitize_excel(sched_val),
                "course_section": sanitize_excel(course_str),
                "subject_code": sanitize_excel(subj_code),
                "semester": sanitize_excel(sem_str),
                "subject_title": sanitize_excel(subj_title),
                "school_year": sanitize_excel(year_str),
                "instructor": sanitize_excel(instructor),
            }
            if info.get("units"):
                try:
                    field_values["units"] = int(info["units"])
                except Exception:
                    pass

            for field_name, val in field_values.items():
                target = recipe.get_header_target(field_name)
                if target:
                    if "!" in target:
                        t_sheet, t_cell = target.split("!", 1)
                        if t_sheet.lower() in name_map:
                            wb[name_map[t_sheet.lower()]][t_cell] = val
                    else:
                        ws[target] = val

            # 2. Scope-Aware Signature Bindings
            for role, sig_binding in recipe.signature_bindings.items():
                target = sig_binding.target
                if not target:
                    continue

                if "!" in target:
                    t_sheet, t_cell = target.split("!", 1)
                    if t_sheet.lower() in name_map:
                        wb[name_map[t_sheet.lower()]][t_cell] = sanitize_excel(instructor)
                else:
                    if role in ("instructor:laboratory", "laboratory") and "laboratory" in sheet_names:
                        wb[name_map["laboratory"]][target] = sanitize_excel(instructor)
                    elif role in ("instructor:consolidated", "consolidated") and "consolidated" in sheet_names:
                        wb[name_map["consolidated"]][target] = sanitize_excel(instructor)
                    elif role in ("instructor", "instructor_signature"):
                        ws[target] = sanitize_excel(instructor)

            # 3. Student Roster Population with Mandatory Capacity Limit
            if len(cleaned_students) > max_capacity:
                logger.warning(
                    f"Roster for {course_str} ({sched_val}) has {len(cleaned_students)} students, "
                    f"exceeding maximum section capacity of {max_capacity}. Clamping to first {max_capacity} students."
                )
                cleaned_students = cleaned_students[:max_capacity]

            total_slots = max_capacity
            name_col = rb.name_col
            id_col = rb.id_col
            if rb.index_col is None:
                raise TemplateError("GradeGenerator recipe is missing mandatory index_col")
            index_col = rb.index_col

            for r_idx in range(total_slots):
                row_num = start_row + r_idx
                if r_idx < len(cleaned_students):
                    name, num = cleaned_students[r_idx]
                    ws.cell(row=row_num, column=index_col).value = r_idx + 1
                    ws.cell(row=row_num, column=name_col).value = sanitize_excel(name)
                    ws.cell(row=row_num, column=id_col).value = int(num) if (num.isdigit() and not num.startswith('0')) else num
                else:
                    ws.cell(row=row_num, column=name_col).value = None
                    ws.cell(row=row_num, column=id_col).value = None

            # 4. Institutional College Banner (Grading Sheet)
            college_target = recipe.get_header_target("college")
            college_val = str(info.get("college") or "").strip() or "COLLEGE OF ENGINEERING AND INFORMATION TECHNOLOGY"
            if college_target:
                if "!" in college_target:
                    t_sheet, t_cell = college_target.split("!", 1)
                    if t_sheet.lower() in name_map:
                        wb[name_map[t_sheet.lower()]][t_cell] = sanitize_excel(college_val)
                elif "grading sheet" in sheet_names:
                    wb[name_map["grading sheet"]][college_target] = sanitize_excel(college_val)

            wb.save(tmp_path)
            wb.close()
            os.replace(tmp_path, output_path)
            logger.info(f"Grades generated at: {output_path}")
            return True
        except Exception as e:
            if os.path.exists(tmp_path):
                try:
                    os.remove(tmp_path)
                except Exception:
                    pass
            raise e
