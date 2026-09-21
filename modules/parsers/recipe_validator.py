#!/usr/bin/env python3
"""
modules/parsers/recipe_validator.py

Authoritative Recipe Validator for Dynamic Template Discovery.
Sole authority permitted to construct ValidatedTemplateRecipe instances.
Enforces schema version 2, profile constraints, structural rules, and 5-point safety verification.
"""

import os
from typing import Dict, List, Optional, Any, Union, Tuple

import openpyxl
from openpyxl.utils import coordinate_to_tuple

from modules.models.recipe import (
    RECIPE_SCHEMA_VERSION,
    _PRIVATE_CONSTRUCTION_SENTINEL,
    TemplateError,
    AmbiguousTemplateError,
    InvalidRecipeError,
    RosterBinding,
    HeaderCellBinding,
    SignatureBinding,
    GeneratorProfile,
    PROFILE_REGISTRY,
    PROFILE_ACADEMIC_DOCX,
    PROFILE_GRADE_SHEET_XLSX,
    PROFILE_ATTENDANCE_DOCX,
    RawTemplateRecipeCandidate,
    RawAttendanceTemplateRecipeCandidate,
    ValidatedRecipeBase,
    ValidatedTemplateRecipe,
    ValidatedAttendanceTemplateRecipe,
    AttendanceInfoBinding,
    AttendanceMatrixBinding,
)
from modules.parsers.semantic_registry import (
    ROLE_INSTRUCTOR_SIGNATURE,
    FIELD_INSTRUCTOR,
)
from modules.common.path_utils import validate_output_folder
from modules.common.docx_utils import load_docx, w


class RecipeValidator:
    """
    Sole authority that transforms RawTemplateRecipeCandidate or schema v2 dict
    into ValidatedTemplateRecipe.
    """

    @classmethod
    def validate(
        cls,
        candidate: Union[RawTemplateRecipeCandidate, RawAttendanceTemplateRecipeCandidate],
        profile: Union[GeneratorProfile, str],
    ) -> ValidatedRecipeBase:
        """
        Validates raw candidate observations against profile requirements and structural invariants.
        Returns authoritative ValidatedRecipeBase (ValidatedTemplateRecipe or ValidatedAttendanceTemplateRecipe).
        """
        if not candidate:
            raise InvalidRecipeError("Cannot validate empty candidate")

        if isinstance(profile, str):
            profile = PROFILE_REGISTRY.get(profile, PROFILE_ACADEMIC_DOCX)

        if isinstance(candidate, RawAttendanceTemplateRecipeCandidate):
            if profile.profile_id not in ("attendance_docx", "attendance"):
                raise TemplateError(
                    f"Template '{os.path.basename(candidate.template_path)}' is an attendance candidate "
                    f"and cannot be validated against profile '{profile.profile_id}'."
                )
            return cls._validate_attendance(candidate, profile)

        # Non-attendance candidate with attendance profile -> E18 fail closed
        if profile.profile_id in ("attendance_docx", "attendance"):
            raise TemplateError(
                f"Template '{os.path.basename(candidate.template_path)}' fails profile '{profile.profile_id}': "
                f"Candidate is {type(candidate).__name__} without required attendance matrix structure (E18)."
            )

        # 1. Resolve and check collisions / ambiguities
        cls._resolve_collisions(candidate, profile)

        # 2. Validate Roster Binding
        roster_binding = cls._validate_roster(candidate.roster_candidate, profile)

        # 3. Validate Header Bindings
        header_bindings = cls._validate_headers(candidate.header_candidates, profile, candidate.template_path)

        # 4. Validate Signatures
        signature_bindings = cls._validate_signatures(candidate.signature_candidates, profile, candidate.template_path)

        # 5. Check Required Fields
        for req_field in profile.required_fields:
            if req_field not in header_bindings and req_field not in candidate.metadata:
                raise TemplateError(
                    f"Template '{os.path.basename(candidate.template_path or 'template')}' fails profile '{profile.profile_id}': "
                    f"Missing required field '{req_field}'."
                )

        # 6. Check Prohibited Fields
        for pro_field in profile.prohibited_fields:
            if pro_field in header_bindings:
                raise TemplateError(
                    f"Template '{os.path.basename(candidate.template_path or 'template')}' violates profile '{profile.profile_id}': "
                    f"Contains prohibited field '{pro_field}'."
                )

        # 7. Fail-Closed Physical & Serialized Geometry Validation (DOCX / XLSX)
        is_xlsx = (
            profile.document_family == "grade_sheet_xlsx"
            or profile.profile_id == "grade_sheet_xlsx"
            or (bool(candidate.template_path) and candidate.template_path.lower().endswith((".xlsx", ".xls")))
        )
        if is_xlsx:
            cls._validate_xlsx_geometry(candidate, roster_binding, header_bindings, signature_bindings)
        else:
            cls._validate_docx_geometry(candidate, roster_binding, header_bindings, signature_bindings)

        # 8. Ensure output_folder is validated in metadata
        metadata = dict(candidate.metadata)
        output_folder = metadata.get("output_folder", "CEIT_Forms")
        metadata["output_folder"] = validate_output_folder(output_folder, default="CEIT_Forms")

        # 9. Construct Authoritative ValidatedTemplateRecipe
        return ValidatedTemplateRecipe(
            schema_version=RECIPE_SCHEMA_VERSION,
            profile_id=profile.profile_id,
            fingerprint=candidate.fingerprint,
            template_path=candidate.template_path,
            roster_binding=roster_binding,
            header_bindings=header_bindings,
            signature_bindings=signature_bindings,
            metadata=metadata,
            verified_safe=True,
            _construction_token=_PRIVATE_CONSTRUCTION_SENTINEL,
        )

    @staticmethod
    def _validate_info_bindings(
        bindings: Dict[str, Any],
        row_cell_counts: Tuple[int, ...],
        context: str,
    ) -> None:
        """Validates that every info table coordinate is within the row and cell counts."""
        if not isinstance(bindings, dict) or not bindings:
            raise TemplateError(f"{context}: information table bindings are empty or invalid.")
        for field, target in bindings.items():
            if not isinstance(target, (list, tuple)) or len(target) != 2:
                raise InvalidRecipeError(
                    f"{context}: binding for field '{field}' must be a (row, col) coordinate pair, got {target}."
                )
            r, c = target[0], target[1]
            if not isinstance(r, int) or r < 0:
                raise InvalidRecipeError(
                    f"{context}: invalid row index {r} for field '{field}' (must be non-negative integer)."
                )
            if not isinstance(c, int) or c < 0:
                raise InvalidRecipeError(
                    f"{context}: invalid col index {c} for field '{field}' (must be non-negative integer)."
                )
            if r >= len(row_cell_counts):
                raise InvalidRecipeError(
                    f"{context}: row index {r} for field '{field}' out of range (info table has {len(row_cell_counts)} rows)."
                )
            if c >= row_cell_counts[r]:
                raise InvalidRecipeError(
                    f"{context}: column index {c} for field '{field}' out of range (row {r} has {row_cell_counts[r]} cells)."
                )

    @classmethod
    def _validate_attendance(
        cls,
        candidate: RawAttendanceTemplateRecipeCandidate,
        profile: GeneratorProfile,
    ) -> ValidatedAttendanceTemplateRecipe:
        if candidate.collisions:
            raise AmbiguousTemplateError(
                f"Template '{os.path.basename(candidate.template_path)}' has conflicting structural candidates: "
                f"{candidate.collisions}"
            )

        if not candidate.info_candidate:
            raise TemplateError(
                f"Template '{os.path.basename(candidate.template_path)}' is missing required information table (E13)."
            )

        if not candidate.matrix_candidate:
            raise TemplateError(
                f"Template '{os.path.basename(candidate.template_path)}' is missing required attendance matrix table (E11)."
            )

        info_cand = candidate.info_candidate
        tbl_idx = info_cand.get("table_index")
        bindings = info_cand.get("bindings", {})
        if tbl_idx is None or not isinstance(tbl_idx, int) or tbl_idx < 0 or not bindings:
            raise TemplateError(
                f"Template '{os.path.basename(candidate.template_path)}' information table bindings are invalid."
            )

        matrix_cand = candidate.matrix_candidate
        m_tbl_idx = matrix_cand.get("table_index")
        if m_tbl_idx is None or not isinstance(m_tbl_idx, int) or m_tbl_idx < 0:
            raise TemplateError(
                f"Template '{os.path.basename(candidate.template_path)}' matrix table index is missing or invalid."
            )

        # 1. Info and matrix tables must remain distinct
        if tbl_idx == m_tbl_idx:
            raise AmbiguousTemplateError(
                f"Table identity collision: info table and matrix table target identical index {tbl_idx}."
            )

        # 2. Check student columns (no, name, id)
        no_col = matrix_cand.get("no_col", 0)
        name_col = matrix_cand.get("name_col")
        id_col = matrix_cand.get("id_col")
        if not isinstance(no_col, int) or no_col < 0:
            raise InvalidRecipeError(f"Template '{os.path.basename(candidate.template_path)}' invalid no_col.")
        if name_col is None or not isinstance(name_col, int) or name_col < 0:
            raise TemplateError(
                f"Template '{os.path.basename(candidate.template_path)}' missing student name or student number column (E12)."
            )
        if id_col is None or not isinstance(id_col, int) or id_col < 0:
            raise TemplateError(
                f"Template '{os.path.basename(candidate.template_path)}' missing student name or student number column (E12)."
            )

        # 3. Check row indices
        h0_idx = matrix_cand.get("header_row0_index", 0)
        h1_idx = matrix_cand.get("header_row1_index", 1)
        student_row_idx = matrix_cand.get("student_template_row_index")
        if not isinstance(h0_idx, int) or h0_idx < 0 or not isinstance(h1_idx, int) or h1_idx < 0:
            raise InvalidRecipeError(f"Template '{os.path.basename(candidate.template_path)}' invalid header row indices.")
        if student_row_idx is None or not isinstance(student_row_idx, int) or student_row_idx < 0:
            raise InvalidRecipeError(
                f"Template '{os.path.basename(candidate.template_path)}' has invalid student_template_row_index / invalid student template row (E16)."
            )

        # 4. Check date columns geometry
        date_start = matrix_cand.get("date_columns_start")
        if date_start is None or not isinstance(date_start, int) or date_start < 0:
            raise InvalidRecipeError(
                f"Template '{os.path.basename(candidate.template_path)}' has invalid date_columns_start / invalid date columns start (E16)."
            )

        session_capacity = matrix_cand.get("template_session_capacity", 0)
        if not isinstance(session_capacity, int) or session_capacity <= 0:
            raise InvalidRecipeError(
                f"Template '{os.path.basename(candidate.template_path)}' template_session_capacity must be an integer > 0 / invalid session capacity (E16)."
            )

        # 5. Check summary columns count and names
        summary_columns_count = matrix_cand.get("summary_columns_count", 3)
        if not isinstance(summary_columns_count, int) or summary_columns_count <= 0:
            raise InvalidRecipeError("Attendance summary_columns_count must be an integer > 0.")

        summary_names = tuple(matrix_cand.get("summary_column_names", ("lb", "lc", "r")))
        if len(summary_names) != summary_columns_count:
            raise InvalidRecipeError(
                f"Summary column names count ({len(summary_names)}) does not match summary_columns_count ({summary_columns_count})."
            )

        total_cols = date_start + session_capacity + summary_columns_count
        default_summary_indices = tuple(range(date_start + session_capacity, total_cols))

        summary_column_indices = tuple(matrix_cand.get("summary_column_indices", default_summary_indices))
        summary_header1_cell_cols = tuple(matrix_cand.get("summary_header1_cell_cols", default_summary_indices))
        student_summary_cell_cols = tuple(matrix_cand.get("student_summary_cell_cols", default_summary_indices))

        default_sum_w_map = {"lb": 212, "lc": 208, "r": 133}
        default_widths = tuple(default_sum_w_map.get(str(sn).lower(), 200) for sn in summary_names)
        raw_widths = matrix_cand.get("summary_column_widths", default_widths)
        if not isinstance(raw_widths, (list, tuple)):
            raise InvalidRecipeError("Attendance summary_column_widths must be a list or tuple of integers.")
        summary_widths = tuple(raw_widths)

        # 6. Validate summary arrays lengths match
        if len(summary_column_indices) != summary_columns_count:
            raise InvalidRecipeError(
                f"Summary column indices length ({len(summary_column_indices)}) does not match summary_columns_count ({summary_columns_count})."
            )
        if len(summary_header1_cell_cols) != summary_columns_count:
            raise InvalidRecipeError(
                f"Summary header1 cell cols length ({len(summary_header1_cell_cols)}) does not match summary_columns_count ({summary_columns_count})."
            )
        if len(student_summary_cell_cols) != summary_columns_count:
            raise InvalidRecipeError(
                f"Student summary cell cols length ({len(student_summary_cell_cols)}) does not match summary_columns_count ({summary_columns_count})."
            )
        if len(summary_widths) != summary_columns_count:
            raise InvalidRecipeError(
                f"Attendance summary_column_widths length ({len(summary_widths)}) must match summary_column_names ({summary_columns_count})."
            )

        # 7. Validate non-negativity and positive widths
        for wval in summary_widths:
            if not isinstance(wval, int) or wval <= 0:
                raise InvalidRecipeError(
                    f"Attendance summary_column_widths contains invalid width: {wval}"
                )

        for c in summary_column_indices:
            if not isinstance(c, int) or c < 0:
                raise InvalidRecipeError(f"Invalid non-negative summary_column_indices column: {c}")
        for c in summary_header1_cell_cols:
            if not isinstance(c, int) or c < 0:
                raise InvalidRecipeError(f"Invalid non-negative summary_header1_cell_cols column: {c}")
        for c in student_summary_cell_cols:
            if not isinstance(c, int) or c < 0:
                raise InvalidRecipeError(f"Invalid non-negative student_summary_cell_cols column: {c}")

        # 8. Validate summary columns do not contain duplicates
        if len(set(summary_column_indices)) != len(summary_column_indices):
            raise InvalidRecipeError("Summary column indices contain duplicate columns.")
        if len(set(summary_header1_cell_cols)) != len(summary_header1_cell_cols):
            raise InvalidRecipeError("Summary header1 cell columns contain duplicate columns.")
        if len(set(student_summary_cell_cols)) != len(student_summary_cell_cols):
            raise InvalidRecipeError("Student summary cell columns contain duplicate columns.")

        # 9. Validate single structural coordinate cells
        week_template_cell_col = matrix_cand.get("week_template_cell_col", date_start)
        if not isinstance(week_template_cell_col, int) or week_template_cell_col < 0:
            raise InvalidRecipeError(f"Invalid week_template_cell_col: {week_template_cell_col}")

        default_h0_sum = default_summary_indices[0] if default_summary_indices else date_start + session_capacity
        summary_header0_cell_col = matrix_cand.get("summary_header0_cell_col", default_h0_sum)
        if not isinstance(summary_header0_cell_col, int) or summary_header0_cell_col < 0:
            raise InvalidRecipeError(f"Invalid summary_header0_cell_col: {summary_header0_cell_col}")

        date_template_cell_col = matrix_cand.get("date_template_cell_col", date_start)
        if not isinstance(date_template_cell_col, int) or date_template_cell_col < 0:
            raise InvalidRecipeError(f"Invalid date_template_cell_col: {date_template_cell_col}")

        student_date_template_cell_col = matrix_cand.get("student_date_template_cell_col", date_start)
        if not isinstance(student_date_template_cell_col, int) or student_date_template_cell_col < 0:
            raise InvalidRecipeError(f"Invalid student_date_template_cell_col: {student_date_template_cell_col}")

        # 10. Date and summary regions do not overlap
        date_cols_set = set(range(date_start, date_start + session_capacity))
        lead_cols_set = {no_col, name_col, id_col}

        if any(c in date_cols_set for c in lead_cols_set):
            raise InvalidRecipeError("Date columns overlap with lead student info columns.")
        if any(c in date_cols_set for c in summary_column_indices):
            raise InvalidRecipeError("Summary columns overlap with date columns region.")
        if any(c in date_cols_set for c in summary_header1_cell_cols):
            raise InvalidRecipeError("Summary header1 cell columns overlap with date columns region.")
        if any(c in date_cols_set for c in student_summary_cell_cols):
            raise InvalidRecipeError("Student summary cell columns overlap with date columns region.")
        if any(c in lead_cols_set for c in summary_column_indices):
            raise InvalidRecipeError("Summary columns overlap with lead columns.")
        if any(c in lead_cols_set for c in summary_header1_cell_cols):
            raise InvalidRecipeError("Summary header1 cell columns overlap with lead columns.")
        if any(c in lead_cols_set for c in student_summary_cell_cols):
            raise InvalidRecipeError("Student summary cell columns overlap with lead columns.")

        # 11. Enforce indices within their respective prototype-row cell counts and matrix row count
        matrix_row_count = None
        row0_cell_count = None
        row1_cell_count = None
        student_row_cell_count = None
        info_row_cell_counts = None

        if candidate.template_path and os.path.exists(candidate.template_path):
            try:
                zin, root, body = load_docx(candidate.template_path)
                zin.close()
                tables = body.findall(w("tbl"))
                if tbl_idx < 0 or tbl_idx >= len(tables):
                    raise InvalidRecipeError(
                        f"Info table index {tbl_idx} out of range for template '{os.path.basename(candidate.template_path)}' (found {len(tables)} tables)."
                    )
                if m_tbl_idx < 0 or m_tbl_idx >= len(tables):
                    raise InvalidRecipeError(
                        f"Matrix table index {m_tbl_idx} out of range for template '{os.path.basename(candidate.template_path)}' (found {len(tables)} tables)."
                    )

                info_rows = tables[tbl_idx].findall(w("tr"))
                info_row_cell_counts = tuple(len(tr.findall(w("tc"))) for tr in info_rows)
                cls._validate_info_bindings(
                    bindings,
                    info_row_cell_counts,
                    f"Physical template '{os.path.basename(candidate.template_path)}'",
                )

                matrix_rows = tables[m_tbl_idx].findall(w("tr"))
                matrix_row_count = len(matrix_rows)
                if h0_idx >= matrix_row_count:
                    raise InvalidRecipeError(
                        f"header_row0_index ({h0_idx}) out of range for template table {m_tbl_idx} (found {matrix_row_count} rows)."
                    )
                if h1_idx >= matrix_row_count:
                    raise InvalidRecipeError(
                        f"header_row1_index ({h1_idx}) out of range for template table {m_tbl_idx} (found {matrix_row_count} rows)."
                    )
                if student_row_idx >= matrix_row_count:
                    raise InvalidRecipeError(
                        f"student_template_row_index ({student_row_idx}) out of range for template table {m_tbl_idx} (found {matrix_row_count} rows)."
                    )

                row0_cell_count = len(matrix_rows[h0_idx].findall(w("tc")))
                row1_cell_count = len(matrix_rows[h1_idx].findall(w("tc")))
                student_row_cell_count = len(matrix_rows[student_row_idx].findall(w("tc")))
            except (TemplateError, InvalidRecipeError):
                raise
            except Exception as e:
                raise TemplateError(
                    f"Physical template inspection failed for '{candidate.template_path}': {e}"
                ) from e
        else:
            # Physical template is unavailable: must have candidate-supplied row-count metadata
            matrix_row_count = matrix_cand.get("matrix_row_count")
            row0_cell_count = matrix_cand.get("row0_cell_count")
            row1_cell_count = matrix_cand.get("row1_cell_count")
            student_row_cell_count = matrix_cand.get("student_row_cell_count")
            raw_info_counts = info_cand.get("row_cell_counts")

            if (
                matrix_row_count is None or not isinstance(matrix_row_count, int) or matrix_row_count <= 0
                or row0_cell_count is None or not isinstance(row0_cell_count, int) or row0_cell_count <= 0
                or row1_cell_count is None or not isinstance(row1_cell_count, int) or row1_cell_count <= 0
                or student_row_cell_count is None or not isinstance(student_row_cell_count, int) or student_row_cell_count <= 0
            ):
                raise InvalidRecipeError(
                    f"Serialized attendance recipe for '{os.path.basename(candidate.template_path or 'template')}' "
                    f"is missing required structural row-count metadata (matrix_row_count, row0_cell_count, "
                    f"row1_cell_count, student_row_cell_count) when physical template is unavailable."
                )

            if (
                raw_info_counts is None
                or not isinstance(raw_info_counts, (list, tuple))
                or len(raw_info_counts) == 0
                or not all(isinstance(c, int) and c >= 0 for c in raw_info_counts)
            ):
                raise InvalidRecipeError(
                    f"Serialized attendance recipe for '{os.path.basename(candidate.template_path or 'template')}' "
                    f"is missing required 'row_cell_counts' in info_binding when physical template is unavailable."
                )
            info_row_cell_counts = tuple(raw_info_counts)

            if h0_idx >= matrix_row_count:
                raise InvalidRecipeError(
                    f"header_row0_index ({h0_idx}) out of range for matrix_row_count ({matrix_row_count})."
                )
            if h1_idx >= matrix_row_count:
                raise InvalidRecipeError(
                    f"header_row1_index ({h1_idx}) out of range for matrix_row_count ({matrix_row_count})."
                )
            if student_row_idx >= matrix_row_count:
                raise InvalidRecipeError(
                    f"student_template_row_index ({student_row_idx}) out of range for matrix_row_count ({matrix_row_count})."
                )

            cls._validate_info_bindings(
                bindings,
                info_row_cell_counts,
                f"Serialized recipe '{os.path.basename(candidate.template_path or 'template')}'",
            )

        if week_template_cell_col >= row0_cell_count:
            raise InvalidRecipeError(
                f"week_template_cell_col ({week_template_cell_col}) out of range for header row 0 cell count ({row0_cell_count})."
            )
        if summary_header0_cell_col >= row0_cell_count:
            raise InvalidRecipeError(
                f"summary_header0_cell_col ({summary_header0_cell_col}) out of range for header row 0 cell count ({row0_cell_count})."
            )

        if date_template_cell_col >= row1_cell_count:
            raise InvalidRecipeError(
                f"date_template_cell_col ({date_template_cell_col}) out of range for header row 1 cell count ({row1_cell_count})."
            )
        for c in summary_header1_cell_cols:
            if c >= row1_cell_count:
                raise InvalidRecipeError(
                    f"summary_header1_cell_cols index ({c}) out of range for header row 1 cell count ({row1_cell_count})."
                )
        for c in summary_column_indices:
            if c >= row1_cell_count:
                raise InvalidRecipeError(
                    f"summary_column_indices index ({c}) out of range for header row 1 cell count ({row1_cell_count})."
                )

        if student_date_template_cell_col >= student_row_cell_count:
            raise InvalidRecipeError(
                f"student_date_template_cell_col ({student_date_template_cell_col}) out of range for student row cell count ({student_row_cell_count})."
            )
        for c in student_summary_cell_cols:
            if c >= student_row_cell_count:
                raise InvalidRecipeError(
                    f"student_summary_cell_cols index ({c}) out of range for student row cell count ({student_row_cell_count})."
                )
        if no_col >= student_row_cell_count:
            raise InvalidRecipeError(
                f"no_col ({no_col}) out of range for student row cell count ({student_row_cell_count})."
            )
        if name_col >= student_row_cell_count:
            raise InvalidRecipeError(
                f"name_col ({name_col}) out of range for student row cell count ({student_row_cell_count})."
            )
        if id_col >= student_row_cell_count:
            raise InvalidRecipeError(
                f"id_col ({id_col}) out of range for student row cell count ({student_row_cell_count})."
            )

        info_binding = AttendanceInfoBinding(
            table_index=tbl_idx,
            bindings=bindings,
            row_cell_counts=info_row_cell_counts,
        )

        matrix_binding = AttendanceMatrixBinding(
            table_index=m_tbl_idx,
            header_row0_index=h0_idx,
            header_row1_index=h1_idx,
            student_template_row_index=student_row_idx,
            no_col=no_col,
            name_col=name_col,
            id_col=id_col,
            date_columns_start=date_start,
            summary_columns_count=summary_columns_count,
            summary_column_names=summary_names,
            template_session_capacity=session_capacity,
            template_student_row_capacity=matrix_cand.get("template_student_row_capacity", 40),
            week_template_cell_col=week_template_cell_col,
            summary_header0_cell_col=summary_header0_cell_col,
            date_template_cell_col=date_template_cell_col,
            summary_column_indices=summary_column_indices,
            summary_header1_cell_cols=summary_header1_cell_cols,
            student_date_template_cell_col=student_date_template_cell_col,
            student_summary_cell_cols=student_summary_cell_cols,
            summary_column_widths=summary_widths,
            row0_cell_count=row0_cell_count,
            row1_cell_count=row1_cell_count,
            student_row_cell_count=student_row_cell_count,
            matrix_row_count=matrix_row_count,
        )

        metadata = dict(candidate.metadata)
        if "output_folder" not in metadata:
            metadata["output_folder"] = "Attendance"
        metadata["output_folder"] = validate_output_folder(metadata.get("output_folder", "Attendance"), default="Attendance")

        verified_safe = bool(metadata.pop("verified_safe", True))

        return ValidatedAttendanceTemplateRecipe(
            schema_version=RECIPE_SCHEMA_VERSION,
            profile_id=profile.profile_id if hasattr(profile, "profile_id") else "attendance_docx",
            fingerprint=candidate.fingerprint,
            template_path=candidate.template_path,
            info_binding=info_binding,
            matrix_binding=matrix_binding,
            metadata=metadata,
            verified_safe=verified_safe,
            _construction_token=_PRIVATE_CONSTRUCTION_SENTINEL,
        )

    @classmethod
    def with_metadata(
        cls,
        recipe: ValidatedRecipeBase,
        extra_metadata: Dict[str, Any],
    ) -> ValidatedRecipeBase:
        """
        Produces a new ValidatedRecipe with extra/overridden metadata without mutating the original.
        """
        if isinstance(recipe, ValidatedTemplateRecipe):
            merged_meta = dict(recipe.metadata)
            merged_meta.update(extra_metadata)
            if "output_folder" in merged_meta:
                merged_meta["output_folder"] = validate_output_folder(merged_meta["output_folder"], default="CEIT_Forms")
            return ValidatedTemplateRecipe(
                schema_version=recipe.schema_version,
                profile_id=recipe.profile_id,
                fingerprint=recipe.fingerprint,
                template_path=recipe.template_path,
                roster_binding=recipe.roster_binding,
                header_bindings=dict(recipe.header_bindings),
                signature_bindings=dict(recipe.signature_bindings),
                metadata=merged_meta,
                verified_safe=recipe.verified_safe,
                _construction_token=_PRIVATE_CONSTRUCTION_SENTINEL,
            )
        elif isinstance(recipe, ValidatedAttendanceTemplateRecipe):
            merged_meta = dict(recipe.metadata)
            merged_meta.update(extra_metadata)
            return ValidatedAttendanceTemplateRecipe(
                schema_version=recipe.schema_version,
                profile_id=recipe.profile_id,
                fingerprint=recipe.fingerprint,
                template_path=recipe.template_path,
                info_binding=recipe.info_binding,
                matrix_binding=recipe.matrix_binding,
                metadata=merged_meta,
                verified_safe=recipe.verified_safe,
                _construction_token=_PRIVATE_CONSTRUCTION_SENTINEL,
            )
        raise TypeError(f"Unsupported recipe type: {type(recipe)}")

    @classmethod
    def validate_dict(
        cls,
        data: Dict[str, Any],
        profile: Optional[GeneratorProfile] = None,
    ) -> ValidatedRecipeBase:
        """
        Validates and deserializes a serialized recipe dictionary.
        Strictly enforces schema_version == 2 and ignores/rejects external construction tokens.
        """
        if not isinstance(data, dict):
            raise InvalidRecipeError("Recipe data must be a dictionary")

        # Strict Schema Policy: accept only schema_version == 2
        schema_version = data.get("schema_version")
        if schema_version != RECIPE_SCHEMA_VERSION:
            raise InvalidRecipeError(
                f"Unsupported recipe schema_version: {schema_version}. "
                f"Expected {RECIPE_SCHEMA_VERSION}. Missing, v1, or future versions are rejected."
            )

        # Reject externally supplied construction token
        if "_construction_token" in data:
            raise InvalidRecipeError("Externally supplied '_construction_token' is prohibited.")
        clean_data = dict(data)

        # 1. Determine requested profile if passed
        requested_profile = None
        if profile is not None:
            if isinstance(profile, GeneratorProfile):
                requested_profile = profile
            elif isinstance(profile, str):
                prof_key = profile.strip().lower()
                if prof_key not in PROFILE_REGISTRY:
                    raise TemplateError(f"Unknown generator profile ID '{profile}'")
                requested_profile = PROFILE_REGISTRY[prof_key]
            else:
                raise InvalidRecipeError(f"Invalid profile specification: {profile}")

        # 2. Process serialized profile_id
        serialized_profile_id = clean_data.get("profile_id")
        if serialized_profile_id is not None and not isinstance(serialized_profile_id, str):
            raise InvalidRecipeError("Serialized recipe profile_id must be a string.")

        serialized_profile = None
        if serialized_profile_id is not None:
            norm_serialized_profile_id = serialized_profile_id.strip().lower()
            serialized_profile = PROFILE_REGISTRY.get(norm_serialized_profile_id)
            if serialized_profile is None:
                raise InvalidRecipeError(f"Unknown serialized profile ID '{serialized_profile_id}'.")

        # 3. Compare canonical profile identity if both are present
        if requested_profile is not None and serialized_profile is not None:
            if serialized_profile.profile_id != requested_profile.profile_id:
                has_attendance_bindings = (
                    "info_binding" in clean_data
                    or "matrix_binding" in clean_data
                    or "info_candidate" in clean_data
                    or "matrix_candidate" in clean_data
                )
                details = ""
                if has_attendance_bindings and requested_profile.profile_id != "attendance_docx":
                    details = f" Attendance data supplied with non-attendance profile '{requested_profile.profile_id}'."
                elif not has_attendance_bindings and requested_profile.profile_id == "attendance_docx":
                    details = f" Academic data supplied with attendance profile '{requested_profile.profile_id}'."
                raise InvalidRecipeError(
                    f"Profile mismatch: serialized profile '{serialized_profile_id}' (canonical '{serialized_profile.profile_id}') "
                    f"conflicts with requested profile '{requested_profile.profile_id}'.{details}"
                )

        if requested_profile is not None:
            profile = requested_profile
        elif serialized_profile is not None:
            profile = serialized_profile
        else:
            profile = PROFILE_ACADEMIC_DOCX

        # 4. Attendance deserialization & profile compatibility enforcement
        is_attendance_profile = (
            profile.profile_id in ("attendance_docx", "attendance")
            or profile.document_family == "attendance_docx"
        )
        has_attendance_bindings = (
            "info_binding" in clean_data
            or "matrix_binding" in clean_data
            or "info_candidate" in clean_data
            or "matrix_candidate" in clean_data
        )

        if is_attendance_profile or has_attendance_bindings:
            if serialized_profile_id is None or not serialized_profile_id.strip():
                raise InvalidRecipeError("Schema v2 attendance recipe requires 'profile_id'.")

        if not is_attendance_profile and has_attendance_bindings:
            raise InvalidRecipeError(
                f"Attendance data supplied with non-attendance profile '{profile.profile_id}'."
            )

        if is_attendance_profile:
            if not has_attendance_bindings:
                raise InvalidRecipeError(
                    f"Academic data supplied with attendance profile '{profile.profile_id}'."
                )

            # Bidirectional profile compatibility check
            if "detected_profile" in clean_data:
                det = clean_data["detected_profile"]
                det_prof = PROFILE_REGISTRY.get(str(det).strip().lower())
                if det_prof is None or det_prof.profile_id != profile.profile_id:
                    raise InvalidRecipeError(
                        f"Profile mismatch: dictionary specifies '{det}', but validation requested '{profile.profile_id}'."
                    )

            info_d = clean_data.get("info_binding") or clean_data.get("info_candidate")
            matrix_d = clean_data.get("matrix_binding") or clean_data.get("matrix_candidate")
            if not isinstance(info_d, dict):
                raise InvalidRecipeError("Attendance recipe requires dictionary 'info_binding'.")
            if not isinstance(matrix_d, dict):
                raise InvalidRecipeError("Attendance recipe requires dictionary 'matrix_binding'.")

            candidate_metadata = dict(clean_data.get("metadata", {}))
            if "verified_safe" in clean_data:
                candidate_metadata["verified_safe"] = clean_data["verified_safe"]

            candidate = RawAttendanceTemplateRecipeCandidate(
                template_path=clean_data.get("template_path", ""),
                profile_id=profile.profile_id,
                fingerprint=clean_data.get("fingerprint", clean_data.get("template_hash", "")),
                info_candidate=info_d,
                matrix_candidate=matrix_d,
                metadata=candidate_metadata,
                collisions=clean_data.get("collisions", []),
            )
            return cls._validate_attendance(candidate, profile)

        # Reconstruct components
        roster_data = clean_data.get("roster_binding") or clean_data.get("roster_table")

        raw_headers = clean_data.get("header_bindings", {})
        header_candidates: List[Dict[str, Any]] = []
        if isinstance(raw_headers, dict):
            header_candidates = [v if isinstance(v, dict) else v.to_dict() for v in raw_headers.values()]
        elif isinstance(raw_headers, list):
            for item in raw_headers:
                if isinstance(item, dict):
                    if item.get("type") == "table_cell":
                        header_candidates.append({
                            "cell_type": "docx_table",
                            "target": (item.get("table_index", 0), item.get("row_index", 0), item.get("cell_index", 1)),
                            "field": item.get("field"),
                            "shrink_threshold": item.get("shrink_threshold", 0),
                            "shrink_sz": item.get("shrink_sz", "18"),
                            "confidence": float(item.get("confidence", 1.0)),
                        })
                    elif item.get("type") == "paragraph_colon":
                        header_candidates.append({
                            "cell_type": "docx_paragraph",
                            "target": item.get("para_index", 0),
                            "field": item.get("field"),
                            "shrink_threshold": item.get("shrink_threshold", 0),
                            "shrink_sz": item.get("shrink_sz", "18"),
                            "confidence": float(item.get("confidence", 1.0)),
                        })
                    else:
                        header_candidates.append(item)

        raw_signatures = clean_data.get("signature_bindings", {})
        signature_candidates: List[Dict[str, Any]] = []
        if isinstance(raw_signatures, dict):
            signature_candidates = [v if isinstance(v, dict) else v.to_dict() for v in raw_signatures.values()]
        elif isinstance(raw_signatures, list):
            signature_candidates = [s for s in raw_signatures if isinstance(s, dict)]

        meta = clean_data.get("metadata")
        if not isinstance(meta, dict):
            meta = {
                "title": clean_data.get("title", ""),
                "suffix": clean_data.get("suffix", ""),
                "output_folder": clean_data.get("output_folder", "CEIT_Forms"),
                "placeholders": clean_data.get("placeholders", []),
            }
        else:
            meta = dict(meta)
            if "output_folder" not in meta and "output_folder" in clean_data:
                meta["output_folder"] = clean_data["output_folder"]

        if "docx_geometry" in clean_data and "docx_geometry" not in meta:
            meta["docx_geometry"] = clean_data["docx_geometry"]
        if "xlsx_geometry" in clean_data and "xlsx_geometry" not in meta:
            meta["xlsx_geometry"] = clean_data["xlsx_geometry"]

        # Construct candidate to run full validation checks
        candidate = RawTemplateRecipeCandidate(
            template_path=clean_data.get("template_path", ""),
            profile_id=profile.profile_id,
            fingerprint=clean_data.get("fingerprint", ""),
            roster_candidate=roster_data,
            header_candidates=header_candidates,
            signature_candidates=signature_candidates,
            collisions=[],
            metadata=meta,
        )

        return cls.validate(candidate, profile)

    @classmethod
    def _resolve_collisions(
        cls,
        candidate: RawTemplateRecipeCandidate,
        profile: GeneratorProfile,
    ) -> None:
        """
        Checks for ambiguity and collisions among candidates.
        Rejects structural ambiguity when multiple distinct physical targets
        satisfy the same semantic field without a unique structural discriminator.
        """
        for collision in candidate.collisions:
            field = collision.get("field")
            c_type = collision.get("type")
            if c_type in ("ambiguous_roster_table", "ambiguous_signature_binding"):
                raise AmbiguousTemplateError(
                    f"Ambiguous template structure in '{os.path.basename(candidate.template_path or 'template')}': {c_type}."
                )

            c_list = collision.get("candidates", [])
            if not c_list:
                continue

            viable = []
            for c in c_list:
                conf = float(c.get("confidence", 0.5))
                if conf < 0.70:
                    continue
                if profile.allowed_fields is not None and field and field not in profile.allowed_fields:
                    continue
                if field == FIELD_INSTRUCTOR and c.get("is_signature_region"):
                    continue
                viable.append(c)

            targets = set()
            for c in viable:
                t = c.get("target")
                t_key = (c.get("cell_type"), tuple(t) if isinstance(t, (list, tuple)) else t)
                targets.add(t_key)

            if len(targets) > 1:
                # Check for unique structural dominance
                ph_cands = [c for c in viable if c.get("derivation_evidence") == "explicit_placeholder" or c.get("has_placeholder")]
                ph_targets = {(c.get("cell_type"), tuple(c.get("target")) if isinstance(c.get("target"), (list, tuple)) else c.get("target")) for c in ph_cands}
                if len(ph_targets) == 1:
                    continue

                tbl_cands = [c for c in viable if c.get("cell_type") == "docx_table"]
                tbl_targets = {(c.get("cell_type"), tuple(c.get("target")) if isinstance(c.get("target"), (list, tuple)) else c.get("target")) for c in tbl_cands}
                if len(tbl_targets) == 1 and all(c.get("cell_type") == "docx_paragraph" for c in viable if c not in tbl_cands):
                    continue

                raise AmbiguousTemplateError(
                    f"Ambiguous template binding for field '{field}' in '{os.path.basename(candidate.template_path or 'template')}': "
                    f"Found {len(targets)} distinct candidate locations without structural dominance."
                )

    @classmethod
    def _validate_roster(
        cls,
        roster_data: Optional[Dict[str, Any]],
        profile: GeneratorProfile,
    ) -> Optional[RosterBinding]:
        """
        Validates roster binding requirements based on generator profile.
        """
        if not roster_data:
            if profile.requires_capacity or profile.document_family in ("academic_docx", "grade_sheet_xlsx"):
                if profile.profile_id != "custom_docx":
                    raise TemplateError(
                        f"Profile '{profile.profile_id}' strictly requires a student roster table, but none was detected."
                    )
            return None

        # Validate mandatory roster coordinates
        table_idx = roster_data.get("table_index")
        first_row = roster_data.get("first_data_row_index")
        name_col = roster_data.get("name_col")
        id_col = roster_data.get("id_col")

        if table_idx is None or first_row is None:
            raise TemplateError("Roster binding missing table_index or first_data_row_index")

        worksheet_name = roster_data.get("worksheet_name")
        if worksheet_name is not None and not isinstance(worksheet_name, str):
            raise TemplateError("Roster binding worksheet_name must be a string")

        if name_col is None and not roster_data.get("has_split_names"):
            raise TemplateError("Roster binding requires name_col or split names")

        if id_col is None:
            raise TemplateError("Roster binding requires id_col")

        # Grade sheet capacity and structural requirements
        capacity = roster_data.get("capacity_limit")
        if profile.requires_capacity or profile.document_family == "grade_sheet_xlsx":
            if capacity is None or capacity <= 0:
                raise TemplateError(
                    f"grade_sheet profile '{profile.profile_id}' requires capacity_limit > 0, "
                    f"got {capacity}."
                )
            if roster_data.get("index_col") is None:
                raise TemplateError(
                    f"grade_sheet profile '{profile.profile_id}' requires roster 'index_col' coordinate."
                )
            if not worksheet_name:
                raise TemplateError(
                    f"grade_sheet profile '{profile.profile_id}' requires roster 'worksheet_name'."
                )

        # Validate multi-row header structure if specified
        header_row_idx = roster_data.get("header_row_index", 0)
        header_row_cnt = roster_data.get("header_row_count", 1)
        if header_row_cnt < 1:
            raise TemplateError(f"header_row_count must be >= 1, got {header_row_cnt}")
        if first_row < header_row_idx + header_row_cnt:
            raise TemplateError(
                f"first_data_row_index ({first_row}) cannot precede headers "
                f"(header_row_index {header_row_idx} + count {header_row_cnt})"
            )

        return RosterBinding.from_dict(roster_data)

    @classmethod
    def _validate_headers(
        cls,
        header_candidates: List[Dict[str, Any]],
        profile: GeneratorProfile,
        template_path: Optional[str] = None,
    ) -> Dict[str, HeaderCellBinding]:
        """
        Validates header candidates into validated bindings.
        Rejects structural ambiguity when multiple distinct physical targets satisfy
        the same semantic field without a unique structural discriminator.
        """
        by_field: Dict[str, List[Dict[str, Any]]] = {}
        for c in header_candidates:
            f = c.get("field")
            if f:
                by_field.setdefault(f, []).append(c)

        bindings: Dict[str, HeaderCellBinding] = {}
        used_targets = set()

        for field, c_list in by_field.items():
            if profile.allowed_fields is not None and field not in profile.allowed_fields:
                continue

            viable = []
            for c in c_list:
                conf = float(c.get("confidence", 0.5))
                if conf < 0.70:
                    continue
                if field == FIELD_INSTRUCTOR and c.get("is_signature_region"):
                    continue
                viable.append(c)

            if not viable:
                continue

            target_to_candidates: Dict[Any, List[Dict[str, Any]]] = {}
            for c in viable:
                t = c.get("target")
                t_key = (c.get("cell_type"), tuple(t) if isinstance(t, (list, tuple)) else t)
                target_to_candidates.setdefault(t_key, []).append(c)

            if len(target_to_candidates) > 1:
                chosen_candidate = None

                ph_cands = [c for c in viable if c.get("derivation_evidence") == "explicit_placeholder" or c.get("has_placeholder")]
                ph_targets = {(c.get("cell_type"), tuple(c.get("target")) if isinstance(c.get("target"), (list, tuple)) else c.get("target")) for c in ph_cands}
                if len(ph_targets) == 1:
                    chosen_candidate = ph_cands[0]
                elif len(ph_targets) > 1:
                    raise AmbiguousTemplateError(
                        f"Ambiguous template binding for field '{field}' in '{os.path.basename(template_path or 'template')}': "
                        f"Multiple explicit placeholder targets detected for '{field}'."
                    )
                else:
                    tbl_cands = [c for c in viable if c.get("cell_type") == "docx_table"]
                    tbl_targets = {(c.get("cell_type"), tuple(c.get("target")) if isinstance(c.get("target"), (list, tuple)) else c.get("target")) for c in tbl_cands}
                    if len(tbl_targets) == 1 and all(c.get("cell_type") == "docx_paragraph" for c in viable if c not in tbl_cands):
                        chosen_candidate = tbl_cands[0]
                    else:
                        raise AmbiguousTemplateError(
                            f"Ambiguous template binding for field '{field}' in '{os.path.basename(template_path or 'template')}': "
                            f"Multiple distinct physical targets satisfy '{field}' without structural dominance."
                        )
            else:
                chosen_candidate = max(viable, key=lambda c: float(c.get("confidence", 0.5)))

            target_key = (chosen_candidate.get("cell_type"), tuple(chosen_candidate.get("target")) if isinstance(chosen_candidate.get("target"), (list, tuple)) else chosen_candidate.get("target"))
            if target_key not in used_targets:
                bindings[field] = HeaderCellBinding.from_dict(chosen_candidate)
                used_targets.add(target_key)

        return bindings

    @classmethod
    def _validate_signatures(
        cls,
        signature_candidates: List[Dict[str, Any]],
        profile: GeneratorProfile,
        template_path: Optional[str] = None,
    ) -> Dict[str, SignatureBinding]:
        """
        Validates signature candidates into validated signature bindings.
        Rejects multiple distinct physical targets for the same role without structural dominance.
        """
        by_role: Dict[str, List[Dict[str, Any]]] = {}
        for c in signature_candidates:
            r = c.get("role")
            if r:
                by_role.setdefault(r, []).append(c)

        bindings: Dict[str, SignatureBinding] = {}

        for role, c_list in by_role.items():
            viable = [c for c in c_list if float(c.get("confidence", 0.5)) >= 0.60]
            if not viable:
                continue

            target_to_candidates: Dict[Any, List[Dict[str, Any]]] = {}
            for c in viable:
                t = c.get("target")
                t_key = tuple(t) if isinstance(t, (list, tuple)) else t
                target_to_candidates.setdefault(t_key, []).append(c)

            if len(target_to_candidates) > 1:
                dom = [c for c in viable if c.get("derivation_evidence") == "structural_merged_box_above_label"]
                dom_targets = {tuple(c.get("target")) if isinstance(c.get("target"), (list, tuple)) else c.get("target") for c in dom}
                if len(dom_targets) == 1:
                    chosen = dom[0]
                else:
                    raise AmbiguousTemplateError(
                        f"Ambiguous signature binding for role '{role}' in '{os.path.basename(template_path or 'template')}': "
                        f"Found {len(target_to_candidates)} distinct signature candidate locations without structural dominance."
                    )
            else:
                chosen = max(viable, key=lambda c: float(c.get("confidence", 0.5)))

            bindings[role] = SignatureBinding.from_dict(chosen)

        return bindings

    @classmethod
    def _validate_docx_geometry(
        cls,
        candidate: RawTemplateRecipeCandidate,
        roster_binding: Optional[RosterBinding],
        header_bindings: Dict[str, HeaderCellBinding],
        signature_bindings: Dict[str, SignatureBinding],
    ) -> None:
        """
        Validates every physical DOCX coordinate (headers, roster, signatures)
        against actual template geometry (or serialized docx_geometry metadata if template file is absent).
        Raises InvalidRecipeError on any out-of-bounds or invalid coordinate.
        """
        table_row_cell_counts = None
        paragraph_count = None

        if candidate.template_path and os.path.exists(candidate.template_path):
            try:
                zin, root, body = load_docx(candidate.template_path)
                zin.close()
                tables = body.findall(w("tbl"))
                paragraphs = body.findall(w("p"))
                paragraph_count = len(paragraphs)
                table_row_cell_counts = []
                for tbl in tables:
                    rows = tbl.findall(w("tr"))
                    row_counts = [len(tr.findall(w("tc"))) for tr in rows]
                    table_row_cell_counts.append(row_counts)
            except (TemplateError, InvalidRecipeError):
                raise
            except Exception as e:
                raise TemplateError(
                    f"Physical template inspection failed for '{candidate.template_path}': {e}"
                ) from e
        else:
            geom = (candidate.metadata or {}).get("docx_geometry")
            has_coords = bool(header_bindings or roster_binding or signature_bindings)
            if not isinstance(geom, dict):
                if has_coords:
                    raise InvalidRecipeError(
                        f"Serialized DOCX recipe for '{os.path.basename(candidate.template_path or 'template')}' "
                        f"is missing required 'docx_geometry' metadata when physical template is unavailable."
                    )
                return
            table_row_cell_counts = geom.get("table_row_cell_counts")
            paragraph_count = geom.get("paragraph_count")
            if table_row_cell_counts is None or paragraph_count is None:
                if has_coords:
                    raise InvalidRecipeError(
                        f"Serialized DOCX recipe for '{os.path.basename(candidate.template_path or 'template')}' "
                        f"has invalid 'docx_geometry' metadata."
                    )
                return

        num_tables = len(table_row_cell_counts)

        # 1. Validate Header Bindings
        for field, binding in header_bindings.items():
            if binding.cell_type == "docx_table":
                target = binding.target
                if not isinstance(target, (list, tuple)) or len(target) != 3:
                    raise InvalidRecipeError(
                        f"Header binding for '{field}' has invalid table target {target} (must be (tbl, row, cell))."
                    )
                t_idx, r_idx, c_idx = target[0], target[1], target[2]
                if not isinstance(t_idx, int) or t_idx < 0 or t_idx >= num_tables:
                    raise InvalidRecipeError(
                        f"Header binding for '{field}' table index {t_idx} out of range (document has {num_tables} tables)."
                    )
                tbl_rows = table_row_cell_counts[t_idx]
                if not isinstance(r_idx, int) or r_idx < 0 or r_idx >= len(tbl_rows):
                    raise InvalidRecipeError(
                        f"Header binding for '{field}' row index {r_idx} out of range for table {t_idx} (has {len(tbl_rows)} rows)."
                    )
                row_cells = tbl_rows[r_idx]
                if not isinstance(c_idx, int) or c_idx < 0 or c_idx >= row_cells:
                    raise InvalidRecipeError(
                        f"Header binding for '{field}' cell index {c_idx} out of range for table {t_idx}, row {r_idx} (has {row_cells} cells)."
                    )
            elif binding.cell_type == "docx_paragraph":
                target = binding.target
                p_idx = target[0] if isinstance(target, (list, tuple)) and len(target) == 1 else target
                if not isinstance(p_idx, int) or p_idx < 0 or p_idx >= paragraph_count:
                    raise InvalidRecipeError(
                        f"Header binding for '{field}' paragraph index {p_idx} out of range (document has {paragraph_count} paragraphs)."
                    )

        # 2. Validate Roster Binding
        if roster_binding:
            t_idx = roster_binding.table_index
            if not isinstance(t_idx, int) or t_idx < 0 or t_idx >= num_tables:
                raise InvalidRecipeError(
                    f"Roster binding table index {t_idx} out of range (document has {num_tables} tables)."
                )
            tbl_rows = table_row_cell_counts[t_idx]
            num_rows = len(tbl_rows)
            h_idx = roster_binding.header_row_index
            h_count = roster_binding.header_row_count
            f_idx = roster_binding.first_data_row_index

            if not isinstance(h_idx, int) or h_idx < 0 or h_idx >= num_rows:
                raise InvalidRecipeError(
                    f"Roster header_row_index {h_idx} out of range for table {t_idx} (has {num_rows} rows)."
                )
            if h_idx + h_count > num_rows:
                raise InvalidRecipeError(
                    f"Roster header_row_index ({h_idx}) + header_row_count ({h_count}) exceeds table {t_idx} rows ({num_rows})."
                )
            if not isinstance(f_idx, int) or f_idx < 0 or f_idx >= num_rows:
                raise InvalidRecipeError(
                    f"Roster first_data_row_index {f_idx} out of range for table {t_idx} (has {num_rows} rows)."
                )

            data_row_cells = tbl_rows[f_idx]
            cols_to_check = [("name_col", roster_binding.name_col), ("id_col", roster_binding.id_col)]
            if roster_binding.index_col is not None:
                cols_to_check.append(("index_col", roster_binding.index_col))
            if roster_binding.signature_col is not None:
                cols_to_check.append(("signature_col", roster_binding.signature_col))
            if roster_binding.has_split_names:
                for col_name, c_val in [
                    ("last_name_col", roster_binding.last_name_col),
                    ("first_name_col", roster_binding.first_name_col),
                    ("middle_name_col", roster_binding.middle_name_col),
                ]:
                    if c_val is not None:
                        cols_to_check.append((col_name, c_val))

            for col_name, col_val in cols_to_check:
                if not isinstance(col_val, int) or col_val < 0 or col_val >= data_row_cells:
                    raise InvalidRecipeError(
                        f"Roster {col_name} ({col_val}) out of range for table {t_idx} row {f_idx} (has {data_row_cells} cells)."
                    )

        # 3. Validate Signature Bindings
        for role, sig in signature_bindings.items():
            target = sig.target
            if isinstance(target, (list, tuple)) and len(target) == 3:
                t_idx, r_idx, c_idx = target[0], target[1], target[2]
                if not isinstance(t_idx, int) or t_idx < 0 or t_idx >= num_tables:
                    raise InvalidRecipeError(
                        f"Signature binding for '{role}' table index {t_idx} out of range (document has {num_tables} tables)."
                    )
                tbl_rows = table_row_cell_counts[t_idx]
                if not isinstance(r_idx, int) or r_idx < 0 or r_idx >= len(tbl_rows):
                    raise InvalidRecipeError(
                        f"Signature binding for '{role}' row index {r_idx} out of range for table {t_idx} (has {len(tbl_rows)} rows)."
                    )
                row_cells = tbl_rows[r_idx]
                if not isinstance(c_idx, int) or c_idx < 0 or c_idx >= row_cells:
                    raise InvalidRecipeError(
                        f"Signature binding for '{role}' cell index {c_idx} out of range for table {t_idx}, row {r_idx} (has {row_cells} cells)."
                    )
            elif isinstance(target, int) or (isinstance(target, (list, tuple)) and len(target) == 1):
                p_idx = target[0] if isinstance(target, (list, tuple)) else target
                if not isinstance(p_idx, int) or p_idx < 0 or p_idx >= paragraph_count:
                    raise InvalidRecipeError(
                        f"Signature binding for '{role}' paragraph index {p_idx} out of range (document has {paragraph_count} paragraphs)."
                    )

    @classmethod
    def _validate_xlsx_geometry(
        cls,
        candidate: RawTemplateRecipeCandidate,
        roster_binding: Optional[RosterBinding],
        header_bindings: Dict[str, HeaderCellBinding],
        signature_bindings: Dict[str, SignatureBinding],
    ) -> None:
        """
        Validates every physical XLSX coordinate (headers, roster, signatures)
        against actual workbook geometry (or serialized xlsx_geometry metadata if workbook is absent).
        Raises InvalidRecipeError on any out-of-bounds or invalid coordinate.
        """
        worksheets_geom: Dict[str, Dict[str, int]] = {}

        if candidate.template_path and os.path.exists(candidate.template_path):
            try:
                wb = openpyxl.load_workbook(candidate.template_path, data_only=True)
                for s_name in wb.sheetnames:
                    ws = wb[s_name]
                    worksheets_geom[s_name] = {
                        "max_row": ws.max_row,
                        "max_column": ws.max_column,
                    }
                wb.close()
            except (TemplateError, InvalidRecipeError):
                raise
            except Exception as e:
                raise TemplateError(
                    f"Physical template inspection failed for '{candidate.template_path}': {e}"
                ) from e
        else:
            geom = (candidate.metadata or {}).get("xlsx_geometry")
            has_coords = bool(header_bindings or roster_binding or signature_bindings)
            if not isinstance(geom, dict) or "worksheets" not in geom:
                if has_coords:
                    raise InvalidRecipeError(
                        f"Serialized XLSX recipe for '{os.path.basename(candidate.template_path or 'template')}' "
                        f"is missing required 'xlsx_geometry' metadata when physical workbook is unavailable."
                    )
                return
            worksheets_geom = geom["worksheets"]

        def get_ws_bounds(sheet_name: Optional[str]) -> Tuple[str, int, int]:
            if not sheet_name:
                if not worksheets_geom:
                    raise InvalidRecipeError("Workbook has no worksheets.")
                first_name = next(iter(worksheets_geom.keys()))
                return first_name, worksheets_geom[first_name]["max_row"], worksheets_geom[first_name]["max_column"]
            for k, v in worksheets_geom.items():
                if k.lower() == sheet_name.lower():
                    return k, v["max_row"], v["max_column"]
            raise InvalidRecipeError(
                f"Worksheet '{sheet_name}' not found in workbook (available: {list(worksheets_geom.keys())})."
            )

        def parse_xlsx_cell(coord_str: str, default_sheet: Optional[str] = None) -> Tuple[str, int, int]:
            target_sheet = default_sheet
            target_cell = coord_str
            if "!" in coord_str:
                target_sheet, target_cell = coord_str.split("!", 1)
            actual_sheet, max_r, max_c = get_ws_bounds(target_sheet)
            try:
                row, col = coordinate_to_tuple(target_cell)
            except Exception as e:
                raise InvalidRecipeError(f"Invalid cell coordinate '{coord_str}': {e}") from e
            return actual_sheet, row, col

        default_roster_sheet = roster_binding.worksheet_name if roster_binding else None

        # 1. Validate Header Bindings
        for field, binding in header_bindings.items():
            if binding.cell_type == "xlsx_cell" or isinstance(binding.target, str):
                coord = str(binding.target)
                actual_sheet, r, c = parse_xlsx_cell(coord, default_sheet=default_roster_sheet)
                _, max_r, max_c = get_ws_bounds(actual_sheet)
                if r < 1 or r > max_r or c < 1 or c > max_c:
                    raise InvalidRecipeError(
                        f"Header binding for '{field}' coordinate '{coord}' out of bounds for sheet '{actual_sheet}' "
                        f"(row {r} not in [1, {max_r}] or col {c} not in [1, {max_c}])."
                    )

        # 2. Validate Roster Binding
        if roster_binding:
            r_sheet = roster_binding.worksheet_name
            if not r_sheet:
                raise InvalidRecipeError("Grade sheet roster binding is missing worksheet_name.")
            actual_sheet, max_r, max_c = get_ws_bounds(r_sheet)

            f_row = roster_binding.first_data_row_index
            if f_row < 1 or f_row > max_r:
                raise InvalidRecipeError(
                    f"Roster first_data_row_index {f_row} out of bounds for sheet '{actual_sheet}' (max_row {max_r})."
                )

            name_col = roster_binding.name_col
            id_col = roster_binding.id_col
            if name_col < 1 or name_col > max_c:
                raise InvalidRecipeError(
                    f"Roster name_col {name_col} out of bounds for sheet '{actual_sheet}' (max_column {max_c})."
                )
            if id_col < 1 or id_col > max_c:
                raise InvalidRecipeError(
                    f"Roster id_col {id_col} out of bounds for sheet '{actual_sheet}' (max_column {max_c})."
                )

            if roster_binding.index_col is not None:
                idx_col = roster_binding.index_col
                if idx_col < 1 or idx_col > max_c:
                    raise InvalidRecipeError(
                        f"Roster index_col {idx_col} out of bounds for sheet '{actual_sheet}' (max_column {max_c})."
                    )

            if roster_binding.capacity_limit is not None:
                cap = roster_binding.capacity_limit
                if cap <= 0:
                    raise InvalidRecipeError(f"Roster capacity_limit must be > 0, got {cap}.")
                if f_row + cap - 1 > max_r:
                    raise InvalidRecipeError(
                        f"Roster capacity range [{f_row}, {f_row + cap - 1}] exceeds sheet '{actual_sheet}' max_row ({max_r})."
                    )

        # 3. Validate Signatures
        for role, sig in signature_bindings.items():
            if isinstance(sig.target, str):
                coord = sig.target
                default_sheet = default_roster_sheet
                if "!" not in coord:
                    r_lower = role.lower()
                    if "laboratory" in r_lower:
                        default_sheet = "laboratory"
                    elif "consolidated" in r_lower:
                        default_sheet = "consolidated"

                actual_sheet, r, c = parse_xlsx_cell(coord, default_sheet=default_sheet)
                _, max_r, max_c = get_ws_bounds(actual_sheet)
                if r < 1 or r > max_r or c < 1 or c > max_c:
                    raise InvalidRecipeError(
                        f"Signature binding for '{role}' coordinate '{coord}' out of bounds for sheet '{actual_sheet}' "
                        f"(row {r} not in [1, {max_r}] or col {c} not in [1, {max_c}])."
                    )
