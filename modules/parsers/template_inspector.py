#!/usr/bin/env python3
"""
modules/parsers/template_inspector.py

Authoritative Template Inspectors and Compatibility Facade.
DocxTemplateInspector and XlsxTemplateInspector emit RawTemplateRecipeCandidate ONLY.
Inspectors NEVER construct or return ValidatedTemplateRecipe.
RecipeValidator owns validation authority.
"""

import os
import re
import hashlib
from typing import Dict, Any, List, Optional, Tuple

import openpyxl
from openpyxl.utils import get_column_letter
from openpyxl.formula.tokenizer import Tokenizer

from modules.common.logger import logger
from modules.common.docx_utils import load_docx, get_full_text, w
from modules.models.recipe import (
    TemplateError,
    AmbiguousTemplateError,
    GeneratorProfile,
    PROFILE_REGISTRY,
    PROFILE_ACADEMIC_DOCX,
    PROFILE_GRADE_SHEET_XLSX,
    PROFILE_ATTENDANCE_DOCX,
    RawTemplateRecipeCandidate,
    RawAttendanceTemplateRecipeCandidate,
    ValidatedTemplateRecipe,
    ValidatedAttendanceTemplateRecipe,
)
from modules.parsers.semantic_registry import (
    SemanticRegistry,
    FIELD_INSTRUCTOR,
    FIELD_COURSE_SECTION,
    FIELD_SCHEDULE_CODE,
    FIELD_SUBJECT,
    FIELD_SEMESTER_AY,
    FIELD_DATE,
    FIELD_TIME_DAYS_ROOM,
    ROLE_INSTRUCTOR_SIGNATURE,
)
from modules.parsers.recipe_validator import RecipeValidator


# Standard font shrink thresholds matching CvSU form specs
FIELD_SHRINK_THRESHOLDS = {
    "instructor": 30,
    "course_section": 0,
    "schedule_code": 0,
    "subject": 40,
    "time_days_room": 0,
    "semester_ay": 0,
    "date": 0,
}

class SheetReferenceSet(set):
    """
    A set of worksheet names extracted from an Excel formula,
    carrying an is_reliable flag indicating whether the formula
    was parsed conclusively without syntax corruption or ambiguity.
    """
    def __init__(self, iterable=(), is_reliable: bool = True):
        super().__init__(iterable)
        self.is_reliable = bool(is_reliable)


CELL_ADDR_PATTERN = re.compile(
    r"^(?:\$?[A-Za-z]+\$?[0-9]+(?::\$?[A-Za-z]+\$?[0-9]+)?|\$?[A-Za-z]+:\$?[A-Za-z]+|\$?[0-9]+:\$?[0-9]+|#REF!|[A-Za-z_][A-Za-z0-9_]*)$"
)

STRICT_CELL_REF = r"(?:\$?[A-Za-z]+\$?[0-9]+(?::\$?[A-Za-z]+\$?[0-9]+)?|\$?[A-Za-z]+:\$?[A-Za-z]+|\$?[0-9]+:\$?[0-9]+|#REF!)"
CONSERVATIVE_SHEET_REF_REGEX = re.compile(
    rf"(?:'((?:[^']|'')*)'|([A-Za-z0-9_]+))!({STRICT_CELL_REF})",
    re.IGNORECASE
)

# Unparseable syntax or illegal characters outside standard Excel formula grammar
CORRUPTED_CHARS_REGEX = re.compile(r"[^A-Za-z0-9_ \t\r\n\+\-\*\/\^\=\<\>\&\:\%\$\#\(\)\,\'\"\.\!]")


def split_sheet_and_cell(val: str) -> Tuple[Optional[str], str]:
    """Splits a RANGE token value into (sheet_part, cell_part)."""
    if val.startswith("'"):
        idx = 1
        while idx < len(val):
            if val[idx] == "'":
                if idx + 1 < len(val) and val[idx + 1] == "'":
                    idx += 2
                    continue
                elif idx + 1 < len(val) and val[idx + 1] == "!":
                    return val[:idx + 1], val[idx + 2:]
            idx += 1
    elif "!" in val:
        parts = val.split("!", 1)
        return parts[0], parts[1]
    return None, val


def extract_referenced_sheets(formula_str: str) -> SheetReferenceSet:
    """
    Extracts all worksheet names referenced in an Excel formula string.
    Supports:
      - Unquoted: Sheet1!A1
      - Quoted with spaces: 'Sheet A'!A1
      - Escaped apostrophes: 'Dean''s Practical Sheet'!C7
      - Ranges: 'Summary 2026'!A1:A20
      - Absolute refs: 'Practical Component'!$B$5
    Uses OpenXML/openpyxl formula tokenizer with conservative validation.
    Returns a SheetReferenceSet with is_reliable flag.
    If the formula cannot be parsed reliably (or contains dynamic/hidden/external references),
    lineage = unknown (is_reliable = False) and no incomplete or false dependency edges are created.
    """
    if not isinstance(formula_str, str) or not formula_str.startswith("="):
        return SheetReferenceSet(set(), is_reliable=True)

    # 1. Dynamic / hidden references check:
    # Functions like INDIRECT construct references dynamically at runtime.
    # Statically, dependency is indeterminate: UNKNOWN != NO DEPENDENCY.
    if re.search(r"\bINDIRECT\s*\(", formula_str, re.IGNORECASE):
        return SheetReferenceSet(set(), is_reliable=False)

    # 2. External workbook references check:
    # References containing brackets like [Book.xlsx]Sheet!A1 point to external files.
    # They must NEVER be converted into local worksheet dependencies.
    if "[" in formula_str or "]" in formula_str:
        return SheetReferenceSet(set(), is_reliable=False)

    try:
        tok = Tokenizer(formula_str)
        referenced = set()
        for item in tok.items:
            # Check for dynamic reference functions in tokens
            if item.type == "FUNC" and item.value.upper().startswith("INDIRECT("):
                return SheetReferenceSet(set(), is_reliable=False)

            if "!" in item.value:
                if item.type == "OPERAND" and item.subtype == "RANGE":
                    # External reference in range token
                    if "[" in item.value or "]" in item.value:
                        return SheetReferenceSet(set(), is_reliable=False)

                    sheet_part, cell_part = split_sheet_and_cell(item.value)
                    if not sheet_part or not cell_part:
                        return SheetReferenceSet(set(), is_reliable=False)
                    # 3D references across multiple sheets: Sheet1:Sheet3!A1
                    if ":" in sheet_part:
                        return SheetReferenceSet(set(), is_reliable=False)
                    if not CELL_ADDR_PATTERN.match(cell_part):
                        return SheetReferenceSet(set(), is_reliable=False)
                    if sheet_part.startswith("'") and sheet_part.endswith("'"):
                        sheet_name = sheet_part[1:-1].replace("''", "'")
                    else:
                        sheet_name = sheet_part
                    if not sheet_name or "[" in sheet_name or "]" in sheet_name or ":" in sheet_name:
                        return SheetReferenceSet(set(), is_reliable=False)
                    referenced.add(sheet_name)
                elif item.type == "OPERAND" and item.subtype == "TEXT":
                    pass
                else:
                    return SheetReferenceSet(set(), is_reliable=False)
        return SheetReferenceSet(referenced, is_reliable=True)
    except Exception:
        pass

    if "!" not in formula_str:
        return SheetReferenceSet(set(), is_reliable=True)

    # Conservative validation when Tokenizer fails:
    # 1. Balanced quotes
    if formula_str.replace("''", "").count("'") % 2 != 0:
        return SheetReferenceSet(set(), is_reliable=False)
    if formula_str.replace('""', '').count('"') % 2 != 0:
        return SheetReferenceSet(set(), is_reliable=False)

    # 2. Check for corrupted/invalid characters outside normal formula grammar
    if CORRUPTED_CHARS_REGEX.search(formula_str):
        return SheetReferenceSet(set(), is_reliable=False)

    # 3. All '!' occurrences must match strict sheet reference pattern
    matches = list(CONSERVATIVE_SHEET_REF_REGEX.finditer(formula_str))
    if len(matches) == 0 or len(matches) != formula_str.count("!"):
        return SheetReferenceSet(set(), is_reliable=False)

    referenced = set()
    for m in matches:
        quoted, unquoted, _ = m.groups()
        name = quoted.replace("''", "'") if quoted is not None else unquoted
        if not name or ":" in name or "[" in name or "]" in name:
            return SheetReferenceSet(set(), is_reliable=False)
        referenced.add(name)
    return SheetReferenceSet(referenced, is_reliable=True)


def is_aggregation_formula(val: str, other_sheet: str, all_sheets: set) -> bool:
    """
    Checks if a formula string physically represents an aggregation formula.
    An aggregation formula:
      - References multiple distinct worksheets (e.g. '=SheetA!A1 + SheetB!B1'), OR
      - References other_sheet and combines it mathematically using arithmetic operators
        (+ - * / ^) or aggregation functions (SUM, AVERAGE, PRODUCT, etc.),
        rather than being a pure single-cell mirror/passthrough (e.g. '=SheetB!A1').
    """
    if not isinstance(val, str) or not val.startswith("="):
        return False

    refs = extract_referenced_sheets(val)
    if not getattr(refs, "is_reliable", True):
        return False

    # 1. Multi-source formula: references multiple distinct worksheets in the workbook
    valid_refs = {r for r in refs if any(r.lower() == s.lower() for s in all_sheets)}
    if len(valid_refs) >= 2:
        return True

    # 2. Mathematical combination / aggregation of other_sheet:
    if any(r.lower() == other_sheet.lower() for r in refs):
        try:
            tok = Tokenizer(val)
            meaningful_tokens = [
                item for item in tok.items
                if item.type not in ("WHITE-SPACE",)
            ]
            if len(meaningful_tokens) == 1 and meaningful_tokens[0].type == "OPERAND" and meaningful_tokens[0].subtype == "RANGE":
                # Single-cell passthrough mirror: not an aggregation
                return False

            has_operators = any(item.type in ("OPERATOR-INFIX", "OPERATOR-PREFIX", "OPERATOR-POSTFIX") for item in meaningful_tokens)
            has_functions = any(item.type == "FUNC" for item in meaningful_tokens)
            if has_operators or has_functions:
                return True
        except Exception:
            pass

    return False


def has_aggregation_structure(ws, other_sheet: str, all_sheets: set) -> bool:
    """
    Scans a worksheet to verify whether it physically demonstrates an aggregation
    structure combining the other candidate assessment sheet with broader instructional
    or student components.
    """
    for row in ws.iter_rows(values_only=True):
        for val in row:
            if isinstance(val, str) and val.startswith("="):
                if is_aggregation_formula(val, other_sheet, all_sheets):
                    return True
    return False


# ═══════════════════════════════════════════════════════════════════════════════
# DocxTemplateInspector
# ═══════════════════════════════════════════════════════════════════════════════

class DocxTemplateInspector:
    """
    Analyzes Word (.docx) templates and emits candidate observations.
    CRITICAL: inspect() returns RawTemplateRecipeCandidate.
    DocxTemplateInspector NEVER constructs or returns ValidatedTemplateRecipe.
    """

    def inspect(
        self,
        template_path: str,
        profile_id: str = "academic_docx",
    ) -> RawTemplateRecipeCandidate:
        """
        Inspects a Word (.docx) template file and produces raw candidate observations.
        """
        if not os.path.exists(template_path):
            raise FileNotFoundError(f"Template file not found: {template_path}")

        zin, root, body = load_docx(template_path)
        zin.close()
        tables = body.findall(w("tbl"))
        paragraphs = body.findall(w("p"))

        # Compute file fingerprint
        with open(template_path, "rb") as f:
            fingerprint = hashlib.sha256(f.read()).hexdigest()

        header_candidates: List[Dict[str, Any]] = []
        signature_candidates: List[Dict[str, Any]] = []
        roster_candidate: Optional[Dict[str, Any]] = None

        total_tables = len(tables)

        # 0. Check all tables to detect candidate roster tables
        candidate_roster_tables = []
        for t_idx, tbl in enumerate(tables):
            r_info = self._detect_roster_table(tbl, t_idx)
            if r_info:
                candidate_roster_tables.append(r_info)

        if len(candidate_roster_tables) > 1:
            # Check for structural dominance:
            # Table with repeated student rows and Word header markup structurally dominates a table without
            dom = [
                c for c in candidate_roster_tables
                if c.get("total_rows", 0) > c.get("first_data_row_index", 0) + 1
            ]
            if len(dom) == 1:
                roster_candidate = dom[0]
            else:
                raise AmbiguousTemplateError(
                    f"Multiple candidate roster tables detected in '{os.path.basename(template_path)}': "
                    f"tables {[c['table_index'] for c in candidate_roster_tables]}. Ambiguous roster tables detected."
                )
        elif candidate_roster_tables:
            roster_candidate = candidate_roster_tables[0]
        else:
            roster_candidate = None

        # 1. Inspect Table Cells for Metadata and Signatures
        for t_idx, tbl in enumerate(tables):
            if roster_candidate and t_idx == roster_candidate["table_index"]:
                continue

            rows = tbl.findall(w("tr"))
            if not rows:
                continue

            # Otherwise, inspect table cells for metadata and signatures
            total_rows = len(rows)
            for r_idx, row in enumerate(rows):
                cells = row.findall(w("tc"))
                for c_idx, cell in enumerate(cells):
                    text = SemanticRegistry.normalize_text(get_full_text(cell))
                    if not text:
                        continue

                    # Determine if cell is in signature context
                    is_sig = SemanticRegistry.is_signature_context(
                        text,
                        table_index=t_idx,
                        total_tables=total_tables,
                        row_index=r_idx,
                        total_rows=total_rows,
                    )

                    # Determine value cell adjacent to label cell
                    if len(cells) == 2:
                        target_c_idx = 1
                    elif len(cells) >= 3 and c_idx == 0:
                        cell1_txt = SemanticRegistry.normalize_text(get_full_text(cells[1]))
                        if cell1_txt == ":":
                            target_c_idx = 2
                        else:
                            target_c_idx = 1
                    elif c_idx + 1 < len(cells):
                        target_c_idx = c_idx + 1
                    else:
                        target_c_idx = c_idx

                    # Metadata candidate
                    meta_matches = SemanticRegistry.match_metadata_candidate(text, is_signature_region=is_sig)
                    for field_name, conf in meta_matches:
                        header_candidates.append({
                            "cell_type": "docx_table",
                            "target": (t_idx, r_idx, target_c_idx),
                            "field": field_name,
                            "confidence": conf,
                            "pattern": text,
                            "shrink_threshold": 0 if profile_id == "custom_docx" else FIELD_SHRINK_THRESHOLDS.get(field_name, 0),
                            "shrink_sz": "18",
                            "is_signature_region": is_sig,
                        })

                    # Signature candidate
                    sig_matches = SemanticRegistry.match_signature_candidate(
                        text,
                        in_signature_table=is_sig,
                        relative_vertical_pos=float(r_idx) / max(total_rows, 1),
                    )
                    for role_name, conf in sig_matches:
                        signature_candidates.append({
                            "role": role_name,
                            "target": (t_idx, r_idx, c_idx),
                            "confidence": conf,
                        })

        # 2. Inspect Paragraphs (Outside tables)
        for p_idx, p in enumerate(paragraphs):
            text = SemanticRegistry.normalize_text(get_full_text(p))
            if not text:
                continue

            # Skip paragraphs that contain explicit template tags (handled by placeholder engine)
            if "{{" in text or "}}" in text or "[[" in text or "]]" in text:
                continue

            # Split paragraph into segments if delimited by '|' or ';'
            segments = [s.strip() for s in re.split(r'[|;]', text)]
            for seg in segments:
                if ":" in seg:
                    lbl_part, _ = seg.split(":", 1)
                    lbl_norm = SemanticRegistry.normalize_text(lbl_part)
                    meta_matches = SemanticRegistry.match_metadata_candidate(lbl_norm, is_signature_region=False)
                    for field_name, conf in meta_matches:
                        header_candidates.append({
                            "cell_type": "docx_paragraph",
                            "target": p_idx,
                            "field": field_name,
                            "confidence": conf * 0.9,
                            "pattern": lbl_norm,
                            "shrink_threshold": 0 if profile_id == "custom_docx" else FIELD_SHRINK_THRESHOLDS.get(field_name, 0),
                            "shrink_sz": "18",
                            "is_signature_region": False,
                        })

        # 3. Detect explicit placeholder tokens {{...}} or [...]
        placeholders = self._detect_placeholders(body)

        # 4. Observe collisions
        collisions = []
        collisions.extend(SemanticRegistry.observe_collisions(header_candidates))

        # Observe signature collisions
        sig_map: Dict[str, List[Dict[str, Any]]] = {}
        for s in signature_candidates:
            r = s.get("role")
            if r:
                sig_map.setdefault(r, []).append(s)
        for r, sc_list in sig_map.items():
            targets = {s.get("target") if not isinstance(s.get("target"), list) else tuple(s.get("target")) for s in sc_list}
            if len(targets) > 1:
                dom = [s for s in sc_list if s.get("derivation_evidence") == "structural_merged_box_above_label"]
                if len(dom) != 1:
                    collisions.append({
                        "type": "ambiguous_signature_binding",
                        "role": r,
                        "candidates": sc_list,
                    })

        # 5. Generate title, suffix, and confidence
        title, suffix = self._suggest_title_and_suffix(template_path, paragraphs)
        confidence_score = self._compute_confidence(roster_candidate, header_candidates, placeholders=placeholders)

        table_row_cell_counts = [
            tuple(len(r.findall(w("tc"))) for r in tbl.findall(w("tr")))
            for tbl in tables
        ]

        metadata = {
            "title": title,
            "suffix": suffix,
            "confidence": confidence_score,
            "total_tables": total_tables,
            "total_paragraphs": len(paragraphs),
            "placeholders": placeholders,
            "docx_geometry": {
                "table_row_cell_counts": table_row_cell_counts,
                "paragraph_count": len(paragraphs),
            },
        }

        return RawTemplateRecipeCandidate(
            template_path=template_path,
            profile_id=profile_id,
            fingerprint=fingerprint,
            roster_candidate=roster_candidate,
            header_candidates=header_candidates,
            signature_candidates=signature_candidates,
            collisions=collisions,
            metadata=metadata,
        )

    def _detect_placeholders(self, body: Any) -> List[Dict[str, Any]]:
        """Finds explicit template tags like {{INSTRUCTOR}} or [INSTRUCTOR]."""
        placeholders = []
        full_text = "".join(t.text or "" for t in body.iter(w("t")))
        placeholder_pattern = re.compile(r"\{\{([A-Z0-9_]+)\}\}|\[([A-Z0-9_]+)\]")
        matches = placeholder_pattern.findall(full_text)

        tag_mapping = {
            "INSTRUCTOR": "instructor",
            "FACULTY": "instructor",
            "COURSE": "course_section",
            "COURSE_SECTION": "course_section",
            "SECTION": "course_section",
            "SCHEDULE_CODE": "schedule_code",
            "SCHED_CODE": "schedule_code",
            "SUBJECT": "subject",
            "SUBJECT_CODE": "subject_code",
            "SUBJECT_TITLE": "subject_title",
            "TIME_DAYS_ROOM": "time_days_room",
            "TIME_ROOM": "time_days_room",
            "SEMESTER": "semester",
            "SEMESTER_AY": "semester_ay",
            "SCHOOL_YEAR": "school_year",
            "ACADEMIC_YEAR": "school_year",
            "COLLEGE": "college",
            "DEPARTMENT": "department",
            "PROGRAM": "program",
            "DATE": "date",
            "PERIOD": "period",
            "UNITS": "units",
        }

        seen = set()
        for match in matches:
            tag = match[0] or match[1]
            if tag in tag_mapping and tag not in seen:
                seen.add(tag)
                placeholders.append({
                    "field": tag_mapping[tag],
                    "tag": f"{{{{{tag}}}}}" if match[0] else f"[{tag}]",
                    "raw_token": tag,
                })
        return placeholders

    def _detect_roster_table(self, tbl, table_index: int) -> Optional[Dict[str, Any]]:
        """
        Detects whether a table is a student roster table by examining header columns.
        """
        rows = tbl.findall(w("tr"))
        if len(rows) < 2:
            return None

        best_score = 0
        best_candidate = None

        # Check rows 0, 1, and 2 for header labels
        for r_idx in range(min(3, len(rows))):
            cells = rows[r_idx].findall(w("tc"))
            if not cells:
                continue

            name_col = None
            id_col = None
            index_col = None
            sig_col = None
            score = 0

            for c_idx, cell in enumerate(cells):
                txt = SemanticRegistry.normalize_text(get_full_text(cell)).lower()
                cleaned = re.sub(r"[^a-z0-9#\.\s]", " ", txt).strip()
                words = set(cleaned.split())

                is_id = any(tok in cleaned for tok in ("student number", "id number", "stud no", "student no", "stud. no", "student id", "studentnumber", "id no", "id.", "numero", "lrn")) or bool(words & {"lrn", "numero", "id"})
                is_index = any(tok in words for tok in ("no", "no.", "#", "item", "bilang", "index")) or cleaned.startswith("#") or cleaned.startswith("no")
                is_sig = any(tok in cleaned for tok in ("signature", "lagda", "sign", "remarks", "grade", "initial", "status", "action taken", "concern", "pirma"))
                is_name = any(tok in cleaned for tok in ("name", "pangalan", "student's name", "name of student", "name of students")) or ("student" in words and not is_id)

                if id_col is None and is_id:
                    id_col = c_idx
                    score += 3
                elif index_col is None and is_index:
                    index_col = c_idx
                    score += 1
                elif sig_col is None and is_sig:
                    sig_col = c_idx
                    score += 1
                elif name_col is None and is_name:
                    name_col = c_idx
                    score += 3

            if name_col is not None and (id_col is not None or sig_col is not None):
                # Scan subsequent rows for secondary header rows (multi-row headers)
                header_row_count = 1
                curr_r = r_idx + 1
                while curr_r < min(r_idx + 4, len(rows)):
                    cand_tr = rows[curr_r]
                    cand_cells = cand_tr.findall(w("tc"))
                    if not cand_cells:
                        break

                    if self._is_secondary_header_row(
                        cand_tr, cand_cells, name_col, id_col, index_col
                    ):
                        header_row_count += 1
                        curr_r += 1
                    else:
                        break

                first_data_row = r_idx + header_row_count
                if first_data_row >= len(rows):
                    # Table has no physical data rows after headers; cannot serve as a roster table
                    continue
                total_cols = len(rows[first_data_row].findall(w("tc")))
                capacity_limit = len(rows) - first_data_row

                if score > best_score:
                    best_score = score
                    best_candidate = {
                        "table_index": table_index,
                        "header_row_index": r_idx,
                        "header_row_count": header_row_count,
                        "template_row_index": first_data_row,
                        "first_data_row": first_data_row,
                        "first_data_row_index": first_data_row,
                        "name_col": name_col,
                        "id_col": id_col,
                        "index_col": index_col,
                        "signature_col": sig_col,
                        "total_cols": total_cols,
                        "total_rows": len(rows),
                        "capacity_limit": capacity_limit,
                        "has_split_names": False,
                        "score": best_score,
                    }

        return best_candidate

    def _is_secondary_header_row(
        self,
        tr_elem: Any,
        cells: List[Any],
        name_col: int,
        id_col: Optional[int],
        index_col: Optional[int],
    ) -> bool:
        """
        Determines whether a subsequent row is a secondary/continuation header row
        or a template data row.

        Signals combined:
          1. Explicit Word header marker (<w:tblHeader/>).
          2. Student data presence (rejects header classification).
          3. Blank row presence (blank row without <w:tblHeader/> is treated as data template row, NOT header).
          4. Semantic/header continuation markers in non-name/id columns or subheader keywords.
        """
        # 1. Explicit Word header element
        trPr = tr_elem.find(w("trPr"))
        has_tbl_header = False
        if trPr is not None and trPr.find(w("tblHeader")) is not None:
            has_tbl_header = True

        # Extract normalized cell text
        cell_texts = [SemanticRegistry.normalize_text(get_full_text(c)) for c in cells]

        # Check if row is completely blank (all cells empty or whitespace)
        all_blank = all(not t for t in cell_texts)
        if all_blank:
            # A blank row without explicit Word tblHeader is a data template row, NOT a header
            return has_tbl_header

        # 2. Check for student data in id_col or name_col
        # Check student ID patterns in id_col
        if id_col is not None and id_col < len(cell_texts):
            id_txt = cell_texts[id_col].strip()
            clean_digits = re.sub(r"[^\d]", "", id_txt)
            if len(clean_digits) >= 5:
                return False  # Real student ID present, definitely data row

        # Check real student name in name_col
        if name_col < len(cell_texts):
            name_txt = cell_texts[name_col].strip()
            if name_txt:
                lower_name = name_txt.lower()
                is_header_kw = any(kw in lower_name for kw in ("name", "student", "signature", "date", "no.", "no"))
                if "," in name_txt and not is_header_kw and len(name_txt) > 3:
                    return False  # Real student name with comma, definitely data row

        # Check index_col for numbered template row (e.g. '1', '2')
        if index_col is not None and index_col < len(cell_texts):
            idx_txt = cell_texts[index_col].strip()
            if idx_txt.isdigit() and int(idx_txt) >= 1:
                # Check if any other cell has distinct subheader keywords
                has_header_keywords = False
                for c_idx, txt in enumerate(cell_texts):
                    if c_idx == index_col:
                        continue
                    clean_w = set(re.sub(r"[^a-z0-9]", " ", txt.lower()).split())
                    if clean_w & {"date", "day", "days", "week", "weel", "time", "score", "quiz", "exam", "lab", "lec", "remarks"}:
                        has_header_keywords = True
                        break
                if not has_header_keywords and not has_tbl_header:
                    return False  # Pre-numbered row template (e.g. '1', '', '', '')

        # If explicit Word header marker is present, trust it
        if has_tbl_header:
            return True

        # 3. Semantic subheader keywords in other columns
        SUBHEADER_KEYWORDS = {
            "date", "day", "days", "week", "weel", "time", "room", "score", "quiz",
            "exam", "lab", "lec", "lecture", "laboratory", "remarks", "remark",
            "signature", "sig", "total", "subtotal", "grade", "equivalent",
            "status", "initial", "item", "items", "unit", "units", "hours", "hrs",
            "topic", "criteria", "percentage", "m", "f", "mon", "tue", "wed", "thu", "fri", "sat", "sun",
            "first", "last", "middle", "suffix", "period", "month", "year"
        }

        # Check non-name/id cells (or sub-divided header cells) for subheader keywords
        for c_idx, txt in enumerate(cell_texts):
            clean_txt = re.sub(r"[^a-z0-9]", " ", txt.lower()).strip()
            words = set(clean_txt.split())
            if words & SUBHEADER_KEYWORDS:
                return True

        # Check vertical merge continuation (vMerge without val="restart")
        has_vmerge_continuation = False
        for c in cells:
            tcPr = c.find(w("tcPr"))
            if tcPr is not None:
                vm = tcPr.find(w("vMerge"))
                if vm is not None:
                    val = vm.attrib.get(w("val"), "")
                    if val != "restart":
                        has_vmerge_continuation = True
                        break

        if has_vmerge_continuation and not any(len(re.sub(r"[^\d]", "", t)) >= 5 for t in cell_texts):
            return True

        return False

    def _suggest_title_and_suffix(self, template_path: str, paragraphs: List[Any]) -> Tuple[str, str]:
        """Infers document title and file suffix from filename or top paragraphs."""
        basename = os.path.splitext(os.path.basename(template_path))[0]
        clean_name = re.sub(r"^(?:template_|cvsu_|form_)", "", basename, flags=re.IGNORECASE)
        clean_name = re.sub(r"[-_\s]+", "_", clean_name).strip("_")

        title = clean_name.replace("_", " ").title()
        suffix = clean_name.upper()

        for p in paragraphs[:5]:
            txt = SemanticRegistry.normalize_text(get_full_text(p))
            if 5 < len(txt) < 80:
                upper = txt.upper()
                if any(k in upper for k in ("FORM", "ACCEPTANCE", "DISCUSSION", "ACKNOWLEDGMENT", "RETURNS", "RECORD", "LOG", "REPORT")):
                    title = txt.title()
                    sub_suffix = re.sub(r"[^A-Za-z0-9\s]", "", upper)
                    tokens = [t for t in sub_suffix.split() if t not in ("OF", "THE", "AND", "FOR", "CVSU", "CEIT")]
                    if tokens:
                        suffix = "_".join(tokens[:4])
                    break

        return title, suffix

    def _compute_confidence(
        self,
        roster_info: Optional[Dict[str, Any]],
        header_candidates: List[Dict[str, Any]],
        placeholders: Optional[List[Dict[str, Any]]] = None,
    ) -> int:
        """Calculates discovery confidence score 0 - 100."""
        score = 0
        if roster_info:
            score += 25
            if roster_info.get("name_col") is not None:
                score += 10
            if roster_info.get("id_col") is not None:
                score += 10

        bound_fields = {c["field"] for c in header_candidates}
        if placeholders:
            for p in placeholders:
                if "field" in p:
                    bound_fields.add(p["field"])

        core_fields = {"instructor", "course_section", "schedule_code", "subject"}
        found_core = bound_fields.intersection(core_fields)
        score += min(len(found_core) * 10, 40)

        if "semester_ay" in bound_fields or "time_days_room" in bound_fields:
            score += 5
        score += 10  # baseline structure score

        return min(max(score, 0), 100)



# ═══════════════════════════════════════════════════════════════════════════════
# XlsxTemplateInspector
# ═══════════════════════════════════════════════════════════════════════════════

class XlsxTemplateInspector:
    """
    Analyzes Excel (.xlsx) grading sheet templates and emits candidate observations.
    CRITICAL: inspect() returns RawTemplateRecipeCandidate ONLY.
    XlsxTemplateInspector NEVER constructs or returns ValidatedTemplateRecipe.
    RecipeValidator owns final validation and construction authority.
    """

    def inspect(
        self,
        template_path: str,
        profile_id: str = "grade_sheet_xlsx",
    ) -> RawTemplateRecipeCandidate:
        if not os.path.exists(template_path):
            raise FileNotFoundError(f"Template file not found: {template_path}")

        with open(template_path, "rb") as f:
            fingerprint = hashlib.sha256(f.read()).hexdigest()

        wb = openpyxl.load_workbook(template_path, data_only=False)
        sheet_names = wb.sheetnames
        sheet_map = {s.lower().strip(): s for s in sheet_names}

        # ── 1. Structural Worksheet Analysis ──────────────────────────────
        def inspect_sheet_roster(ws) -> Optional[Dict[str, Any]]:
            header_row = None
            index_col = None
            name_col = None
            id_col = None

            for r in range(1, min(ws.max_row + 1, 25)):
                for c in range(1, min(ws.max_column + 1, 30)):
                    c_val = ws.cell(r, c).value
                    if not c_val or not isinstance(c_val, str):
                        continue
                    norm_c = c_val.strip().lower()
                    if norm_c in ("#", "no.", "no", "item", "bilang"):
                        index_col = c
                        header_row = r
                    elif "student name" in norm_c or "name of student" in norm_c or "surname" in norm_c or norm_c in ("name", "pangalan", "full name", "names of students"):
                        name_col = c
                    elif "student number" in norm_c or "student no" in norm_c or "id number" in norm_c or "id no" in norm_c or norm_c in ("id", "student id", "lrn", "id no."):
                        id_col = c

                if header_row is not None and name_col is not None and id_col is not None and index_col is not None:
                    break

            if header_row is not None and name_col is not None and id_col is not None and index_col is not None:
                first_data_row = None
                capacity_limit = 0
                check_col = index_col
                for r in range(header_row + 1, ws.max_row + 1):
                    v = ws.cell(r, check_col).value
                    if v == 1 or str(v).strip() == "1":
                        first_data_row = r
                        break

                if first_data_row is not None:
                    expected_num = 1
                    curr_r = first_data_row
                    while curr_r <= ws.max_row:
                        v = ws.cell(curr_r, check_col).value
                        if v == expected_num or str(v).strip() == str(expected_num):
                            capacity_limit += 1
                            expected_num += 1
                            curr_r += 1
                        else:
                            break

                if first_data_row is not None and capacity_limit > 0:
                    return {
                        "table_index": 0,
                        "worksheet_name": ws.title,
                        "first_data_row_index": first_data_row,
                        "name_col": name_col,
                        "id_col": id_col,
                        "index_col": index_col,
                        "capacity_limit": capacity_limit,
                        "has_split_names": False,
                    }
            return None

        # Discover all candidate roster sheets
        roster_candidates_by_sheet: Dict[str, Dict[str, Any]] = {}
        for s_name in sheet_names:
            r_cand = inspect_sheet_roster(wb[s_name])
            if r_cand is not None:
                roster_candidates_by_sheet[s_name] = r_cand

        # Select primary roster sheet based on structural class metadata density in header rows
        def score_roster_sheet(s_name: str) -> int:
            ws_cand = wb[s_name]
            score = 0
            # Metadata presence in rows 1-10 is the primary structural discriminator
            for r in range(1, min(ws_cand.max_row + 1, 10)):
                for c in range(1, min(ws_cand.max_column + 1, 15)):
                    val = ws_cand.cell(r, c).value
                    if val and isinstance(val, str):
                        m = SemanticRegistry.match_metadata_candidate(val.strip())
                        if m:
                            score += 25
            return score

        primary_roster_sheet: Optional[str] = None
        if roster_candidates_by_sheet:
            roster_scores = {s: score_roster_sheet(s) for s in roster_candidates_by_sheet}
            max_r_score = max(roster_scores.values())
            top_rosters = [s for s, sc in roster_scores.items() if sc == max_r_score]
            if len(top_rosters) > 1:
                # Multiple candidate roster worksheets with equal structural evidence
                raise AmbiguousTemplateError(
                    f"Grade sheet template '{os.path.basename(template_path)}' contains multiple ambiguous candidate roster worksheets with equal structural evidence: {top_rosters}"
                )
            primary_roster_sheet = top_rosters[0]

        if not primary_roster_sheet:
            raise TemplateError(
                f"Grade sheet template '{os.path.basename(template_path)}' is missing required student roster worksheet."
            )

        # Select summary rating sheet structurally by scoring rating table and institutional banners
        def score_summary_sheet(s_name: str) -> int:
            if s_name == primary_roster_sheet:
                return -100
            ws_s = wb[s_name]
            score = 0
            has_student_id = False
            has_rating_col = False
            has_banner = False
            for r in range(1, min(ws_s.max_row + 1, 30)):
                for c in range(1, min(ws_s.max_column + 1, 25)):
                    raw_val = ws_s.cell(r, c).value
                    if not raw_val or not isinstance(raw_val, str):
                        continue
                    t = raw_val.strip().lower()
                    if len(t) > 100:
                        continue
                    if any(k in t for k in ("college", "faculty", "department", "university", "republic", "official grades")):
                        has_banner = True
                    # Table headers must be concise column labels, not prose instructions
                    if len(t) <= 30:
                        if t in ("student number", "id number", "student no", "student no.", "stud no", "lrn", "id", "id."):
                            has_student_id = True
                        elif t in ("grade", "rating", "mark", "final grade", "final rating", "semestral grade", "numerical rating", "remarks"):
                            has_rating_col = True

            if has_student_id and has_rating_col:
                score += 100
            elif has_rating_col:
                score += 50
            if has_banner:
                score += 25
            return score

        candidate_summary_sheets = [s for s in sheet_names if s != primary_roster_sheet]
        summary_sheet: Optional[str] = None
        if candidate_summary_sheets:
            summary_scores = {s: score_summary_sheet(s) for s in candidate_summary_sheets}
            max_s_score = max(summary_scores.values()) if summary_scores else 0
            if max_s_score > 0:
                top_summaries = [s for s, sc in summary_scores.items() if sc == max_s_score]
                if len(top_summaries) > 1:
                    raise AmbiguousTemplateError(
                        f"Grade sheet template '{os.path.basename(template_path)}' contains multiple ambiguous summary rating worksheets with equal structural evidence: {top_summaries}"
                    )
                summary_sheet = top_summaries[0]

        if not summary_sheet:
            raise TemplateError(
                f"Grade sheet template '{os.path.basename(template_path)}' is missing required 'Grading Sheet' worksheet or structural summary rating sheet."
            )

        # Detect secondary / laboratory component structurally from remaining assessment sheets
        remaining_assessment_sheets = [
            s for s in sheet_names
            if s in roster_candidates_by_sheet and s != primary_roster_sheet and s != summary_sheet
        ]
        has_lab = bool(len(remaining_assessment_sheets) >= 1)
        lab_sheet: Optional[str] = None
        has_consolidated = False
        con_sheet: Optional[str] = None

        if len(remaining_assessment_sheets) == 1:
            lab_sheet = remaining_assessment_sheets[0]
            has_lab = True
            has_consolidated = False
        elif len(remaining_assessment_sheets) == 2:
            s1, s2 = remaining_assessment_sheets[0], remaining_assessment_sheets[1]
            ws1, ws2 = wb[s1], wb[s2]

            def count_cross_sheet_refs(ws) -> int:
                refs = 0
                for row in ws.iter_rows(values_only=True):
                    for val in row:
                        if isinstance(val, str) and val.startswith("=") and "!" in val:
                            refs += 1
                return refs

            refs1 = count_cross_sheet_refs(ws1)
            refs2 = count_cross_sheet_refs(ws2)
            cols1 = ws1.max_column
            cols2 = ws2.max_column
            rows1 = ws1.max_row
            rows2 = ws2.max_row

            # Determine laboratory vs consolidated based on unique physical lineage:
            # Construct directed dependency relation: A -> B meaning worksheet A
            # physically contains formulas referencing worksheet B.
            def get_referenced_sheets(ws_source) -> Tuple[set, bool]:
                referenced = set()
                is_reliable = True
                for row in ws_source.iter_rows(values_only=True):
                    for val in row:
                        if isinstance(val, str) and val.startswith("="):
                            sheet_refs = extract_referenced_sheets(val)
                            if not getattr(sheet_refs, "is_reliable", True):
                                is_reliable = False
                            referenced.update(sheet_refs)
                return referenced, is_reliable

            s1_referenced, s1_reliable = get_referenced_sheets(ws1)
            s2_referenced, s2_reliable = get_referenced_sheets(ws2)

            # Harden fallback formula parsing: if candidate secondary sheets contain
            # unparseable, corrupted, or indeterminate formula references, fail closed.
            if not s1_reliable or not s2_reliable:
                unreliable = [s for s, r in [(s1, s1_reliable), (s2, s2_reliable)] if not r]
                raise AmbiguousTemplateError(
                    f"Grade sheet template '{os.path.basename(template_path)}' contains secondary assessment worksheets with unparseable or indeterminate formula lineage: {unreliable}"
                )

            non_summary_sheets = {s.lower() for s in sheet_names if s.lower() != summary_sheet.lower()}
            s1_non_summary_refs = {r.lower() for r in s1_referenced if r.lower() in non_summary_sheets and r.lower() != s1.lower()}
            s2_non_summary_refs = {r.lower() for r in s2_referenced if r.lower() in non_summary_sheets and r.lower() != s2.lower()}

            s1_refs_s2 = any(r.lower() == s2.lower() for r in s1_referenced)
            s2_refs_s1 = any(r.lower() == s1.lower() for r in s2_referenced)

            s1_has_agg = has_aggregation_structure(ws1, s2, set(sheet_names))
            s2_has_agg = has_aggregation_structure(ws2, s1, set(sheet_names))

            # Case D: s1 -> s2 and s2 -> s1 (competing/cyclic lineage)
            if s1_refs_s2 and s2_refs_s1:
                raise AmbiguousTemplateError(
                    f"Grade sheet template '{os.path.basename(template_path)}' contains multiple secondary assessment worksheets with competing/cyclic formula lineage referencing each other: {[s1, s2]}"
                )
            # Case A: s1 -> s2 and s2 !-> s1 (s1 candidate for Consolidated)
            elif s1_refs_s2 and not s2_refs_s1:
                # Directed lineage alone is not automatic semantic truth:
                # To resolve s1 as Consolidated, s1 must exhibit role-consistent
                # physical aggregation topology:
                # 1. It must reference multiple instructional/student components (>= 2 non-summary sheets)
                # 2. It must physically demonstrate an aggregation structure combining components
                if len(s1_non_summary_refs) < 2 or not s1_has_agg:
                    raise AmbiguousTemplateError(
                        f"Grade sheet template '{os.path.basename(template_path)}' contains secondary assessment worksheets with directed lineage ({s1} -> {s2}) but lacking role-consistent physical aggregation topology: {[s1, s2]}"
                    )
                con_sheet = s1
                lab_sheet = s2
                has_lab = True
                has_consolidated = True
            # Case B: s2 -> s1 and s1 !-> s2 (s2 candidate for Consolidated)
            elif s2_refs_s1 and not s1_refs_s2:
                # Symmetrically, s2 must exhibit role-consistent physical aggregation topology:
                if len(s2_non_summary_refs) < 2 or not s2_has_agg:
                    raise AmbiguousTemplateError(
                        f"Grade sheet template '{os.path.basename(template_path)}' contains secondary assessment worksheets with directed lineage ({s2} -> {s1}) but lacking role-consistent physical aggregation topology: {[s1, s2]}"
                    )
                con_sheet = s2
                lab_sheet = s1
                has_lab = True
                has_consolidated = True
            # Case C: s1 !-> s2 and s2 !-> s1 (no unique lineage)
            else:
                if cols1 == cols2 and rows1 == rows2 and refs1 == refs2:
                    raise AmbiguousTemplateError(
                        f"Grade sheet template '{os.path.basename(template_path)}' contains multiple secondary assessment worksheets with identical structural dimensions and evidence: {[s1, s2]}"
                    )
                else:
                    raise AmbiguousTemplateError(
                        f"Grade sheet template '{os.path.basename(template_path)}' contains multiple secondary assessment worksheets with different physical dimensions but no unique structural lineage distinguishing laboratory from consolidated roles: {[s1, s2]}"
                    )
        elif len(remaining_assessment_sheets) > 2:
            raise AmbiguousTemplateError(
                f"Grade sheet template '{os.path.basename(template_path)}' contains {len(remaining_assessment_sheets)} secondary assessment worksheets, exceeding supported dual-component capacity: {remaining_assessment_sheets}"
            )


        ws_lec = wb[primary_roster_sheet]
        header_candidates: List[Dict[str, Any]] = []
        signature_candidates: List[Dict[str, Any]] = []
        roster_candidate: Optional[Dict[str, Any]] = roster_candidates_by_sheet.get(primary_roster_sheet)

        # ── 2. Header Metadata Discovery in Primary Roster Sheet ──────────
        lec_merges = list(ws_lec.merged_cells.ranges)

        def find_target_for_label(row_idx: int, col_idx: int) -> str:
            row_merges = [
                m for m in lec_merges
                if m.min_row == row_idx and m.min_col > col_idx
            ]
            if row_merges:
                closest_merge = min(row_merges, key=lambda m: m.min_col)
                return f"{get_column_letter(closest_merge.min_col)}{closest_merge.min_row}"
            return f"{get_column_letter(col_idx + 1)}{row_idx}"

        for r in range(1, ws_lec.max_row + 1):
            for c in range(1, ws_lec.max_column + 1):
                val = ws_lec.cell(r, c).value
                if not val or not isinstance(val, str):
                    continue
                text = val.strip()
                if not text:
                    continue

                # Signature labels are discovered separately from metadata bindings.
                if "INSTRUCTOR" in text.upper() and r > 6:
                    continue

                matches = SemanticRegistry.match_metadata_candidate(text)
                if matches:
                    matched_field, conf = matches[0]
                    target_coord = find_target_for_label(r, c)
                    header_candidates.append({
                        "cell_type": "xlsx_cell",
                        "target": target_coord,
                        "field": matched_field,
                        "confidence": conf,
                        "pattern": text,
                        "label_cell": ws_lec.cell(r, c).coordinate,
                        "is_signature_region": False,
                    })

        # ── 3. Institutional College Banner Discovery ─────────────────────
        if summary_sheet and summary_sheet in wb.sheetnames:
            ws_grd = wb[summary_sheet]
            for r in range(1, ws_grd.max_row + 1):
                for c in range(1, ws_grd.max_column + 1):
                    val = ws_grd.cell(r, c).value
                    if val and isinstance(val, str) and "COLLEGE OF" in val.upper():
                        header_candidates.append({
                            "cell_type": "xlsx_cell",
                            "target": f"{summary_sheet}!{ws_grd.cell(r, c).coordinate}",
                            "field": "college",
                            "confidence": 1.0,
                            "pattern": val.strip(),
                            "label_cell": None,
                            "is_signature_region": False,
                        })
                        break

        # ── 4. Scope-Aware Signature Box Discovery ────────────────────────
        def find_signature_target(ws, sheet_display_name: str, scope_name: str) -> Optional[Dict[str, Any]]:
            label_rng = None
            for rng in ws.merged_cells.ranges:
                top_val = ws.cell(rng.min_row, rng.min_col).value
                if top_val and isinstance(top_val, str) and any(k in top_val.upper() for k in ("INSTRUCTOR", "PROFESSOR", "FACULTY", "TEACHER")):
                    label_rng = rng
                    break

            if not label_rng:
                # No structural merged label found; do NOT invent a target
                return None

            for rng in ws.merged_cells.ranges:
                if rng.min_col == label_rng.min_col and rng.max_col == label_rng.max_col:
                    if rng.max_row == label_rng.min_row - 1:
                        target_cell = f"{get_column_letter(rng.min_col)}{rng.min_row}"
                        target_coord = target_cell if sheet_display_name == primary_roster_sheet else f"{sheet_display_name}!{target_cell}"
                        return {
                            "role": "instructor" if scope_name == "lecture" else f"instructor:{scope_name}",
                            "scope": scope_name,
                            "target": target_coord,
                            "confidence": 1.0,
                            "derivation_evidence": "structural_merged_box_above_label",
                        }

            # If no structural merged target box directly above label, fail discovery rather than guessing an offset
            return None

        lec_sig = find_signature_target(ws_lec, primary_roster_sheet, "lecture")
        if lec_sig:
            signature_candidates.append(lec_sig)
            signature_candidates.append({
                "role": ROLE_INSTRUCTOR_SIGNATURE,
                "scope": "lecture",
                "target": lec_sig["target"],
                "confidence": lec_sig["confidence"],
                "derivation_evidence": lec_sig["derivation_evidence"],
            })

        if lab_sheet and lab_sheet in wb.sheetnames:
            lab_sig = find_signature_target(wb[lab_sheet], lab_sheet, "laboratory")
            if lab_sig:
                signature_candidates.append(lab_sig)

        if con_sheet and con_sheet in wb.sheetnames:
            con_sig = find_signature_target(wb[con_sheet], con_sheet, "consolidated")
            if con_sig:
                signature_candidates.append(con_sig)

        xlsx_geometry = {
            "worksheets": {
                s: {
                    "max_row": wb[s].max_row,
                    "max_column": wb[s].max_column,
                }
                for s in wb.sheetnames
            },
            "roster_sheet": primary_roster_sheet,
            "summary_sheet": summary_sheet,
            "lab_sheet": lab_sheet,
            "con_sheet": con_sheet,
        }

        return RawTemplateRecipeCandidate(
            template_path=template_path,
            profile_id=profile_id,
            fingerprint=fingerprint,
            roster_candidate=roster_candidate,
            header_candidates=header_candidates,
            signature_candidates=signature_candidates,
            collisions=[],
            metadata={
                "sheet_names": sheet_names,
                "roster_sheet": primary_roster_sheet,
                "summary_sheet": summary_sheet,
                "lab_sheet": lab_sheet,
                "con_sheet": con_sheet,
                "has_lab": has_lab,
                "has_consolidated": has_consolidated,
                "xlsx_geometry": xlsx_geometry,
            },
        )


# ═══════════════════════════════════════════════════════════════════════════════
# Backward-Compatibility Facade: TemplateInspector
# ═══════════════════════════════════════════════════════════════════════════════

class TemplateInspector:
    """
    Backward-compatibility facade matching the legacy TemplateInspector interface.
    Internally delegates to DocxTemplateInspector / XlsxTemplateInspector -> RecipeValidator.
    """

    CANONICAL_ACADEMIC_TEMPLATES = {
        "template_syllabus.docx",
        "template_exam_midterm.docx",
        "template_exam_finals.docx",
        "template_tos_midterm.docx",
        "template_tos_finals.docx",
        "Midterm-Grade-Discussion_LATEST.docx",
        "Final-Grade-Discussion_LATEST.docx",
    }

    def __init__(self):
        self._docx_inspector = DocxTemplateInspector()
        self._xlsx_inspector = XlsxTemplateInspector()

    def inspect(self, template_path: str) -> Dict[str, Any]:
        """Legacy inspect() method returning a dictionary representation."""
        if template_path.lower().endswith((".xlsx", ".xls")):
            return self.inspect_xlsx(template_path)
        return self.inspect_docx(template_path)

    def inspect_xlsx(self, template_path: str, profile_id: str = "grade_sheet_xlsx") -> Dict[str, Any]:
        candidate = self._xlsx_inspector.inspect(template_path, profile_id=profile_id)
        profile = PROFILE_REGISTRY.get(profile_id, PROFILE_GRADE_SHEET_XLSX)
        validated_recipe = RecipeValidator.validate(candidate, profile)
        return validated_recipe.to_dict()

    def inspect_xlsx_recipe(
        self,
        template_path: str,
        profile_id: str = "grade_sheet_xlsx",
    ) -> ValidatedTemplateRecipe:
        """Public API returning an authoritative ValidatedTemplateRecipe for XLSX."""
        candidate = self._xlsx_inspector.inspect(template_path, profile_id=profile_id)
        profile = PROFILE_REGISTRY.get(profile_id, PROFILE_GRADE_SHEET_XLSX)
        return RecipeValidator.validate(candidate, profile)

    def inspect_docx(self, template_path: str, profile_id: Optional[str] = None) -> Dict[str, Any]:
        """
        Inspects a Word (.docx) file and returns a dictionary matching the legacy API format.
        Internally runs DocxTemplateInspector -> RecipeValidator -> legacy dict format.
        """
        base = os.path.basename(template_path)
        if profile_id is None:
            profile_id = "academic_docx" if base in self.CANONICAL_ACADEMIC_TEMPLATES else "custom_docx"

        candidate = self._docx_inspector.inspect(template_path, profile_id=profile_id)
        from modules.models.recipe import PROFILE_CUSTOM_DOCX
        profile = PROFILE_REGISTRY.get(profile_id, PROFILE_ACADEMIC_DOCX if profile_id == "academic_docx" else PROFILE_CUSTOM_DOCX)
        validated_recipe = RecipeValidator.validate(candidate, profile)

        # Build legacy format expected by existing tests and frontend
        roster_dict = None
        if validated_recipe.roster_binding:
            rb = validated_recipe.roster_binding
            roster_dict = {
                "table_index": rb.table_index,
                "first_data_row": rb.first_data_row_index,
                "first_data_row_index": rb.first_data_row_index,
                "name_col": rb.name_col,
                "id_col": rb.id_col,
                "index_col": candidate.roster_candidate.get("index_col") if candidate.roster_candidate else None,
                "signature_col": candidate.roster_candidate.get("signature_col") if candidate.roster_candidate else None,
                "total_cols": candidate.roster_candidate.get("total_cols") if candidate.roster_candidate else None,
                "total_rows": candidate.roster_candidate.get("total_rows") if candidate.roster_candidate else None,
                "capacity_limit": rb.capacity_limit,
                "has_split_names": rb.has_split_names,
            }

        header_bindings_list = []
        detected_fields = []
        for field_name, b in validated_recipe.header_bindings.items():
            t = b.target
            if b.cell_type == "docx_table" and isinstance(t, (tuple, list)) and len(t) == 3:
                legacy_b = {
                    "field": field_name,
                    "type": "table_cell",
                    "table_index": t[0],
                    "row_index": t[1],
                    "cell_index": t[2],
                    "shrink_threshold": b.shrink_threshold,
                    "shrink_sz": b.shrink_sz,
                }
            elif b.cell_type == "docx_paragraph":
                legacy_b = {
                    "field": field_name,
                    "type": "paragraph_colon",
                    "para_index": t,
                    "shrink_threshold": b.shrink_threshold,
                    "shrink_sz": b.shrink_sz,
                }
            else:
                legacy_b = b.to_dict()
            header_bindings_list.append(legacy_b)
            detected_fields.append(field_name)

        confidence = candidate.metadata.get("confidence", 95)
        placeholders = candidate.metadata.get("placeholders", [])

        return {
            "template_path": template_path,
            "title": candidate.metadata.get("title", ""),
            "suffix": candidate.metadata.get("suffix", ""),
            "confidence": confidence,
            "summary": {
                "roster_table_found": roster_dict is not None,
                "detected_fields": detected_fields,
            },
            "roster_table": roster_dict,
            "header_bindings": header_bindings_list,
            "signature_bindings": [s.to_dict() for s in validated_recipe.signature_bindings.values()],
            "placeholders": placeholders,
            "schema_version": validated_recipe.schema_version,
            "fingerprint": validated_recipe.fingerprint,
            "docx_geometry": candidate.metadata.get("docx_geometry"),
        }

    def inspect_docx_recipe(
        self,
        template_path: str,
        profile_id: str = "academic_docx",
    ) -> ValidatedTemplateRecipe:
        """
        Public API returning an authoritative ValidatedTemplateRecipe.
        """
        candidate = self._docx_inspector.inspect(template_path, profile_id=profile_id)
        profile = PROFILE_REGISTRY.get(profile_id, PROFILE_ACADEMIC_DOCX)
        return RecipeValidator.validate(candidate, profile)


# ═══════════════════════════════════════════════════════════════════════════════
# AttendanceTemplateInspector
# ═══════════════════════════════════════════════════════════════════════════════

class AttendanceTemplateInspector:
    """
    Analyzes Word (.docx) templates specifically for Attendance sheets
    and emits raw candidate observations.
    CRITICAL: inspect() returns RawAttendanceTemplateRecipeCandidate.
    Inspectors NEVER construct or return ValidatedAttendanceTemplateRecipe.
    RecipeValidator owns validation authority.
    """

    INFO_PATTERNS = {
        "course_code_title": re.compile(r"(?i)\b(course\s*(?:code)?(?:\s*&)?\s*title|subject|course)\b"),
        "month_year": re.compile(r"(?i)\b(month\s*(?:&)?\s*year|for\s+the\s+month|month)\b"),
        "class_schedule": re.compile(r"(?i)\b(class\s*schedule|schedule|time)\b"),
        "semester_ay": re.compile(r"(?i)\b(semester\s*(?:&)?\s*ay|semester|a\.?y\.?|academic\s*year)\b"),
        "room_assignment": re.compile(r"(?i)\b(room\s*assignment|room)\b"),
        "instructor": re.compile(r"(?i)\b(name\s*of\s*instructor|instructor|professor|faculty)\b"),
    }

    def inspect(
        self,
        template_path: str,
        profile_id: str = "attendance_docx",
    ) -> RawAttendanceTemplateRecipeCandidate:
        """
        Inspects an Attendance Word (.docx) template file and produces raw candidate observations.
        """
        if not os.path.exists(template_path):
            raise FileNotFoundError(f"Template file not found: {template_path}")

        h = hashlib.sha256()
        with open(template_path, "rb") as f:
            for chunk in iter(lambda: f.read(65536), b""):
                h.update(chunk)
        fingerprint = h.hexdigest()

        zin, root, body = load_docx(template_path)
        zin.close()
        tables = body.findall(w("tbl"))

        info_candidates: List[Dict[str, Any]] = []
        matrix_candidates: List[Dict[str, Any]] = []
        collisions: List[Dict[str, Any]] = []

        for tbl_idx, tbl in enumerate(tables):
            # 1. Evaluate for info table
            info_score, bindings, row_cell_counts = self._evaluate_info_table(tbl, tbl_idx)
            if info_score >= 3:
                info_candidates.append({
                    "table_index": tbl_idx,
                    "score": info_score,
                    "bindings": bindings,
                    "row_cell_counts": row_cell_counts,
                })

            # 2. Evaluate for matrix table
            matrix_cand = self._evaluate_matrix_table(tbl, tbl_idx)
            if matrix_cand is not None:
                matrix_candidates.append(matrix_cand)

        selected_info = None
        if len(info_candidates) == 1:
            selected_info = info_candidates[0]
        elif len(info_candidates) > 1:
            info_candidates.sort(key=lambda x: x["score"], reverse=True)
            if info_candidates[0]["score"] == info_candidates[1]["score"]:
                collisions.append({
                    "type": "ambiguous_info_table",
                    "candidates": [c["table_index"] for c in info_candidates],
                })
            else:
                selected_info = info_candidates[0]

        selected_matrix = None
        if len(matrix_candidates) == 1:
            selected_matrix = matrix_candidates[0]
        elif len(matrix_candidates) > 1:
            collisions.append({
                "type": "ambiguous_matrix_table",
                "candidates": [c["table_index"] for c in matrix_candidates],
            })

        if selected_info and selected_matrix and selected_info["table_index"] == selected_matrix["table_index"]:
            collisions.append({
                "type": "table_identity_collision",
                "table_index": selected_info["table_index"],
            })
            selected_info = None
            selected_matrix = None

        return RawAttendanceTemplateRecipeCandidate(
            template_path=template_path,
            profile_id=profile_id,
            fingerprint=fingerprint,
            info_candidate=selected_info,
            matrix_candidate=selected_matrix,
            metadata={"output_folder": "Attendance"},
            collisions=collisions,
        )

    def _evaluate_info_table(self, tbl, tbl_idx: int) -> Tuple[int, Dict[str, Tuple[int, int]], Tuple[int, ...]]:
        rows = tbl.findall(w("tr"))
        if len(rows) < 2:
            return 0, {}, ()

        row_cell_counts = tuple(len(tr.findall(w("tc"))) for tr in rows)
        bindings: Dict[str, Tuple[int, int]] = {}
        matched_fields = set()

        for r_idx, tr in enumerate(rows):
            cells = tr.findall(w("tc"))
            for c_idx, tc in enumerate(cells):
                txt = get_full_text(tc).strip()
                if not txt:
                    continue

                for field_name, pattern in self.INFO_PATTERNS.items():
                    if field_name in matched_fields:
                        continue
                    if pattern.search(txt):
                        if c_idx + 1 < len(cells):
                            target_cell = (r_idx, c_idx + 1)
                        else:
                            target_cell = (r_idx, c_idx)
                        bindings[field_name] = target_cell
                        matched_fields.add(field_name)

        return len(matched_fields), bindings, row_cell_counts

    def _evaluate_matrix_table(self, tbl, tbl_idx: int) -> Optional[Dict[str, Any]]:
        rows = tbl.findall(w("tr"))
        if len(rows) < 3:
            return None

        no_col = None
        name_col = None
        id_col = None
        header_row0_idx = None
        header_row1_idx = None
        student_template_row_idx = None

        for r_idx, tr in enumerate(rows[:6]):
            cells = tr.findall(w("tc"))
            row_texts = [get_full_text(c).strip().upper() for c in cells]

            has_no = any(bool(re.match(r"^(\s*NO\.?|\s*#|\s*ITEM)\s*$", t)) for t in row_texts)
            has_name = any("NAME" in t or "PANGALAN" in t for t in row_texts)
            has_id = any("STUDENT NUMBER" in t or "STUDENT NO" in t or "NUMERO" in t or "ID" in t for t in row_texts)
            has_week = any("WEEK" in t for t in row_texts)

            if (has_no and has_name) or (has_week and (has_name or has_no or has_id)):
                if header_row0_idx is None:
                    header_row0_idx = r_idx
                elif header_row1_idx is None:
                    header_row1_idx = r_idx

        if header_row0_idx is None:
            return None

        if header_row1_idx is None and header_row0_idx + 1 < len(rows):
            r1_cells = rows[header_row0_idx + 1].findall(w("tc"))
            r1_texts = [get_full_text(c).strip().upper() for c in r1_cells]
            if (
                any(t in ("LB", "LC", "R") for t in r1_texts)
                or any("DATE" in t for t in r1_texts)
                or any(t.isdigit() for t in r1_texts)
            ):
                header_row1_idx = header_row0_idx + 1
            else:
                header_row1_idx = header_row0_idx

        eval_rows = [header_row0_idx]
        if header_row1_idx is not None and header_row1_idx != header_row0_idx:
            eval_rows.append(header_row1_idx)

        summary_cols = []
        summary_names = []

        for r_idx in eval_rows:
            cells = rows[r_idx].findall(w("tc"))
            for c_idx, c in enumerate(cells):
                txt = get_full_text(c).strip().upper()
                if no_col is None and re.match(r"^(\s*NO\.?|\s*#|\s*ITEM)\s*$", txt):
                    no_col = c_idx
                if name_col is None and ("NAME" in txt or "PANGALAN" in txt):
                    name_col = c_idx
                if id_col is None and ("STUDENT NUMBER" in txt or "STUDENT NO" in txt or "NUMERO" in txt or "ID" in txt):
                    id_col = c_idx
                if txt in ("LB", "LC", "R") or txt in ("REMARKS", "TOTAL", "ABSENT"):
                    if c_idx not in summary_cols:
                        summary_cols.append(c_idx)
                        summary_names.append(txt.lower())

        if name_col is None or id_col is None:
            return None

        if no_col is None:
            no_col = 0

        start_search = max(eval_rows) + 1
        student_rows_count = 0

        for r_idx in range(start_search, len(rows)):
            tr = rows[r_idx]
            cells = tr.findall(w("tc"))
            if len(cells) < 3:
                continue

            cell_texts = [get_full_text(c).strip() for c in cells]
            
            # Check for decorative / guidance rows (e.g. non-numeric text in no_col like "GUIDE", "NOTE" with empty name/id)
            no_text = cell_texts[no_col] if no_col < len(cell_texts) else ""
            name_text = cell_texts[name_col] if name_col < len(cell_texts) else ""
            id_text = cell_texts[id_col] if id_col < len(cell_texts) else ""
            is_decorative = (
                (no_text and not no_text.isdigit() and not name_text and not id_text)
                or (no_text.upper() in ("GUIDE", "NOTE", "INSTRUCTIONS", "REMARKS"))
            )
            if is_decorative and student_template_row_idx is None:
                continue

            if student_template_row_idx is None:
                # Discovered prototype student row
                student_template_row_idx = r_idx
                student_rows_count += 1
            else:
                student_rows_count += 1

        if student_template_row_idx is None:
            student_template_row_idx = max(eval_rows) + 1

        student_cols = [c for c in (no_col, name_col, id_col) if c is not None]

        # Discover actual starting column of date/session columns
        first_date_col = None
        for r_idx in eval_rows:
            cells = rows[r_idx].findall(w("tc"))
            for c_idx, c in enumerate(cells):
                txt = get_full_text(c).strip().upper()
                if "WEEK" in txt or txt.startswith("DATE") or (r_idx == header_row1_idx and txt.isdigit() and int(txt) <= 31):
                    if c_idx not in student_cols:
                        if first_date_col is None or c_idx < first_date_col:
                            first_date_col = c_idx

        if first_date_col is not None:
            date_columns_start = first_date_col
        else:
            date_columns_start = max(student_cols) + 1

        if student_template_row_idx < len(rows):
            target_row_for_len = rows[student_template_row_idx]
        elif header_row1_idx is not None and header_row1_idx < len(rows):
            target_row_for_len = rows[header_row1_idx]
        elif header_row0_idx < len(rows):
            target_row_for_len = rows[header_row0_idx]
        else:
            raise TemplateError("Matrix table has no rows to determine column count.")
        num_cols = len(target_row_for_len.findall(w("tc")))

        summary_cols_sorted = sorted(summary_cols)
        summary_count = len(summary_cols_sorted) if summary_cols_sorted else 3

        # An attendance matrix MUST physically contain date/session columns or summary columns
        if first_date_col is None and not summary_cols:
            return None

        if summary_cols_sorted:
            first_summary_col = summary_cols_sorted[0]
            template_session_capacity = first_summary_col - date_columns_start
        else:
            template_session_capacity = num_cols - date_columns_start - summary_count

        if template_session_capacity <= 0:
            return None

        if not summary_names:
            summary_names = ["lb", "lc", "r"]
            if summary_cols_sorted:
                summary_column_indices = tuple(summary_cols_sorted)
            else:
                summary_column_indices = tuple(range(max(date_columns_start + 1, num_cols - summary_count), num_cols))
        else:
            summary_column_indices = tuple(summary_cols)

        # Discover cell column spans for header row 0 to find summary header cell
        row0_tcs = rows[header_row0_idx].findall(w("tc"))
        r0_col = 0
        r0_spans = []
        for idx, tc in enumerate(row0_tcs):
            span = 1
            tcPr = tc.find(w("tcPr"))
            if tcPr is not None:
                gs = tcPr.find(w("gridSpan"))
                if gs is not None:
                    try:
                        span = int(gs.attrib.get(w("val"), "1"))
                    except ValueError:
                        span = 1
            r0_spans.append((idx, r0_col, r0_col + span))
            r0_col += span

        first_sum_col = summary_cols_sorted[0] if summary_cols_sorted else (date_columns_start + template_session_capacity)
        summary_header0_cell_col = len(row0_tcs) - 1
        for tc_idx, sc, ec in r0_spans:
            if sc <= first_sum_col < ec:
                summary_header0_cell_col = tc_idx
                break
            elif sc >= first_sum_col and tc_idx > 0:
                summary_header0_cell_col = tc_idx
                break

        week_template_cell_col = date_columns_start if date_columns_start < len(row0_tcs) else (len(row0_tcs) - 1)
        r1_tcs = rows[header_row1_idx].findall(w("tc")) if header_row1_idx is not None else row0_tcs
        date_template_cell_col = date_columns_start if date_columns_start < len(r1_tcs) else 0

        st_tcs = rows[student_template_row_idx].findall(w("tc")) if student_template_row_idx < len(rows) else r1_tcs
        student_date_template_cell_col = date_columns_start if date_columns_start < len(st_tcs) else 0

        # Discover actual/validated widths for each semantic summary column
        # Priority 1: Authoritative discovery - inspect actual prototype summary cell width from XML (<w:tcW>)
        # Priority 2: Documented legacy compatibility fallback - only used if cell lacks explicit XML width
        legacy_semantic_fallback_w_map = {"lb": 212, "lc": 208, "r": 133}
        summary_widths = []
        for name, col_idx in zip(summary_names, summary_column_indices):
            w_val = None
            # 1. Authoritative: Inspect header row 1 prototype cell
            if col_idx < len(r1_tcs):
                tcPr = r1_tcs[col_idx].find(w("tcPr"))
                if tcPr is not None:
                    tcW = tcPr.find(w("tcW"))
                    if tcW is not None:
                        try:
                            val = int(tcW.attrib.get(w("w"), "0"))
                            if val > 0:
                                w_val = val
                        except (ValueError, TypeError):
                            w_val = None
            # 2. Authoritative: If missing in row 1, check student row prototype cell
            if w_val is None and col_idx < len(st_tcs):
                tcPr = st_tcs[col_idx].find(w("tcPr"))
                if tcPr is not None:
                    tcW = tcPr.find(w("tcW"))
                    if tcW is not None:
                        try:
                            val = int(tcW.attrib.get(w("w"), "0"))
                            if val > 0:
                                w_val = val
                        except (ValueError, TypeError):
                            w_val = None
            # 3. Documented legacy compatibility fallback: only if cell lacks explicit XML width
            if w_val is None:
                s_name = str(name).lower()
                w_val = legacy_semantic_fallback_w_map.get(s_name, 200)
            summary_widths.append(w_val)

        return {
            "table_index": tbl_idx,
            "header_row0_index": header_row0_idx,
            "header_row1_index": header_row1_idx if header_row1_idx is not None else header_row0_idx,
            "student_template_row_index": student_template_row_idx,
            "no_col": no_col,
            "name_col": name_col,
            "id_col": id_col,
            "date_columns_start": date_columns_start,
            "summary_columns_count": summary_count,
            "summary_column_names": tuple(summary_names),
            "template_session_capacity": template_session_capacity,
            "template_student_row_capacity": max(student_rows_count, 1),
            "week_template_cell_col": week_template_cell_col,
            "summary_header0_cell_col": summary_header0_cell_col,
            "date_template_cell_col": date_template_cell_col,
            "summary_column_indices": summary_column_indices,
            "summary_header1_cell_cols": summary_column_indices,
            "student_date_template_cell_col": student_date_template_cell_col,
            "student_summary_cell_cols": summary_column_indices,
            "summary_column_widths": tuple(summary_widths),
            "row0_cell_count": len(row0_tcs),
            "row1_cell_count": len(r1_tcs),
            "student_row_cell_count": len(st_tcs),
            "matrix_row_count": len(rows),
        }

