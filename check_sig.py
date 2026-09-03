import openpyxl
wb = openpyxl.load_workbook("c:/Users/danjo/OneDrive/CVSU GENERATORS/Lecture and Lab.xlsx")
print("--- Laboratory ---")
for r in range(50, 65):
    val = wb["Laboratory"].cell(row=r, column=41).value # 41 is AO
    print(f"AO{r}: {val}")
print("--- Consolidated ---")
for r in range(50, 65):
    val = wb["Consolidated"].cell(row=r, column=10).value # 10 is J
    print(f"J{r}: {val}")
print("--- Lecture ---")
for r in range(50, 65):
    val = wb["Lecture"].cell(row=r, column=61).value # 61 is BI
    print(f"BI{r}: {val}")
