import openpyxl
wb = openpyxl.load_workbook("c:/Users/danjo/OneDrive/CVSU GENERATORS/Lecture only.xlsx", data_only=True)
for r in range(50, 65):
    for c in range(50, 75):
        val = wb["Lecture"].cell(row=r, column=c).value
        if val is not None:
            print(f"Lecture cell(row={r}, col={c}) ({wb['Lecture'].cell(row=r, column=c).coordinate}): {val}")
