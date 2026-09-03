import os
import re
import shutil
import openpyxl
import logging

logger = logging.getLogger(__name__)

class GradeGenerator:
    def __init__(self, templates_dir: str):
        self.templates_dir = templates_dir
        self.ll_template = os.path.join(templates_dir, "GRADING_LECTURE_LAB_TEMPLATE.xlsx")
        self.l_template = os.path.join(templates_dir, "GRADING_LECTURE_TEMPLATE.xlsx")

    def _parse_semester_and_year(self, raw_sem: str) -> tuple[str, str]:
        """
        Normalize semester and school year to match the standard format used in
        the dropdown validations on Sheet1:
        e.g. Semester: '1st Semester' or '2nd Semester'
             School Year: '2025-2026', '2026-2027', etc.
        """
        raw_sem = str(raw_sem).strip()
        sem_lower = raw_sem.lower()
        if "second" in sem_lower or "2nd" in sem_lower:
            sem_str = "2nd Semester"
        elif "first" in sem_lower or "1st" in sem_lower:
            sem_str = "1st Semester"
        elif "midyear" in sem_lower or "summer" in sem_lower:
            sem_str = "Midyear"
        else:
            sem_str = raw_sem

        year_match = re.search(r'(20\d{2})\s*[-–/]\s*(20\d{2})', raw_sem)
        if year_match:
            year_str = f"{year_match.group(1)}-{year_match.group(2)}"
        else:
            years = re.findall(r'20\d{2}', raw_sem)
            if len(years) >= 2:
                year_str = f"{years[0]}-{years[1]}"
            elif len(years) == 1:
                y = int(years[0])
                year_str = f"{y}-{y+1}"
            else:
                raise ValueError(f"Could not parse year from '{raw_sem}'")

        return sem_str, year_str

    def _parse_subject(self, raw_subject: str) -> tuple[str, str]:
        """
        Parse subject string into (subject_code, subject_title).
        e.g. 'DCIT 21A - INTRODUCTION TO COMPUTING' -> ('DCIT 21A', 'INTRODUCTION TO COMPUTING')
        """
        raw_subject = str(raw_subject).strip()
        if " - " in raw_subject:
            parts = raw_subject.split(" - ", 1)
            return parts[0].strip(), parts[1].strip()
        elif "-" in raw_subject:
            parts = raw_subject.split("-", 1)
            return parts[0].strip(), parts[1].strip()
        return raw_subject, ""

    def _normalize_course_section(self, raw_course: str) -> str:
        """
        Normalize course section:
        e.g. 'BSCS1-4' -> 'BSCS 1-4'
        """
        raw = str(raw_course).strip()
        norm = re.sub(r'([a-zA-Z]+)(\d)', r'\1 \2', raw)
        return re.sub(r'\s+', ' ', norm).strip()

    def generate(self, info: dict, students: list, output_path: str) -> bool:
        """
        info dictionary mapping:
            - "instructor": Instructor name
            - "course": Course & Section
            - "sched": Schedule code
            - "subject": Subject Code - Title
            - "semester": "FIRST SEMESTER, AY 2026 - 2027"
            - "time": Full time string
            - "has_lab": Optional boolean flag
            - "template_type": Optional ("ll" or "l")
            - "units": Optional int
        students: list of (name, student_number) tuples OR list of dicts
        output_path: Target .xlsx path
        """
        # 1. Determine Template
        subject_type = str(info.get("subject_type", "")).lower()
        template_type = str(info.get("template_type", "")).lower()
        time_str = str(info.get("time", "")).upper()
        subj_str = str(info.get("subject", "")).upper()
        has_lab_flag = info.get("has_lab")

        is_lab = False
        if subject_type:
            is_lab = ("lab" in subject_type)
        elif has_lab_flag is not None:
            is_lab = bool(has_lab_flag)
        elif "ll" in template_type or "lab" in template_type:
            is_lab = True
        elif "l" in template_type or "lec" in template_type:
            is_lab = False
        elif "LAB" in time_str or "LAB:" in time_str or "(LAB)" in subj_str or "LABORATORY" in subj_str:
            is_lab = True

        template_path = self.ll_template if is_lab else self.l_template
        start_row = 12 if is_lab else 11
        max_rows = 40 if is_lab else 60

        if not os.path.exists(template_path):
            print(f"  [ERROR]  Missing grade template: {template_path}")
            return False

        # 2. Parse metadata
        sem_str, year_str = self._parse_semester_and_year(info.get("semester", ""))
        subj_code, subj_title = self._parse_subject(info.get("subject", ""))
        course_str = self._normalize_course_section(info.get("course", ""))
        raw_sched = str(info.get("sched", "")).strip()
        sched_val = int(raw_sched) if raw_sched.isdigit() else raw_sched
        instructor = str(info.get("instructor", "")).strip().upper()

        # 3. Clean and unpack students
        cleaned_students = []
        for s in students:
            if isinstance(s, (list, tuple)):
                s_name = str(s[0]).strip() if len(s) > 0 and s[0] is not None else ""
                s_num = str(s[1]).strip() if len(s) > 1 and s[1] is not None else ""
            elif isinstance(s, dict):
                s_name = str(s.get("student_name") or s.get("name") or "").strip()
                s_num = str(s.get("student_number") or s.get("number") or "").strip()
            else:
                s_name = str(s).strip()
                s_num = ""
            if s_name:
                cleaned_students.append((s_name, s_num))

        # 4. Copy template to a temporary path for atomic writes
        os.makedirs(os.path.dirname(os.path.abspath(output_path)), exist_ok=True)
        import uuid
        tmp_path = output_path + f".{uuid.uuid4().hex[:8]}.tmp.xlsx"
        shutil.copy2(template_path, tmp_path)

        # 5. Populate workbook natively
        try:
            wb = openpyxl.load_workbook(tmp_path, data_only=False)
            
            sheet_names = [s.lower() for s in wb.sheetnames]
            name_map = {s.lower(): s for s in wb.sheetnames}

            # Populate 'Lecture' sheet
            if "lecture" not in sheet_names:
                raise ValueError(f"Grade template is missing required 'Lecture' sheet.")
            ws = wb[name_map["lecture"]]
            ws['C1'] = sched_val
            ws['M1'] = course_str
            ws['C2'] = subj_code
            ws['M2'] = sem_str
            ws['C3'] = subj_title
            ws['M3'] = year_str
            ws['M4'] = instructor
            if info.get("units"):
                try:
                    ws['C4'] = int(info["units"])
                except Exception:
                    pass

            # Update instructor signature cell in Lecture & Lab
            if is_lab:
                # Scan adjacent cells near BI57 (col 61, row 57) for 'Instructor'
                found_label = False
                for r_offset in range(-5, 6):
                    for c_offset in range(-20, 5):
                        val = ws.cell(row=57 + r_offset, column=61 + c_offset).value
                        if val and "instructor" in str(val).lower():
                            found_label = True
                            break
                    if found_label:
                        break
                if found_label:
                    ws['BI57'] = instructor
                else:
                    logger.warning(f"Could not find 'Instructor' anchor cell near BI57 for {course_str} ({sched_val}). Signature omitted.")

            # Inject active student roster without touching formulas in other columns
            total_slots = max(max_rows, len(cleaned_students))
            for r_idx in range(total_slots):
                row_num = start_row + r_idx
                if r_idx < len(cleaned_students):
                    name, num = cleaned_students[r_idx]
                    
                    def sanitize_excel(val):
                        if isinstance(val, str) and val.startswith(('=', '+', '-', '@')):
                            return "'" + val
                        return val
                        
                    ws.cell(row=row_num, column=1).value = r_idx + 1
                    ws.cell(row=row_num, column=2).value = sanitize_excel(name)
                    ws.cell(row=row_num, column=3).value = int(num) if (num.isdigit() and not num.startswith('0')) else num
                else:
                    # Clear remaining student name and number cells
                    ws.cell(row=row_num, column=2).value = None
                    ws.cell(row=row_num, column=3).value = None

            # Populate 'Laboratory' sheet if present
            if "laboratory" in sheet_names:
                ws_lab = wb[name_map["laboratory"]]
                ws_lab['AO59'] = instructor
    
            # Populate 'Consolidated' sheet if present
            if "consolidated" in sheet_names:
                ws_con = wb[name_map["consolidated"]]
                ws_con['J56'] = instructor
    
            # Populate 'Grading Sheet'
            if "grading sheet" not in sheet_names:
                raise ValueError(f"Grade template is missing required 'Grading Sheet' sheet.")
            ws_grd = wb[name_map["grading sheet"]]
            if ws_grd['A9'].value in (None, 'NAME OF COLLEGE'):
                ws_grd['A9'] = 'COLLEGE OF ENGINEERING AND INFORMATION TECHNOLOGY'
    
            wb.save(tmp_path)
            wb.close()
            os.replace(tmp_path, output_path)
            print(f"  [SUCCESS]  Grades generated at: {output_path}")
            return True
        except Exception as e:
            if os.path.exists(tmp_path):
                try:
                    os.remove(tmp_path)
                except Exception:
                    pass
            raise e
